#!/usr/bin/env python3
"""Tạo assets/style_template.xlsx từ một file chào giá TDI mẫu (đã làm xong).

Giữ nguyên theme màu, font, fill, viền, định dạng số của file mẫu:
  _TPL : dòng 1-12 (đầu sheet + tiêu đề bảng) chép nguyên, kèm các "dòng mẫu" style:
         13 section cấp 0 | 14 cấp 1 | 15 cấp 2 | 16 cấp 3 | 17 ghi chú | 18 ghi chú đậm
         19 hạng mục | 20 dòng tổng cộng
         (mỗi dòng mẫu = style xuất hiện nhiều nhất ở từng cột của loại dòng đó trong file mẫu)
  _SUM : sheet Summary chép nguyên.

Dùng khi TDI đổi mẫu màu:
  python3 make_style_template.py "<file mẫu>.xlsx" --sheet "B2_SERVICE " --out ../assets/style_template.xlsx
"""
import argparse
import collections
import copy
import sys
from pathlib import Path

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList

sys.path.insert(0, str(Path(__file__).resolve().parent))
from boq_tdi import detect_kind  # noqa: E402

ROW = {"sec0": 13, "sec1": 14, "sec2": 15, "sec3": 16, "note": 17, "note_bold": 18, "item": 19, "total": 20}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--sheet", default="B2_SERVICE ")
    ap.add_argument("--summary", default="Summary")
    ap.add_argument("--first-data-row", type=int, default=13)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.file)
    ws = wb[a.sheet]
    ncol = 55  # A..BC của template 3 khu

    # 1) Gom style theo loại dòng
    stats = collections.defaultdict(lambda: collections.defaultdict(collections.Counter))
    keep = {}
    total_row = None
    prev_letter = None
    for r in range(a.first_data_row, ws.max_row + 1):
        stt, desc, unit = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 7).value
        if desc is None and stt is None:
            continue
        if isinstance(desc, str) and desc.strip().lower().startswith("tổng cộng"):
            total_row = r
            continue
        qty = [ws.cell(r, c).value for c in (8, 9, 10)]
        qty = [q if isinstance(q, (int, float)) else None for q in qty]
        kind, level = detect_kind(stt, unit, qty, prev_letter)
        if kind == "section":
            key = f"sec{level}"
            if level == 0:
                prev_letter = str(stt).strip().rstrip(".")
        elif kind == "note":
            key = "note_bold" if ws.cell(r, 2).font.b else "note"
        else:
            key = "item"
        for c in range(1, ncol + 1):
            st = ws.cell(r, c)._style
            k = tuple(st)
            stats[key][c][k] += 1
            keep[k] = st

    def modal(key, c):
        cnt = stats[key][c]
        return copy.copy(keep[cnt.most_common(1)[0][0]]) if cnt else None

    fallback = {"sec0": ["sec1"], "sec3": ["sec2", "sec1"], "sec2": ["sec1"], "note_bold": ["note"],
                "note": ["item"]}
    styles = {}
    for key in ROW:
        if key == "total":
            continue
        chain = [key] + fallback.get(key, [])
        src = next((k for k in chain if stats.get(k)), "item")
        styles[key] = {c: modal(src, c) for c in range(1, ncol + 1)}
    if total_row:
        styles["total"] = {c: copy.copy(ws.cell(total_row, c)._style) for c in range(1, ncol + 1)}
    else:
        styles["total"] = styles["sec1"]

    # 2) Dọn workbook: chỉ giữ sheet mẫu + Summary
    for other in list(wb.worksheets):
        if other.title not in (a.sheet, a.summary):
            wb.remove(other)
    for n in list(wb.defined_names.keys()):
        del wb.defined_names[n]
    for s in wb.worksheets:
        for n in list(s.defined_names.keys()):
            del s.defined_names[n]
        s.data_validations.dataValidation = []
        s.conditional_formatting = ConditionalFormattingList()
        s.auto_filter.ref = None
    wb._external_links = []

    for m in [m for m in ws.merged_cells.ranges if m.min_row >= a.first_data_row]:
        ws.merged_cells.remove(m)
    ws.delete_rows(a.first_data_row, ws.max_row)
    for key, r in ROW.items():
        for c, st in styles[key].items():
            if st is not None:
                ws.cell(r, c)._style = st
        ws.cell(r, 2).value = key  # nhãn để dễ nhìn khi mở file mẫu
        ws.row_dimensions[r].height = None
    ws.title = "_TPL"
    wb[a.summary].title = "_SUM"
    found = {k: bool(stats.get(k)) for k in ROW if k != "total"}
    print("Loại dòng tìm thấy trong file mẫu:", found, "| dòng tổng:", total_row)
    wb.save(a.out)
    print("Đã ghi", a.out)


if __name__ == "__main__":
    main()
