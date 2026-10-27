#!/usr/bin/env python3
"""tv_loader.py - nạp Thư viện giá TDI (thư mục các file BG) vào file BOQ thành sheet "2 Thu vien gia", bảng Excel tên TV.

  python3 tv_loader.py LIBDIR BOQ.xlsx OUT.xlsx --code TDI25538 [--mtimes mtimes.json] [--sheet "2 Thu vien gia"]

- Chỉ lấy file .xlsx có tên bắt đầu bằng mã dự án (--code). Không có file nào -> dừng (mã dự án không khớp).
- Nhóm VT = phần sau ".BG" trong tên file, cắt "-", "_", khoảng trắng (cột tên "Column1" để tương thích công thức cũ).
- Đọc mọi sheet đang hiện; dòng tiêu đề = dòng đầu (trong 30 dòng) có "mã hiệu"; ghép tiêu đề 2 tầng.
- Giữ dòng có Mã hiệu và ĐVT; giá "-", 0, chữ -> trống. Ảnh chụp giá (giá trị), không liên kết ngược file nguồn.
- Thêm cột "File nguồn", "Ngày cập nhật" (mtime file trên máy người dùng, truyền qua --mtimes {tên file: mtimeMs}).
- Ghi bằng XML (không mở-lưu file BOQ bằng openpyxl). Sheet thêm vào cuối. Đã có bảng TV -> dừng, hỏi người dùng.
"""
import sys, os, re, glob, json, zipfile, datetime, argparse
import openpyxl
from lxml import etree
from openpyxl.utils import get_column_letter as CL
from xml.sax.saxutils import escape

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PR = 'http://schemas.openxmlformats.org/package/2006/relationships'
q = lambda t: '{%s}%s' % (NS, t)
COLS = ['Column1', 'Mã hiệu', 'Tên công việc', 'ĐVT', 'Giá VL gốc', 'Giá NC gốc', 'Giá NC ước tính',
        'Description', 'Ghi chú TV', 'Nhãn Hiệu Xuất Xứ', 'Sheet nguồn', 'File nguồn', 'Ngày cập nhật']
NUM = {'Giá VL gốc', 'Giá NC gốc', 'Giá NC ước tính'}


def norm(s): return re.sub(r'\s+', ' ', str(s or '')).strip().lower()


def std_name(top, sub):
    t, s = norm(top), norm(sub)
    if s == 'vật liệu' and ('số liệu gốc' in t or t == ''): return 'Giá VL gốc'
    if s == 'nhân công' and ('số liệu gốc' in t or t == ''): return 'Giá NC gốc'
    if 'nhân công' in t and 'ước' in t: return 'Giá NC ước tính'
    if 'mã hiệu' in t and 'mã hiệu' not in s: return 'Mã hiệu'
    if t == 'tên công việc': return 'Tên công việc'
    if t.startswith('unit') or 'đ. vị' in t or t == 'đvt': return 'ĐVT'
    if t == 'description': return 'Description'
    if t == 'ghi chú': return 'Ghi chú TV'
    if 'nhãn hiệu' in t and 'xuất xứ' in t: return 'Nhãn Hiệu Xuất Xứ'
    return None


def num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool) and v != 0: return float(v)
    return None


def group_of(fname):
    b = os.path.splitext(os.path.basename(fname))[0]
    m = re.search(r'\.BG(.*)$', b, re.I)
    return (m.group(1) if m else b).strip(' -_')


def load(libdir, code, mtimes):
    files = sorted(f for f in glob.glob(os.path.join(libdir, '*.xlsx'))
                   if not os.path.basename(f).startswith('~$') and os.path.basename(f).upper().startswith(code.upper()))
    if not files:
        raise SystemExit('DỪNG: không có file thư viện nào bắt đầu bằng mã "%s" trong %s. Hỏi người dùng.' % (code, libdir))
    rows, stat = [], []
    for f in files:
        fn = os.path.basename(f); g = group_of(fn)
        ms = mtimes.get(fn) or os.path.getmtime(f) * 1000
        day = datetime.datetime.fromtimestamp(ms / 1000).strftime('%d/%m/%Y')
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        n0 = len(rows)
        for ws in wb.worksheets:
            if ws.sheet_state != 'visible': continue
            data = list(ws.iter_rows(values_only=True))
            h = next((i for i, r in enumerate(data[:30]) if any('mã hiệu' in norm(v) for v in r)), None)
            if h is None: continue
            top, sub = list(data[h]), list(data[h + 1]) if h + 1 < len(data) else []
            last = ''
            colmap = {}
            for j in range(len(top)):
                if top[j] not in (None, ''): last = top[j]
                name = std_name(last if top[j] in (None, '') else top[j], sub[j] if j < len(sub) else None)
                if name and name not in colmap.values(): colmap[j] = name
            start = h + 2 if any(norm(v) in ('vật liệu', 'nhân công') for v in sub) else h + 1
            for r in data[start:]:
                rec = {colmap[j]: r[j] for j in colmap if j < len(r)}
                ma, dvt = rec.get('Mã hiệu'), rec.get('ĐVT')
                if ma in (None, '') or dvt in (None, '') or str(ma).strip() == '' or str(dvt).strip() == '': continue
                out = {'Column1': g, 'Sheet nguồn': ws.title, 'File nguồn': fn, 'Ngày cập nhật': day}
                for k, v in rec.items():
                    if k in NUM: out[k] = num(v)
                    elif v not in (None, ''): out[k] = str(v).strip() if k != 'Mã hiệu' else (str(v).strip() if not isinstance(v, float) or v != int(v) else str(int(v)))
                rows.append(out)
        stat.append((fn, g, len(rows) - n0, day))
    return rows, stat


def cell(ref, v):
    if v is None: return ''
    if isinstance(v, float): return '<c r="%s"><v>%r</v></c>' % (ref, v)
    return '<c r="%s" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (ref, escape(str(v)))


def add_sheet(src, out, rows, sheet_name):
    z = zipfile.ZipFile(src); files = {i.filename: z.read(i.filename) for i in z.infolist()}; infos = z.infolist()
    for n in files:
        if n.startswith('xl/tables/') and n.endswith('.xml') and re.search(rb'(name|displayName)="TV"', files[n]):
            raise SystemExit('DỪNG: file đã có bảng TV (%s). Hỏi người dùng có thay thư viện không.' % n)
    wb = etree.fromstring(files['xl/workbook.xml'])
    if any(s.get('name') == sheet_name for s in wb.find(q('sheets'))):
        raise SystemExit('DỪNG: đã có sheet "%s".' % sheet_name)
    n_ws = 1 + max([int(m.group(1)) for k in files for m in [re.match(r'xl/worksheets/sheet(\d+)\.xml$', k)] if m] or [0])
    n_tb = 1 + max([int(re.findall(r'table(\d+)\.xml$', k)[0]) for k in files if re.match(r'xl/tables/table\d+\.xml$', k)] or [0])
    ids = [int(t.get('id')) for k in files if re.match(r'xl/tables/table\d+\.xml$', k) for t in [etree.fromstring(files[k])]]
    tid = max(ids or [0]) + 1
    ncol, nrow = len(COLS), len(rows) + 1
    ref = 'A1:%s%d' % (CL(ncol), nrow)
    # sheet xml
    xs = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
          '<worksheet xmlns="%s" xmlns:r="%s"><dimension ref="%s"/>' % (NS, RNS, ref),
          '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>',
          '<sheetFormatPr defaultRowHeight="15"/><cols><col min="1" max="1" width="12" customWidth="1"/><col min="2" max="2" width="22" customWidth="1"/>'
          '<col min="3" max="3" width="45" customWidth="1"/><col min="4" max="8" width="13" customWidth="1"/><col min="9" max="%d" width="18" customWidth="1"/></cols><sheetData>' % ncol]
    xs.append('<row r="1">' + ''.join(cell(CL(j + 1) + '1', c) for j, c in enumerate(COLS)) + '</row>')
    for i, r in enumerate(rows, 2):
        xs.append('<row r="%d">' % i + ''.join(cell(CL(j + 1) + str(i), r.get(c)) for j, c in enumerate(COLS)) + '</row>')
    xs.append('</sheetData><pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
              '<tableParts count="1"><tablePart r:id="rIdTV1"/></tableParts></worksheet>')
    wsp = 'xl/worksheets/sheet%d.xml' % n_ws
    new = {wsp: ''.join(xs).encode('utf8')}
    new['xl/worksheets/_rels/sheet%d.xml.rels' % n_ws] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="%s">'
        '<Relationship Id="rIdTV1" Type="%s/table" Target="../tables/table%d.xml"/></Relationships>' % (PR, RNS, n_tb)).encode('utf8')
    new['xl/tables/table%d.xml' % n_tb] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><table xmlns="%s" id="%d" name="TV" displayName="TV" ref="%s">'
        '<autoFilter ref="%s"/><tableColumns count="%d">%s</tableColumns>'
        '<tableStyleInfo name="TableStyleLight9" showFirstColumn="0" showLastColumn="0" showRowStripes="1" showColumnStripes="0"/></table>'
        % (NS, tid, ref, ref, ncol, ''.join('<tableColumn id="%d" name="%s"/>' % (j + 1, escape(c)) for j, c in enumerate(COLS)))).encode('utf8')
    # workbook rels + sheets
    rels = etree.fromstring(files['xl/_rels/workbook.xml.rels'])
    rid = 'rIdTVs%d' % n_ws
    etree.SubElement(rels, '{%s}Relationship' % PR, Id=rid, Type=RNS + '/worksheet', Target='worksheets/sheet%d.xml' % n_ws)
    new['xl/_rels/workbook.xml.rels'] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
    shs = wb.find(q('sheets'))
    sid = 1 + max(int(s.get('sheetId')) for s in shs)
    s = etree.SubElement(shs, q('sheet')); s.set('name', sheet_name); s.set('sheetId', str(sid)); s.set('{%s}id' % RNS, rid)
    cp = wb.find(q('calcPr'))
    if cp is None: cp = etree.SubElement(wb, q('calcPr'))
    cp.set('fullCalcOnLoad', '1')
    new['xl/workbook.xml'] = etree.tostring(wb, xml_declaration=True, encoding='UTF-8', standalone=True)
    ct = files['[Content_Types].xml'].decode('utf8')
    add = ('<Override PartName="/%s" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' % wsp +
           '<Override PartName="/xl/tables/table%d.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml"/>' % n_tb)
    new['[Content_Types].xml'] = ct.replace('</Types>', add + '</Types>').encode('utf8')
    zo = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    for it in infos: zo.writestr(it, new.pop(it.filename, files[it.filename]))
    for k, v in new.items(): zo.writestr(k, v)
    zo.close()
    return ref


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('libdir'); ap.add_argument('boq'); ap.add_argument('out')
    ap.add_argument('--code', required=True); ap.add_argument('--mtimes'); ap.add_argument('--sheet', default='2 Thu vien gia')
    a = ap.parse_args()
    mt = json.load(open(a.mtimes)) if a.mtimes else {}
    rows, stat = load(a.libdir, a.code, mt)
    for fn, g, n, d in stat: print('%-40s nhóm %-12s %5d dòng  cập nhật %s' % (fn, g, n, d))
    ref = add_sheet(a.boq, a.out, rows, a.sheet)
    print('Đã nạp %d nhóm, %d dòng vào sheet "%s", bảng TV %s' % (len(stat), len(rows), a.sheet, ref))
