#!/usr/bin/env python3
"""boq_ncc_tools.py - bước "ghi giá báo giá NCC vào BOQ template TDI" (mọi dự án, mọi sheet). Bản 19/10/2026.

Tự chứa (không import boq_xml_tools). Ghi trực tiếp XML, KHÔNG dùng openpyxl để ghi file BOQ.

NGUYÊN TẮC AN TOÀN - chỉ 3 trường hợp được ghi vào ô Đơn giá gốc VL / NC:
  (1) ô TRỐNG hoàn toàn (không giá trị, không công thức)  -> ghi số;
  (2) ô có CÔNG THỨC nhưng kết quả hiện tại = 0 hoặc rỗng (thư viện có mã nhưng chưa có giá)
      -> giữ công thức, bọc lại:  =IF(N((cũ))=0, <giá NCC>, (cũ))   => khi thư viện có giá > 0, giá thư viện luôn thắng;
  (3) ô VL đang là GIÁ THAM KHẢO WEB do skill boq-gia-tham-khao-web ghi (số cứng + cột `Nguồn giá` = "Giá tham khảo"
      + ghi chú gán mã chứa "GIÁ THAM KHẢO WEB") -> thay bằng giá NCC, đổi nền xanh sang cam, `Nguồn giá` = "NCC",
      ghi chú đổi thành "GIÁ WEB CŨ (đã thay bằng giá NCC): ..." (giữ nguyên số liệu cũ để truy vết).
Ô đã có giá > 0 khác (số nhập tay, công thức thư viện, công thức nhập tay) KHÔNG BAO GIỜ bị ghi.
Dòng đã có giá MỘT thành phần (ví dụ có NC, thiếu VL): thành phần đã có giữ nguyên; thành phần còn thiếu chỉ được điền khi ASSIGN.json có giá của đúng thành phần đó lấy từ báo giá NCC (không tự suy từ thành phần kia). Số 0 gõ tay: bỏ qua (trừ khi --ghi-so-0).
Mỗi mục ASSIGN.json BẮT BUỘC đủ nguồn: ncc, file, date (YYYY-MM-DD hoặc dd/mm/yyyy), ledger_id của từng thành phần, note có vị trí trong file.
Lệnh `audit` chứng minh điều đó.

Lệnh:
  sheets  FILE [--map MAP.json]            : phân loại sheet (có nhận được cột TDI hay không), in cột đã nhận + thống kê trạng thái ô giá
  todo    FILE [SHEET...] [--map MAP.json] [--out todo.json]
                                           : dòng đủ điều kiện nhận giá NCC (trống, công thức 0, hoặc đang là giá tham khảo web) + thống kê
  ledger-init OUT.xlsx                     : tạo Sổ cái báo giá NCC rỗng
  ledger-check LEDGER.xlsx                 : kiểm tra sổ cái (thiếu trường, VAT không rõ, hết hạn, ảnh/email, trùng); cột theo tiêu đề, cột tùy chọn thiếu thì bỏ qua
  apply   FILE OUT ASSIGN.json [--map MAP.json] [--dry] [--ghi-so-0]
                                           : ghi giá NCC + 4 cột nguồn + cập nhật `Nguồn giá`; --dry chỉ báo, không ghi
  audit   BEFORE AFTER [--map MAP.json]    : so từng ô; chỉ cho phép 3 trường hợp ghi ở trên + cột nguồn NCC + `Nguồn giá` / ghi chú gán mã của đúng dòng đã ghi

MAP.json (chỉ khi sheet có bố cục lạ, tiêu đề không nhận được): {"<sheet>": {"name":"R","unit":"U","vl":"V","nc":"W",
   "hdr_row":11, "code":"S", "kl":"AA"}}   (name/unit/vl/nc bắt buộc; hdr_row = dòng tiêu đề cuối, dữ liệu bắt đầu sau dòng này)

ASSIGN.json: [{"sheet":..., "row":N, "desc_prefix":"<vài ký tự đầu Nội dung>",
               "vl":{"price":123456, "ledger_id":"BG-0012"},   # tùy chọn
               "nc":{"price":..., "ledger_id":"..."},          # tùy chọn
               "ncc":"Taisin", "date":"2025-11-14", "file":"<đường dẫn tương đối file báo giá>",
               "status":"Hiệu lực|Cần kiểm tra|Hết hạn", "calc":"<diễn giải chiết tính, bắt buộc với ĐVT lô/bộ/HT/tấn/10m/100m>",
               "note":"<sheet NCC>!dòng <n>; STT YCBG <x>; đầu mục NCC: ... (vị trí trong file báo giá là bắt buộc)"}]
"""
import zipfile, re, json, sys, os, copy, collections, datetime
from lxml import etree
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.formula.translate import Translator

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
q = lambda t: '{%s}%s' % (NS, t)
SRC_COLS = ['Nguồn giá NCC', 'Ngày BG NCC', 'Trạng thái giá NCC', 'Chiết tính NCC']
HELPERS = ['Nhóm VT', 'Độ tin cậy', 'Ghi chú gán mã', 'Hệ số ĐV', 'Loại dòng']
ORIGIN_COL = 'Nguồn giá'                       # cột do boq-gan-ma-gia-thu-vien (6c) tạo; web skill và bước này cập nhật
WEB_MARK = 'GIÁ THAM KHẢO WEB'                 # mở đầu ghi chú do boq-gia-tham-khao-web ghi
WEB_OLD = 'GIÁ WEB CŨ (đã thay bằng giá NCC)'
ORIGIN_KNOWN = ('', 'chưa có giá', 'giá tham khảo', 'giá thư viện')
FILL_NCC = 'FFFCE4D6'   # cam nhạt: ô giá lấy từ NCC
SPECIAL_UNITS = {'lô', 'lot', 'bộ', 'ht', 'hệ thống', 'gói', 'tấn', 'tan', '10m', '100m', '10 m', '100 m',
                 'tổ', 'cụm', 'hệ', 'ls', 'l.s', 'lumpsum', 'lump sum', 'set', 'bo'}
LEDGER_COLS = ['ID', 'Hệ', 'Nhóm VT', 'NCC', 'Ngày BG', 'Hiệu lực đến', 'Loại file', 'STT YCBG', 'Đầu mục nguyên văn',
               'Tên hàng', 'Hãng / Model / Quy cách', 'ĐVT',
               'Đơn giá', 'VAT (Chưa/Đã/Không rõ)', '% VAT', 'Giá chưa VAT', 'Loại giá (VL/NC/VL+NC)',
               'Chiết khấu / ghi chú giá', 'Điều kiện giao hàng', 'File nguồn', 'Vị trí trong file',
               'Độ tin cậy trích', 'Trạng thái', 'Ghi chú']
LEDGER_FILE_TYPES = ('YCBG điền giá', 'Báo giá NCC')
WRAP_RE = re.compile(r'^=IF\(N\(\((.*)\)\)=0,([0-9.eE+\-]+),\((.*)\)\)$', re.S)
POS_RE = re.compile(r'dòng\s*\d+|trang\s*\d+|email|ô\s*[A-Z]+\d+', re.I)


# ------------------------------------------------------------------ package I/O
class Pkg:
    def __init__(self, path):
        z = zipfile.ZipFile(path)
        self.infos = z.infolist()
        self.files = {i.filename: z.read(i.filename) for i in self.infos}
        self.new = {}
        self.wb = etree.fromstring(self.files['xl/workbook.xml'])
        rels = etree.fromstring(self.files['xl/_rels/workbook.xml.rels'])
        rid = {r.get('Id'): r.get('Target') for r in rels}
        self.sheets = []
        for s in self.wb.find(q('sheets')):
            t = rid.get(s.get('{%s}id' % RNS))
            if t is None: continue
            t = t.lstrip('/'); t = t if t.startswith('xl/') else 'xl/' + t
            self.sheets.append((s.get('name'), t))
        self.path = dict(self.sheets)
        self.ss = []
        if 'xl/sharedStrings.xml' in self.files:
            sst = etree.fromstring(self.files['xl/sharedStrings.xml'])
            self.ss = [''.join(t.text or '' for t in si.iter(q('t'))) for si in sst]
        self._roots, self._masters, self._rows = {}, {}, {}

    def root(self, sheet):
        p = self.path[sheet]
        if p not in self._roots:
            self._roots[p] = etree.fromstring(self.new.get(p) or self.files[p])
        return self._roots[p]

    def rows(self, sheet):
        if sheet not in self._rows:
            self._rows[sheet] = {int(r.get('r')): r for r in self.root(sheet).find(q('sheetData'))}
        return self._rows[sheet]

    def touch(self, sheet):
        self.new[self.path[sheet]] = None

    def masters(self, sheet):
        if sheet not in self._masters:
            m = {}
            for row in self.root(sheet).find(q('sheetData')):
                for c in row:
                    f = c.find(q('f'))
                    if f is not None and f.get('t') == 'shared' and f.text:
                        m[f.get('si')] = (f.text, c.get('r'))
            self._masters[sheet] = m
        return self._masters[sheet]

    def formula(self, sheet, c):
        """Công thức đã bung (không còn shared), không có dấu '='. None nếu ô không có công thức. Lỗi dịch -> '<LOI>'."""
        f = c.find(q('f'))
        if f is None: return None
        if f.get('t') == 'shared' and not f.text:
            m = self.masters(sheet).get(f.get('si'))
            if not m: return '<LOI>'
            try: return Translator('=' + m[0], origin=m[1]).translate_formula(c.get('r'))[1:]
            except Exception: return '<LOI>'
        return f.text or ''

    def val(self, c, sheet=None):
        t = c.get('t'); v = c.find(q('v'))
        f = c.find(q('f'))
        if f is not None and sheet is not None: return '=' + (self.formula(sheet, c) or '')
        if f is not None: return '=' + (f.text or '<shared>')
        if t == 's' and v is not None: return self.ss[int(v.text)]
        if t == 'inlineStr': return ''.join(x.text or '' for x in c.iter(q('t')))
        return v.text if v is not None else None

    def grid(self, sheet, r1=1, r2=10 ** 7):
        out = {}
        for r, row in self.rows(sheet).items():
            if r1 <= r <= r2:
                for c in row: out[c.get('r')] = self.val(c, sheet)
        return out

    def save(self, out):
        for p, rt in self._roots.items():
            if p in self.new:
                self.new[p] = etree.tostring(rt, xml_declaration=True, encoding='UTF-8', standalone=True)
        cp = self.wb.find(q('calcPr'))
        if cp is None: cp = etree.SubElement(self.wb, q('calcPr'))
        cp.set('fullCalcOnLoad', '1')
        self.new['xl/workbook.xml'] = etree.tostring(self.wb, xml_declaration=True, encoding='UTF-8', standalone=True)
        drop = {'xl/calcChain.xml'}
        rels = etree.fromstring(self.files['xl/_rels/workbook.xml.rels'])
        for r in list(rels):
            if 'calcChain' in r.get('Target'): rels.remove(r)
        self.new['xl/_rels/workbook.xml.rels'] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
        ct = self.files['[Content_Types].xml'].decode('utf8')
        self.new['[Content_Types].xml'] = re.sub(r'<Override[^>]*calcChain[^>]*/>', '', ct).encode('utf8')
        z = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
        for it in self.infos:
            if it.filename in drop: continue
            z.writestr(it, self.new.get(it.filename) or self.files[it.filename])
        z.close()


def colnum(ref): return CI(re.match(r'[A-Z]+', ref).group(0))


def setc(row, col, val=None, num=None, formula=None, style=None):
    ref = CL(col) + row.get('r'); cell = None; pos = len(row)
    for i, c in enumerate(row):
        cc = colnum(c.get('r'))
        if cc == col: cell = c; break
        if cc > col: pos = i; break
    if cell is None:
        cell = etree.Element(q('c')); cell.set('r', ref); row.insert(pos, cell)
    s = cell.get('s') if style is None else str(style)
    for ch in list(cell): cell.remove(ch)
    for a in list(cell.attrib):
        if a != 'r': del cell.attrib[a]
    if s: cell.set('s', s)
    if formula is not None: etree.SubElement(cell, q('f')).text = formula
    elif num is not None: etree.SubElement(cell, q('v')).text = repr(num)
    elif val not in (None, ''):
        cell.set('t', 'inlineStr')
        etree.SubElement(etree.SubElement(cell, q('is')), q('t')).text = str(val)
    return cell


def cell_of(row, col):
    if row is None or not col: return None
    ref = CL(col)
    for c in row:
        if re.match(r'[A-Z]+', c.get('r')).group(0) == ref: return c
    return None


# ------------------------------------------------------------------ nhận diện cột (theo tiêu đề, có MAP ghi đè)
def norm(s): return re.sub(r'\s+', ' ', str(s or '')).strip().lower()


def detect_cols(pkg, sheet, override=None):
    ov = (override or {}).get(sheet)
    if ov:
        for k in ('name', 'unit', 'vl', 'nc'):
            if k not in ov: raise SystemExit('MAP của sheet %s thiếu khóa "%s"' % (sheet, k))
        res = {'name': CI(ov['name']), 'unit_tdi': CI(ov['unit']), 'vl0': CI(ov['vl']), 'nc0': CI(ov['nc']),
               'code': CI(ov['code']) if ov.get('code') else None, 'kl': CI(ov['kl']) if ov.get('kl') else None,
               'hdr_row': int(ov.get('hdr_row', 11)), 'from_map': True}
        g = pkg.grid(sheet, 1, res['hdr_row'] + 2)
    else:
        g = pkg.grid(sheet, 1, 25)
        by = collections.defaultdict(list)
        for ref, v in g.items():
            if v is None or str(v).startswith('='): continue
            by[norm(v)].append((int(re.search(r'\d+', ref).group(0)), colnum(ref)))
        cands = sorted(by.get('nội dung công việc', []), key=lambda x: (x[0], x[1]))
        if not cands: raise SystemExit('Không thấy cột "Nội Dung Công Việc" (phần TDI) trong %s' % sheet)
        hr, name = None, None
        for r, c in cands:   # chọn tiêu đề có "Đơn giá gốc" ở bên phải cùng dòng
            if any(rr == r and name_c > c and name_c - c <= 12 for rr, name_c in by.get('đơn giá gốc', [])):
                hr, name = r, c; break
        if name is None: raise SystemExit('Sheet %s không có tiêu đề "Đơn giá gốc" cạnh "Nội Dung Công Việc" (không phải sheet BOQ TDI)' % sheet)
        near = lambda key, lo, hi, rows=(hr,): next(iter(sorted(c for k, lst in by.items() if key(k) for r, c in lst
                                                                if r in rows and lo <= c - name <= hi)), None)
        vl = near(lambda k: k == 'đơn giá gốc', 1, 12)
        unit = near(lambda k: k in ('đơn vị', 'đvt', 'đơn vị tính'), 1, 8) or name + 3
        code = near(lambda k: 'mã hiệu' in k, 1, 6) or name + 1
        klc = near(lambda k: k in ('tổng kl tạm tính', 'tổng khối lượng', 'tổng kl'), 4, 80, rows=(hr, hr + 1))
        res = {'name': name, 'code': code, 'unit_tdi': unit, 'vl0': vl, 'nc0': vl + 1, 'kl': klc, 'hdr_row': hr}
        # dòng phụ ngay dưới tiêu đề (Vật liệu / Nhân công) nằm trong vùng tiêu đề
        sub = {norm(g.get('%s%d' % (CL(vl), hr + 1))), norm(g.get('%s%d' % (CL(vl + 1), hr + 1)))}
        if sub & {'vật liệu', 'nhân công'}: res['hdr_row'] = hr + 1
    rowh = res['hdr_row']
    for h in HELPERS + SRC_COLS + [ORIGIN_COL]:
        res[h] = None
        for ref, v in g.items():
            if v and not str(v).startswith('=') and norm(v) == norm(h) and colnum(ref) > res['name'] \
                    and int(re.search(r'\d+', ref).group(0)) <= rowh:
                res[h] = colnum(ref); break
    res['title_row'] = rowh
    return res


def letters(C): return {k: CL(v) for k, v in C.items() if isinstance(v, int) and k not in ('hdr_row', 'title_row')}


def price_state(pkg, sheet, c):
    """('empty'|'priced'|'zero_formula'|'zero_const'|'unknown'|'text', số)"""
    if c is None: return 'empty', None
    f = c.find(q('f')); v = c.find(q('v')); t = c.get('t')
    if f is not None:
        if t == 'e': return 'unknown', None
        if t in ('str', 'inlineStr'):
            s = (v.text if v is not None else '') or ''
            return ('zero_formula', 0) if s.strip() == '' else ('unknown', None)
        if v is None or v.text is None: return 'unknown', None
        try: n = float(v.text)
        except ValueError: return 'unknown', None
        return ('zero_formula', 0) if n == 0 else ('priced', n)
    val = pkg.val(c)
    if val is None or str(val).strip() == '': return 'empty', None
    if t in ('s', 'inlineStr', 'str'): return 'text', None
    try: n = float(val)
    except ValueError: return 'text', None
    return ('zero_const', 0) if n == 0 else ('priced', n)


def cell_text(pkg, row, col):
    c = cell_of(row, col)
    return '' if c is None else str(pkg.val(c) or '')


def cached_text(pkg, c):
    """Chữ hiển thị của ô: ô công thức (vd =B5) lấy kết quả đã tính <v>, không lấy công thức."""
    if c is None: return ''
    if c.find(q('f')) is not None:
        v = c.find(q('v')); return '' if v is None or v.text is None else str(v.text)
    return str(pkg.val(c) or '')


def is_web_ref(pkg, sh, C, row):
    """Ô VL đang là giá tham khảo web: số cứng > 0 + `Nguồn giá` = Giá tham khảo + ghi chú gán mã chứa WEB_MARK."""
    if row is None or not C.get(ORIGIN_COL) or not C.get('Ghi chú gán mã'): return False
    c = cell_of(row, C['vl0'])
    if c is None or c.find(q('f')) is not None: return False
    if price_state(pkg, sh, c)[0] != 'priced': return False
    return norm(cell_text(pkg, row, C[ORIGIN_COL])) == 'giá tham khảo' and WEB_MARK in cell_text(pkg, row, C['Ghi chú gán mã'])


def comp_state(pkg, sh, C, row, key):
    c = cell_of(row, C['vl0'] if key == 'vl' else C['nc0'])
    s = price_state(pkg, sh, c)[0]
    if key == 'vl' and s == 'priced' and is_web_ref(pkg, sh, C, row): return 'web_ref'
    return s


def new_origin(cur, web_replaced):
    cur = (cur or '').strip(); n = norm(cur)
    if n in ('', 'chưa có giá'): return 'NCC'
    if n == 'giá tham khảo': return 'NCC' if web_replaced else 'Giá tham khảo + NCC'
    if 'ncc' in n: return cur
    return cur + ' + NCC'


def unit_special(u): return norm(u) in SPECIAL_UNITS


def load_map(a):
    if '--map' in a: return json.load(open(a[a.index('--map') + 1]))
    return None


def fmt_price(x):
    x = float(x)
    return str(int(x)) if x == int(x) and abs(x) < 1e15 else repr(x)


def boq_sheets(pkg, override=None, only=None):
    out = []
    for sh, _ in pkg.sheets:
        if only and sh not in only: continue
        try: out.append((sh, detect_cols(pkg, sh, override), None))
        except SystemExit as e: out.append((sh, None, str(e)))
    return out


def data_rows(pkg, sh, C):
    """Duyệt dòng dữ liệu: có Nội dung và ĐVT, sau dòng tiêu đề."""
    rows = pkg.rows(sh)
    for r in sorted(rows):
        if r <= C['hdr_row']: continue
        row = rows[r]
        n, u = cell_of(row, C['name']), cell_of(row, C['unit_tdi'])
        name = pkg.val(n, sh) if n is not None else None
        unit = pkg.val(u, sh) if u is not None else None
        if not str(name or '').strip() or not str(unit or '').strip(): continue
        if str(name).startswith('=') and not str(unit).strip(): continue
        yield r, row, str(name), str(unit)


# ------------------------------------------------------------------ sheets / todo
def stats_for(pkg, sh, C):
    st = collections.Counter()
    for r, row, name, unit in data_rows(pkg, sh, C):
        lc = cell_of(row, C['Loại dòng']) if C.get('Loại dòng') else None
        if lc is not None and pkg.val(lc) == 'Cụm': st['dong_cum'] += 1; continue
        st['dong_cong_viec'] += 1
        for lab, key in (('VL', 'vl'), ('NC', 'nc')):
            st['%s_%s' % (lab, comp_state(pkg, sh, C, row, key))] += 1
    return dict(st)


def cmd_sheets(f, override):
    pkg = Pkg(f)
    for sh, C, err in boq_sheets(pkg, override):
        if C is None: print('BỎ QUA  %-34s %s' % (sh, err[:90])); continue
        helpers = all(C[h] for h in HELPERS)
        print('BOQ     %-34s cột %s | cột phụ gán mã: %s | cột Nguồn giá: %s | %s' % (
            sh, letters(C), 'có' if helpers else 'CHƯA CÓ', 'có' if C.get(ORIGIN_COL) else 'CHƯA CÓ', stats_for(pkg, sh, C)))


def cmd_todo(f, sheets, outfile, override):
    pkg = Pkg(f); res, summary = [], {}
    for sh, C, err in boq_sheets(pkg, override, sheets or None):
        if C is None: summary[sh] = {'bo_qua': err}; continue
        summary[sh] = {'cot': letters(C), 'co_cot_gan_ma': all(C[h] for h in HELPERS), 'co_cot_nguon_gia': bool(C.get(ORIGIN_COL)),
                       **stats_for(pkg, sh, C)}
        for r, row, name, unit in data_rows(pkg, sh, C):
            g = lambda col: (lambda c: pkg.val(c, sh) if c is not None else None)(cell_of(row, col)) if col else None
            loai = g(C.get('Loại dòng'))
            if loai == 'Cụm': continue
            sv, sn = comp_state(pkg, sh, C, row, 'vl'), comp_state(pkg, sh, C, row, 'nc')
            thieu = [lab for lab, s in (('VL', sv), ('NC', sn)) if s in ('empty', 'zero_formula', 'web_ref')]
            if not thieu: continue
            kl = None
            if C.get('kl'):
                cc = cell_of(row, C['kl'])
                try: kl = float(pkg.val(cc)) if cc is not None and cc.find(q('v')) is not None else None
                except (TypeError, ValueError): kl = None
            res.append({'sheet': sh, 'row': r, 'name': name[:140], 'unit': unit, 'ma': g(C['code']),
                        'nhom': g(C.get('Nhóm VT')), 'conf': g(C.get('Độ tin cậy')), 'loai': loai,
                        'thieu': '+'.join(thieu), 'state_vl': sv, 'state_nc': sn,
                        'dang_la_gia_tham_khao_web': sv == 'web_ref',
                        'co_gia_mot_phan': 'priced' in (sv, sn), 'kl': kl,
                        'lo_dac_biet': unit_special(unit), 'ghi_chu': g(C.get('Ghi chú gán mã'))})
    out = {'summary': summary, 'rows': res}
    s = json.dumps(out, ensure_ascii=False, indent=1)
    if outfile:
        open(outfile, 'w').write(s)
        print('Ghi %d dòng nhận được giá NCC (trong đó %d dòng đang là giá tham khảo web, %d dòng đã có giá một thành phần) -> %s' % (
            len(res), sum(1 for x in res if x['dang_la_gia_tham_khao_web']), sum(1 for x in res if x['co_gia_mot_phan']), outfile))
        for sh, d in summary.items(): print(' ', sh, d if 'bo_qua' not in d else 'BỎ QUA')
    else: print(s)


# ------------------------------------------------------------------ apply
def ensure_src_cols(pkg, sh, C):
    if all(C[h] for h in SRC_COLS): return {h: C[h] for h in SRC_COLS}
    rows = pkg.rows(sh)
    last = max(colnum(x.get('r')) for r in rows.values() for x in r)    # sau cột có dữ liệu xa nhất của cả sheet
    hdr = rows[C['title_row']]; out = {}
    for h in SRC_COLS:
        if C[h]: out[h] = C[h]; continue
        last += 1; setc(hdr, last, h); out[h] = last
    return out


def check_source(a):
    """Danh sách trường nguồn còn thiếu / sai (nguyên tắc 5b). Rỗng = đủ."""
    miss = []
    for k in ('ncc', 'file', 'date'):
        if not str(a.get(k) or '').strip(): miss.append(k)
    ds = str(a.get('date') or '').strip()
    if ds:
        for fm in ('%Y-%m-%d', '%d/%m/%Y'):
            try: datetime.datetime.strptime(ds, fm); break
            except ValueError: pass
        else: miss.append('date không đúng dạng YYYY-MM-DD hoặc dd/mm/yyyy')
    comps = [k for k in ('vl', 'nc') if a.get(k)]
    if not comps: miss.append('không có vl/nc')
    for k in comps:
        if not str(a[k].get('ledger_id') or '').strip(): miss.append('%s.ledger_id' % k)
    if not POS_RE.search(str(a.get('note') or '')): miss.append('note thiếu vị trí trong file báo giá (sheet!dòng n / trang n / email)')
    return miss


def cmd_apply(f, out, afile, override, dry=False, allow_zero_const=False):
    A = json.load(open(afile)); pkg = Pkg(f)
    by = collections.defaultdict(list)
    for a in A: by[a['sheet']].append(a)
    st = etree.fromstring(pkg.files['xl/styles.xml'])
    fills, xfs = st.find(q('fills')), st.find(q('cellXfs'))
    fl = etree.SubElement(fills, q('fill')); pf = etree.SubElement(fl, q('patternFill'), patternType='solid')
    etree.SubElement(pf, q('fgColor'), rgb=FILL_NCC); etree.SubElement(pf, q('bgColor'), indexed='64')
    fills.set('count', str(len(fills))); FID = len(fills) - 1; xc = {}

    def xf(s):
        if s in xc: return xc[s]
        x = copy.deepcopy(xfs[int(s or 0)]); x.set('fillId', str(FID)); x.set('applyFill', '1')
        xfs.append(x); xfs.set('count', str(len(xfs))); xc[s] = len(xfs) - 1; return xc[s]

    def unshare_group(sh, si):
        for row in pkg.root(sh).find(q('sheetData')):
            for c in row:
                fx = c.find(q('f'))
                if fx is not None and fx.get('t') == 'shared' and fx.get('si') == si:
                    txt = pkg.formula(sh, c)
                    for a2 in ('t', 'si', 'ref'):
                        if a2 in fx.attrib: del fx.attrib[a2]
                    fx.text = txt
        pkg.masters(sh).pop(si, None)

    ok, skipped, nwrap, nweb = 0, [], 0, 0
    for sh, items in by.items():
        if sh not in pkg.path: skipped.append((sh, 0, 'không có sheet này')); continue
        C = detect_cols(pkg, sh, override); rows = pkg.rows(sh)
        S = ensure_src_cols(pkg, sh, C) if not dry else {h: 0 for h in SRC_COLS}
        for a in items:
            r = a['row']; row = rows.get(r)
            skip = lambda why: skipped.append((sh, r, why))
            if row is None: skip('không có dòng'); continue
            miss = check_source(a)
            if miss: skip('thiếu nguồn, KHÔNG ghi: ' + '; '.join(miss)); continue
            nc_ = cell_of(row, C['name']); name = cached_text(pkg, nc_)
            if a.get('desc_prefix') and not norm(name).startswith(norm(a['desc_prefix'])):   # so không phân biệt hoa/thường, khoảng trắng
                skip('lệch nội dung (số dòng đã đổi?): thấy "%s"' % str(name)[:40]); continue
            lc = cell_of(row, C['Loại dòng']) if C.get('Loại dòng') else None
            if lc is not None and pkg.val(lc) == 'Cụm': skip('dòng Cụm: ghi vào dòng con'); continue
            uc = cell_of(row, C['unit_tdi']); unit = pkg.val(uc, sh) if uc is not None else ''
            if unit_special(unit) and not str(a.get('calc') or '').strip():
                skip('ĐVT "%s" cần chiết tính (thiếu trường calc)' % unit); continue
            wrote, web_replaced, web_old = [], False, None
            for key, col in (('vl', C['vl0']), ('nc', C['nc0'])):
                if not a.get(key): continue
                c = cell_of(row, col); stt = comp_state(pkg, sh, C, row, key)
                price = a[key].get('price')
                where = '%s%d' % (CL(col), r)
                if not isinstance(price, (int, float)) or isinstance(price, bool) or price <= 0:
                    skip('giá %s không hợp lệ: %r' % (key.upper(), price)); continue
                if stt == 'priced': skip('ô %s đã có giá (thư viện/nhập tay) - KHÔNG ghi' % where); continue
                if stt == 'text': skip('ô %s chứa chữ - KHÔNG ghi' % where); continue
                if stt == 'unknown': skip('ô %s là công thức chưa có kết quả tính (mở file bằng Excel, lưu, rồi chạy lại) - KHÔNG ghi' % where); continue
                if stt == 'zero_const' and not allow_zero_const: skip('ô %s đang là số 0 gõ tay - bỏ qua (dùng --ghi-so-0 nếu đúng ý)' % where); continue
                s0 = c.get('s') if c is not None else None
                if stt in ('empty', 'zero_const'):
                    if not dry: setc(row, col, num=float(price), style=xf(s0))
                elif stt == 'web_ref':   # trường hợp 3: thay giá tham khảo web bằng giá NCC (xf mới chép từ xf cũ, đổi nền xanh -> cam)
                    web_old = price_state(pkg, sh, c)[1]
                    if not dry: setc(row, col, num=float(price), style=xf(s0))
                    web_replaced = True; nweb += 1
                else:   # zero_formula: bọc công thức, giữ công thức thư viện
                    old = pkg.formula(sh, c)
                    if old in (None, '<LOI>'): skip('ô %s: không bung được công thức dùng chung' % where); continue
                    fx = c.find(q('f'))
                    if fx.get('t') == 'shared' and fx.text and fx.get('si') is not None and not dry:
                        unshare_group(sh, fx.get('si')); old = pkg.formula(sh, c)
                    if not dry: setc(row, col, formula='IF(N((%s))=0,%s,(%s))' % (old, fmt_price(price), old), style=xf(s0))
                    nwrap += 1
                wrote.append(key)
            if not wrote: continue
            if not dry:
                src = 'NCC %s | %s' % (a.get('ncc'), a.get('file'))
                ids = ','.join(dict.fromkeys(a[k]['ledger_id'] for k in ('vl', 'nc') if a.get(k) and k in wrote and a[k].get('ledger_id')))
                if ids: src += ' | ' + ids
                setc(row, S['Nguồn giá NCC'], src); setc(row, S['Ngày BG NCC'], a.get('date', ''))
                setc(row, S['Trạng thái giá NCC'], a.get('status', 'Cần kiểm tra') + (' | ' + a['note'] if a.get('note') else ''))
                if a.get('calc'): setc(row, S['Chiết tính NCC'], a['calc'])
                if C.get(ORIGIN_COL):
                    setc(row, C[ORIGIN_COL], new_origin(cell_text(pkg, row, C[ORIGIN_COL]), web_replaced))
                if web_replaced and C.get('Ghi chú gán mã'):
                    note = cell_text(pkg, row, C['Ghi chú gán mã'])
                    setc(row, C['Ghi chú gán mã'], note.replace(WEB_MARK, WEB_OLD, 1))
            ok += 1
        if not dry: pkg.touch(sh)
    print('%s: %d dòng được ghi (trong đó %d ô công thức được bọc IF, %d ô thay giá tham khảo web), %d mục bỏ qua' % (
        'DRY-RUN' if dry else 'GHI', ok, nwrap, nweb, len(skipped)))
    for s in skipped: print('  BỎ QUA %s!%s: %s' % s)
    if not dry:
        pkg.new['xl/styles.xml'] = etree.tostring(st, xml_declaration=True, encoding='UTF-8', standalone=True)
        pkg.save(out); print('->', out, '(chạy `audit` trước khi giao; nhắc đóng/mở lại file)')


# ------------------------------------------------------------------ audit
def cmd_audit(fb, fa, override):
    B, A = Pkg(fb), Pkg(fa); bad, changed, wrapped, webrep = [], 0, 0, 0
    if [n for n, _ in B.sheets] != [n for n, _ in A.sheets]: print('LỖI: danh sách sheet khác nhau'); sys.exit(1)
    for sh, _ in B.sheets:
        gb, ga = B.grid(sh), A.grid(sh)
        try: Ca = detect_cols(A, sh, override); Cb = detect_cols(B, sh, override)
        except SystemExit: Ca = Cb = None
        srcc = {Ca[h] for h in SRC_COLS if Ca[h]} if Ca else set()
        diffs = [ref for ref in set(gb) | set(ga) if (gb.get(ref) or '') != (ga.get(ref) or '')]
        rowof = lambda ref: int(re.search(r'\d+', ref).group(0))
        # dòng đã được ghi giá NCC = dòng có ô "Nguồn giá NCC" đổi từ trống sang có chữ
        touched = set()
        if Ca and Ca.get('Nguồn giá NCC'):
            touched = {rowof(ref) for ref in diffs if colnum(ref) == Ca['Nguồn giá NCC'] and (gb.get(ref) in (None, ''))}
        for ref in diffs:
            vb, va = gb.get(ref), ga.get(ref)
            changed += 1
            col, r = colnum(ref), rowof(ref)
            if col in srcc and (Ca is None or r == Ca['title_row'] or vb in (None, '')): continue
            if Ca and col in (Ca['vl0'], Ca['nc0']):
                if vb in (None, '') and va not in (None, '') and not str(va).startswith('='): continue   # ô trống được điền số
                m = WRAP_RE.match(str(va or ''))
                if m and str(vb or '').startswith('=') and m.group(1) == m.group(3) == str(vb)[1:]:
                    cb = cell_of(B.rows(sh).get(r), col)
                    if cb is not None and price_state(B, sh, cb)[0] == 'zero_formula': wrapped += 1; continue
                if col == Ca['vl0'] and r in touched and is_web_ref(B, sh, Cb, B.rows(sh).get(r)):   # thay giá tham khảo web
                    try:
                        if float(va) > 0 and not str(va).startswith('='): webrep += 1; continue
                    except (TypeError, ValueError): pass
            if Ca and r in touched and Ca.get('Ghi chú gán mã') and col == Ca['Ghi chú gán mã'] \
                    and WEB_MARK in str(vb or '') and str(va or '') == str(vb).replace(WEB_MARK, WEB_OLD, 1) \
                    and is_web_ref(B, sh, Cb, B.rows(sh).get(r)): continue
            if Ca and r in touched and Ca.get(ORIGIN_COL) and col == Ca[ORIGIN_COL] and 'NCC' in str(va or '') \
                    and (norm(vb) in ORIGIN_KNOWN or str(va).startswith(str(vb or ''))): continue
            bad.append((sh, ref, vb, va))
    print('Tổng ô thay đổi: %d (trong đó %d ô công thức được bọc IF, %d ô thay giá tham khảo web) | vi phạm: %d' % (changed, wrapped, webrep, len(bad)))
    for b in bad[:50]: print('  VI PHẠM %s!%s: %r -> %r' % b)
    print('PASS: chỉ ghi vào ô giá trống / công thức đang cho kết quả 0 / ô giá tham khảo web, cột nguồn NCC, Nguồn giá và ghi chú của dòng đã ghi' if not bad else 'FAIL')
    sys.exit(1 if bad else 0)


# ------------------------------------------------------------------ ledger
def cmd_ledger_init(out):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'SoCai'
    ws.append(LEDGER_COLS)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='FFD9E1F2')
    ws.freeze_panes = 'A2'
    widths = {'ID': 9, 'Hệ': 10, 'Nhóm VT': 10, 'NCC': 18, 'Ngày BG': 11, 'Hiệu lực đến': 11, 'Loại file': 14, 'STT YCBG': 9,
              'Đầu mục nguyên văn': 40, 'Tên hàng': 40, 'Hãng / Model / Quy cách': 30, 'ĐVT': 7, 'Đơn giá': 14,
              'VAT (Chưa/Đã/Không rõ)': 12, '% VAT': 7, 'Giá chưa VAT': 14, 'Loại giá (VL/NC/VL+NC)': 10,
              'Chiết khấu / ghi chú giá': 28, 'Điều kiện giao hàng': 24, 'File nguồn': 50, 'Vị trí trong file': 18,
              'Độ tin cậy trích': 10, 'Trạng thái': 14, 'Ghi chú': 30}
    for i, h in enumerate(LEDGER_COLS, 1): ws.column_dimensions[CL(i)].width = widths.get(h, 14)
    ws2 = wb.create_sheet('DaDoc'); ws2.append(['File', 'mtime', 'Số dòng trích', 'Ghi chú'])
    wb.save(out); print('Đã tạo sổ cái:', out)


LEDGER_ALIAS = {'Tên hàng': ('Tên hàng', 'Đầu mục nguyên văn'), 'VAT (Chưa/Đã/Không rõ)': ('VAT (Chưa/Đã/Không rõ)', 'VAT'),
                'Độ tin cậy trích': ('Độ tin cậy trích', 'Độ tin cậy')}


def cmd_ledger_check(f):
    import openpyxl
    ws = openpyxl.load_workbook(f, data_only=True)['SoCai']
    H = [c.value for c in ws[1]]; ix = {h: i for i, h in enumerate(H) if h}
    for k, alts in LEDGER_ALIAS.items():          # chấp nhận tên cột khác của sổ cái do ncc_rasoat tạo
        for alt in alts:
            if k not in ix and alt in ix: ix[k] = ix[alt]
    has_pos = 'Vị trí trong file' in ix or ('Sheet' in ix and 'Dòng' in ix)
    need = ['NCC', 'Ngày BG', 'Tên hàng', 'ĐVT', 'Đơn giá', 'VAT (Chưa/Đã/Không rõ)', 'File nguồn']
    opt = ['Loại file', 'STT YCBG', 'Đầu mục nguyên văn', 'Hiệu lực đến']
    absent = [k for k in need if k not in ix] + ([] if has_pos else ['Vị trí trong file (hoặc Sheet + Dòng)'])
    if absent: print('Sổ cái THIẾU cột bắt buộc: %s -> không kiểm tra được đầy đủ' % absent)
    absent_opt = [k for k in opt if k not in ix]
    if absent_opt: print('Sổ cái chưa có cột mở rộng (bỏ qua kiểm tra các cột này): %s' % absent_opt)
    get = lambda r, k: r[ix[k]] if k in ix else None
    today = datetime.date.today(); issues = collections.defaultdict(list); seen = collections.defaultdict(list); n = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r): continue
        n += 1; rid = get(r, 'ID') or '?'
        for k in need:
            if k in ix and get(r, k) in (None, ''): issues['Thiếu ' + k].append(rid)
        if has_pos and 'Vị trí trong file' not in ix and (get(r, 'Sheet') in (None, '') or get(r, 'Dòng') in (None, '')): issues['Thiếu vị trí (Sheet/Dòng)'].append(rid)
        if 'Vị trí trong file' in ix and get(r, 'Vị trí trong file') in (None, ''): issues['Thiếu Vị trí trong file'].append(rid)
        lf = get(r, 'Loại file')
        if 'Loại file' in ix:
            if lf in (None, ''): issues['Thiếu Loại file'].append(rid)
            elif lf not in LEDGER_FILE_TYPES: issues['Loại file không hợp lệ (%s)' % ' / '.join(LEDGER_FILE_TYPES)].append(rid)
            elif lf == 'YCBG điền giá' and 'Đầu mục nguyên văn' in ix and get(r, 'Đầu mục nguyên văn') in (None, ''):
                issues['YCBG điền giá thiếu Đầu mục nguyên văn'].append(rid)
        if str(get(r, 'VAT (Chưa/Đã/Không rõ)') or '').lower().startswith('không'): issues['VAT không rõ - cần xác nhận'].append(rid)
        if 'Hiệu lực đến' in ix:
            hl = get(r, 'Hiệu lực đến')
            if isinstance(hl, datetime.datetime): hl = hl.date()
            if isinstance(hl, datetime.date) and hl < today: issues['Hết hiệu lực'].append(rid)
            if get(r, 'Hiệu lực đến') in (None, ''): issues['Không ghi hiệu lực'].append(rid)
        if get(r, 'Độ tin cậy trích') in ('Thấp', 'Cần kiểm tra') or str(get(r, 'File nguồn') or '').lower().endswith(('.jpg', '.jpeg', '.png', '.msg')):
            issues['Trích từ ảnh/email hoặc độ tin cậy thấp - người kiểm tra'].append(rid)
        seen[(get(r, 'NCC'), str(get(r, 'Tên hàng')).strip().lower(), get(r, 'ĐVT'))].append((rid, get(r, 'Ngày BG'), get(r, 'Giá chưa VAT') or get(r, 'Đơn giá')))
    for k, v in seen.items():
        if len(v) > 1: issues['Cùng NCC+hàng có nhiều báo giá (chọn bản mới nhất, cảnh báo nếu lệch giá)'].append(str(v))
    print('Sổ cái: %d dòng' % n)
    for k, v in issues.items(): print(' - %s: %d %s' % (k, len(v), v[:8]))


if __name__ == '__main__':
    a = sys.argv[1:]
    if not a: print(__doc__); sys.exit()
    cmd, ov = a[0], None
    if '--map' in a: ov = load_map(a); i = a.index('--map'); a = a[:i] + a[i + 2:]
    if cmd == 'sheets': cmd_sheets(a[1], ov)
    elif cmd == 'todo':
        o = a[a.index('--out') + 1] if '--out' in a else None
        rest = [x for x in a[2:] if x != '--out' and x != o]
        cmd_todo(a[1], rest, o, ov)
    elif cmd == 'apply': cmd_apply(a[1], a[2], a[3], ov, '--dry' in a, '--ghi-so-0' in a)
    elif cmd == 'audit': cmd_audit(a[1], a[2], ov)
    elif cmd == 'ledger-init': cmd_ledger_init(a[1])
    elif cmd == 'ledger-check': cmd_ledger_check(a[1])
    else: print(__doc__)
