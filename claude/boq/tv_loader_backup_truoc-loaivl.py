#!/usr/bin/env python3
"""tv_loader.py v2 - nạp ĐỦ và ĐÚNG Thư viện giá TDI (thư mục các file BG) vào file BOQ: sheet "2 Thu vien gia", bảng Excel TV.

  python3 tv_loader.py load  LIBDIR BOQ.xlsx OUT.xlsx --code TDI25538 [--mtimes mtimes.json] [--replace]
  python3 tv_loader.py check LIBDIR BOQ.xlsx --code TDI25538          # đối chiếu bảng TV trong BOQ với file nguồn, từng dòng

Quy tắc (v2, sau lỗi thiếu 685 dòng ở bản v1 / Power Query):
- Đọc MỌI sheet, KỂ CẢ SHEET ẨN (bản v1 bỏ sheet ẩn -> mất 591 dòng "Cua gio" của file OGio). Sheet ẩn ghi "(sheet ẩn)" ở cột Sheet nguồn.
- Dòng tiêu đề = dòng đầu (trong 30 dòng) có ô chứa "mã hiệu"; ghép tiêu đề 2 tầng.
- GIỮ dòng khi: có (Mã hiệu khác rỗng/0 HOẶC Tên công việc) VÀ có (ĐVT HOẶC ít nhất 1 ô giá là số);
  dòng THIẾU ĐVT chỉ giữ khi có ĐỦ cả Mã và Tên (dòng nhóm có tên nhóm ở cột Mã, giá là hệ số 1/1,165… thì bỏ).
  -> giữ cả dòng KHÔNG CÓ MÃ (v1 mất 73 dòng CTKhac: cắt đục, giá đỡ, thử áp, đầu cốt...) và dòng THIẾU ĐVT có giá (v1 mất 21 dòng WE-…-HDPE).
- BỎ: dòng nhóm/tiêu đề (không ĐVT, không mã, giá chỉ là cờ 1/1 hoặc trống), dòng mẫu trống (mã 0 / trống và tên trống).
- Giá giữ NGUYÊN giá trị số kể cả 0 (không đổi 0 thành trống); ô chữ ("-", "liên hệ") -> trống, có liệt kê.
- Mã hiệu số nguyên ghi dạng chữ không ".0"; mã 0 ghi "0".
- Cột Ghi chú loader: "Thiếu ĐVT trong file nguồn", "Không có mã hiệu - tra theo tên".
- Sau khi nạp: tự đối chiếu (số dòng, tổng VL, tổng NC từng sheet nguồn) bằng một lần đọc độc lập; lệch -> DỪNG (exit 2).
- Báo trùng khóa (Nhóm+Mã, Nhóm+Tên) để skill gán mã tránh SUMIFS cộng dồn.
- Kiểm tra tên vùng (defined name) trùng "TV" (không phân biệt hoa thường): tên rác ẩn ({…}/#REF!) -> xóa và báo; tên thật -> DỪNG.
  (Đã gặp: tên ẩn "tv" làm Excel đổi bảng thành TV_1, mọi công thức SUMIFS(TV[...]) lỗi.)
- --replace: thay nội dung bảng TV đã có (giữ tên bảng/sheet), dùng khi nạp lại thư viện.
- Ghi bằng XML, không mở-lưu file BOQ bằng openpyxl.
"""
import sys, os, re, glob, json, zipfile, datetime, argparse, collections
import openpyxl
from lxml import etree
from openpyxl.utils import get_column_letter as CL
from xml.sax.saxutils import escape

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PR = 'http://schemas.openxmlformats.org/package/2006/relationships'
q = lambda t: '{%s}%s' % (NS, t)
COLS = ['Column1', 'Mã hiệu', 'Tên công việc', 'ĐVT', 'Giá VL gốc', 'Giá NC gốc', 'Giá VL chưa CK', 'Giá NC ước tính',
        'Description', 'Ghi chú TV', 'Nhãn Hiệu Xuất Xứ', 'Sheet nguồn', 'File nguồn', 'Ngày cập nhật']
NUM = ['Giá VL gốc', 'Giá NC gốc', 'Giá VL chưa CK', 'Giá NC ước tính']
SHEET = '2 Thu vien gia'


def norm(s): return re.sub(r'\s+', ' ', str(s or '')).strip().lower()


def std_name(top, sub):
    t, s = norm(top), norm(sub)
    if s == 'vật liệu' and ('số liệu gốc' in t or t == ''): return 'Giá VL gốc'
    if s == 'nhân công' and ('số liệu gốc' in t or t == ''): return 'Giá NC gốc'
    if 'nhân công' in t and 'ước' in t: return 'Giá NC ước tính'
    if ('báo giá' in t or 'báo gía' in t or 'b.giá' in t or 'vật tư theo báo giá' in t or 'chưa chiết khấu' in t) and 'nhân công' not in t:
        return 'Giá VL chưa CK'
    if 'mã hiệu' in t and 'mã hiệu' not in s: return 'Mã hiệu'
    if t == 'tên công việc': return 'Tên công việc'
    if t.startswith('unit') or 'đ. vị' in t or t == 'đvt': return 'ĐVT'
    if t == 'description': return 'Description'
    if t == 'ghi chú': return 'Ghi chú TV'
    if 'nhãn hiệu' in t and 'xuất xứ' in t: return 'Nhãn Hiệu Xuất Xứ'
    return None


def isnum(v): return isinstance(v, (int, float)) and not isinstance(v, bool)


def blank(v): return v is None or str(v).strip() == ''


def code_str(v):
    if blank(v): return None
    if isnum(v): return str(int(v)) if float(v) == int(v) else str(v)
    return str(v).strip()


def group_of(fname):
    b = os.path.splitext(os.path.basename(fname))[0]
    m = re.search(r'\.BG(.*)$', b, re.I)
    return (m.group(1) if m else b).strip(' -_')


def lib_files(libdir, code):
    files = sorted(f for f in glob.glob(os.path.join(libdir, '*.xlsx'))
                   if not os.path.basename(f).startswith('~$') and os.path.basename(f).upper().startswith(code.upper()))
    if not files:
        raise SystemExit('DỪNG: không có file thư viện nào bắt đầu bằng mã "%s" trong %s. Hỏi người dùng.' % (code, libdir))
    return files


def parse_sheet(ws):
    """Trả về (colmap, start, data) hoặc None nếu sheet không có bảng giá."""
    data = list(ws.iter_rows(values_only=True))
    h = next((i for i, r in enumerate(data[:30]) if any('mã hiệu' in norm(v) for v in r)), None)
    if h is None: return None
    top, sub = list(data[h]), list(data[h + 1]) if h + 1 < len(data) else []
    last, colmap = '', {}
    for j in range(len(top)):
        if top[j] not in (None, ''): last = top[j]
        name = std_name(last if top[j] in (None, '') else top[j], sub[j] if j < len(sub) else None)
        if name and name not in colmap.values(): colmap[j] = name
    start = h + 2 if any(norm(v) in ('vật liệu', 'nhân công') for v in sub) else h + 1
    return colmap, start, data


def keep_row(rec):
    ma, ten, dvt = code_str(rec.get('Mã hiệu')), rec.get('Tên công việc'), rec.get('ĐVT')
    has_id = (ma not in (None, '0')) or not blank(ten)
    prices = [rec.get(k) for k in NUM]
    has_val = not blank(dvt) or any(isnum(p) and p != 0 for p in prices[:2]) or any(isnum(p) and p != 0 for p in prices[2:])
    if not has_id: return False
    if blank(dvt) and (ma in (None, '0') or blank(ten)):
        return False                        # dòng nhóm/tiêu đề (tên nhóm nằm ở cột Mã, giá là hệ số 1/1,165...), không phải đầu mục
    return has_val


def load(libdir, code, mtimes):
    rows, stat, warn = [], [], []
    for f in lib_files(libdir, code):
        fn = os.path.basename(f); g = group_of(fn)
        ms = mtimes.get(fn) or os.path.getmtime(f) * 1000
        day = datetime.datetime.fromtimestamp(ms / 1000).strftime('%d/%m/%Y')
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        for ws in wb.worksheets:
            p = parse_sheet(ws)
            if p is None:
                warn.append('%s | %s: không có dòng tiêu đề "Mã hiệu" - bỏ qua' % (fn, ws.title)); continue
            colmap, start, data = p
            hidden = ws.sheet_state != 'visible'
            sname = ws.title + (' (sheet ẩn)' if hidden else '')
            n0 = len(rows); skipped = collections.Counter()
            for i, r in enumerate(data[start:], start + 1):
                rec = {colmap[j]: r[j] for j in colmap if j < len(r)}
                if not keep_row(rec):
                    if any(not blank(v) for v in r): skipped['dòng nhóm/mẫu trống'] += 1
                    continue
                out = {'Column1': g, 'Sheet nguồn': sname, 'File nguồn': fn, 'Ngày cập nhật': day, '_row': i}
                notes = []
                for k, v in rec.items():
                    if k in NUM:
                        if isnum(v): out[k] = float(v)
                        elif not blank(v): notes.append('%s="%s" không phải số' % (k, str(v).strip()[:15]))
                    elif k == 'Mã hiệu': out[k] = code_str(v)
                    elif not blank(v): out[k] = str(v).strip()
                if blank(rec.get('ĐVT')): notes.append('Thiếu ĐVT trong file nguồn')
                if out.get('Mã hiệu') in (None, '0'): notes.append('Không có mã hiệu - tra theo tên')
                if notes: out['Ghi chú TV'] = '; '.join(([out['Ghi chú TV']] if out.get('Ghi chú TV') else []) + notes)
                rows.append(out)
            stat.append(dict(file=fn, sheet=sname, group=g, rows=len(rows) - n0, day=day, skipped=dict(skipped)))
    return rows, stat, warn


def recount(libdir, code):
    """Đếm lại độc lập (không dùng keep_row): số dòng + tổng giá theo (file, sheet)."""
    res = {}
    for f in lib_files(libdir, code):
        wb = openpyxl.load_workbook(f, read_only=False, data_only=True)
        for ws in wb.worksheets:
            hr = None
            for row in ws.iter_rows(min_row=1, max_row=30):
                if any(isinstance(c.value, str) and 'mã hiệu' in c.value.lower() for c in row): hr = row[0].row; break
            if hr is None: continue
            hdr = {c.column: norm(c.value) for c in ws[hr]}
            sub = {c.column: norm(c.value) for c in ws[hr + 1]}
            cm = next(c for c, v in hdr.items() if 'mã hiệu' in v)
            cn = next((c for c, v in hdr.items() if v == 'tên công việc'), None)
            cu = next((c for c, v in hdr.items() if v.startswith('unit') or 'đ. vị' in v or v == 'đvt'), None)
            cvl = next((c for c, v in sub.items() if v == 'vật liệu'), None); cnc = next((c for c, v in sub.items() if v == 'nhân công'), None)
            first = hr + 2 if (cvl or cnc) else hr + 1
            n = 0; svl = snc = 0.0
            for r in range(first, ws.max_row + 1):
                ma = ws.cell(r, cm).value; ten = ws.cell(r, cn).value if cn else None; u = ws.cell(r, cu).value if cu else None
                vl = ws.cell(r, cvl).value if cvl else None; nc = ws.cell(r, cnc).value if cnc else None
                mas = code_str(ma)
                ident = (mas not in (None, '0')) or not blank(ten)
                if not ident: continue
                if blank(u) and (mas in (None, '0') or blank(ten)): continue
                if blank(u) and not any(isnum(x) and x != 0 for x in (vl, nc)) and not any(isnum(ws.cell(r, c).value) and ws.cell(r, c).value != 0 for c in hdr if c not in (cm, cn, cu)): continue
                n += 1; svl += vl if isnum(vl) else 0; snc += nc if isnum(nc) else 0
            res[(os.path.basename(f), ws.title + ('' if ws.sheet_state == 'visible' else ' (sheet ẩn)'))] = (n, round(svl, 2), round(snc, 2))
    return res


def verify(rows, libdir, code):
    got = collections.defaultdict(lambda: [0, 0.0, 0.0])
    for r in rows:
        k = (r['File nguồn'], r['Sheet nguồn']); got[k][0] += 1
        got[k][1] += r.get('Giá VL gốc') or 0; got[k][2] += r.get('Giá NC gốc') or 0
    exp = recount(libdir, code); bad = []
    print('\nĐỐI CHIẾU VỚI FILE NGUỒN (đếm độc lập):')
    print('%-36s %-26s %7s %7s %16s %16s' % ('File', 'Sheet', 'Nguồn', 'Đã nạp', 'ΣVL nguồn', 'ΣVL nạp'))
    for k in sorted(set(exp) | set(got)):
        e = exp.get(k, (0, 0, 0)); g = got.get(k, [0, 0, 0])
        ok = e[0] == g[0] and abs(e[1] - g[1]) < 0.5 and abs(e[2] - g[2]) < 0.5
        print('%-36s %-26s %7d %7d %16.0f %16.0f %s' % (k[0][:36], k[1][:26], e[0], g[0], e[1], g[1], '' if ok else '<-- LỆCH'))
        if not ok: bad.append(k)
    return bad


def dup_report(rows):
    c1 = collections.Counter((r['Column1'], r.get('Mã hiệu')) for r in rows if r.get('Mã hiệu') not in (None, '0'))
    c2 = collections.Counter((r['Column1'], r.get('Tên công việc')) for r in rows if r.get('Tên công việc'))
    d1 = [k for k, v in c1.items() if v > 1]; d2 = [k for k, v in c2.items() if v > 1]
    if d1: print('CẢNH BÁO trùng Nhóm+Mã (%d): %s' % (len(d1), d1[:15]))
    if d2: print('CẢNH BÁO trùng Nhóm+Tên (%d) - không tra theo tên các mục này: %s' % (len(d2), [k[1][:40] for k in d2[:15]]))


# ------------------------------------------------------------ ghi XML
def cell(ref, v):
    if v is None: return ''
    if isinstance(v, float): return '<c r="%s"><v>%r</v></c>' % (ref, v)
    return '<c r="%s" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (ref, escape(str(v)))


def sheet_xml(rows):
    ncol, nrow = len(COLS), len(rows) + 1
    ref = 'A1:%s%d' % (CL(ncol), nrow)
    xs = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
          '<worksheet xmlns="%s" xmlns:r="%s"><dimension ref="%s"/>' % (NS, RNS, ref),
          '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>',
          '<sheetFormatPr defaultRowHeight="15"/><cols><col min="1" max="1" width="12" customWidth="1"/><col min="2" max="2" width="22" customWidth="1"/>'
          '<col min="3" max="3" width="45" customWidth="1"/><col min="4" max="8" width="13" customWidth="1"/><col min="9" max="%d" width="18" customWidth="1"/></cols><sheetData>' % ncol]
    xs.append('<row r="1">' + ''.join(cell(CL(j + 1) + '1', c) for j, c in enumerate(COLS)) + '</row>')
    for i, r in enumerate(rows, 2):
        xs.append('<row r="%d">' % i + ''.join(cell(CL(j + 1) + str(i), r.get(c)) for j, c in enumerate(COLS)) + '</row>')
    xs.append('</sheetData><pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
              '<tableParts count="1"><tablePart r:id="rIdTV1"/></tableParts></worksheet>')
    return ''.join(xs).encode('utf8'), ref


def table_xml(tid, ref):
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><table xmlns="%s" id="%d" name="TV" displayName="TV" ref="%s">'
            '<autoFilter ref="%s"/><tableColumns count="%d">%s</tableColumns>'
            '<tableStyleInfo name="TableStyleLight9" showFirstColumn="0" showLastColumn="0" showRowStripes="1" showColumnStripes="0"/></table>'
            % (NS, tid, ref, ref, len(COLS), ''.join('<tableColumn id="%d" name="%s"/>' % (j + 1, escape(c)) for j, c in enumerate(COLS)))).encode('utf8')


def fix_names(wbx):
    """Tên vùng trùng 'TV' (không phân biệt hoa thường): rác -> xóa; thật -> dừng."""
    dns = wbx.find(q('definedNames'))
    if dns is None: return []
    out = []
    for dn in list(dns):
        if (dn.get('name') or '').lower() == 'tv':
            txt = dn.text or ''
            if dn.get('hidden') == '1' or txt.startswith('{') or '#REF' in txt:
                dns.remove(dn); out.append('%s=%s' % (dn.get('name'), txt[:40]))
            else:
                raise SystemExit('DỪNG: file có tên vùng "%s" = %s trùng tên bảng TV. Hỏi người dùng.' % (dn.get('name'), txt))
    return out


def write(src, out, rows, replace=False):
    z = zipfile.ZipFile(src); files = {i.filename: z.read(i.filename) for i in z.infolist()}; infos = z.infolist()
    wb = etree.fromstring(files['xl/workbook.xml'])
    removed = fix_names(wb)
    if removed: print('Đã xóa tên vùng rác trùng TV:', removed)
    rels = etree.fromstring(files['xl/_rels/workbook.xml.rels'])
    rid = {r.get('Id'): r.get('Target') for r in rels}
    tv_tables = [n for n in files if re.match(r'xl/tables/table\d+\.xml$', n) and re.search(rb'name="TV"|displayName="TV(_\d+)?"', files[n])]
    sx, ref = sheet_xml(rows)
    new = {}
    if tv_tables:
        if not replace: raise SystemExit('DỪNG: file đã có bảng TV (%s). Hỏi người dùng; nạp lại thì chạy với --replace.' % tv_tables[0])
        tpath = tv_tables[0]; tid = int(etree.fromstring(files[tpath]).get('id'))
        # tìm sheet chứa bảng
        wsp = None
        for s in wb.find(q('sheets')):
            t = rid[s.get('{%s}id' % RNS)].lstrip('/'); t = t if t.startswith('xl/') else 'xl/' + t
            rp = t.replace('worksheets/', 'worksheets/_rels/') + '.rels'
            if rp in files and os.path.basename(tpath).encode() in files[rp]: wsp, wrp = t, rp
        if not wsp: raise SystemExit('DỪNG: không tìm được sheet chứa bảng TV')
        rr = etree.fromstring(files[wrp]); trid = next(r.get('Id') for r in rr if os.path.basename(tpath) in r.get('Target'))
        new[wsp] = sx.replace(b'r:id="rIdTV1"', ('r:id="%s"' % trid).encode())
        new[tpath] = table_xml(tid, ref)
        print('Thay nội dung bảng TV có sẵn (%s, %s)' % (wsp, tpath))
    else:
        if any(s.get('name') == SHEET for s in wb.find(q('sheets'))): raise SystemExit('DỪNG: đã có sheet "%s" nhưng không có bảng TV.' % SHEET)
        n_ws = 1 + max([int(m.group(1)) for k in files for m in [re.match(r'xl/worksheets/sheet(\d+)\.xml$', k)] if m] or [0])
        n_tb = 1 + max([int(re.findall(r'table(\d+)\.xml$', k)[0]) for k in files if re.match(r'xl/tables/table\d+\.xml$', k)] or [0])
        ids = [int(etree.fromstring(files[k]).get('id')) for k in files if re.match(r'xl/tables/table\d+\.xml$', k)]
        wsp = 'xl/worksheets/sheet%d.xml' % n_ws
        new[wsp] = sx
        new['xl/worksheets/_rels/sheet%d.xml.rels' % n_ws] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="%s">'
            '<Relationship Id="rIdTV1" Type="%s/table" Target="../tables/table%d.xml"/></Relationships>' % (PR, RNS, n_tb)).encode('utf8')
        new['xl/tables/table%d.xml' % n_tb] = table_xml(max(ids or [0]) + 1, ref)
        r_id = 'rIdTVs%d' % n_ws
        etree.SubElement(rels, '{%s}Relationship' % PR, Id=r_id, Type=RNS + '/worksheet', Target='worksheets/sheet%d.xml' % n_ws)
        new['xl/_rels/workbook.xml.rels'] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
        shs = wb.find(q('sheets')); sid = 1 + max(int(s.get('sheetId')) for s in shs)
        s = etree.SubElement(shs, q('sheet')); s.set('name', SHEET); s.set('sheetId', str(sid)); s.set('{%s}id' % RNS, r_id)
        ct = files['[Content_Types].xml'].decode('utf8')
        add = ('<Override PartName="/%s" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' % wsp +
               '<Override PartName="/xl/tables/table%d.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml"/>' % n_tb)
        new['[Content_Types].xml'] = ct.replace('</Types>', add + '</Types>').encode('utf8')
    cp = wb.find(q('calcPr'))
    if cp is None: cp = etree.SubElement(wb, q('calcPr'))
    cp.set('fullCalcOnLoad', '1')
    new['xl/workbook.xml'] = etree.tostring(wb, xml_declaration=True, encoding='UTF-8', standalone=True)
    drop = {'xl/calcChain.xml'} if replace else set()
    if drop and 'xl/calcChain.xml' in files:
        rels2 = etree.fromstring(new.get('xl/_rels/workbook.xml.rels', files['xl/_rels/workbook.xml.rels']))
        for r in list(rels2):
            if 'calcChain' in r.get('Target'): rels2.remove(r)
        new['xl/_rels/workbook.xml.rels'] = etree.tostring(rels2, xml_declaration=True, encoding='UTF-8', standalone=True)
        ct = new.get('[Content_Types].xml', files['[Content_Types].xml']).decode('utf8')
        new['[Content_Types].xml'] = re.sub(r'<Override[^>]*calcChain[^>]*/>', '', ct).encode('utf8')
    zo = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    for it in infos:
        if it.filename in drop: continue
        zo.writestr(it, new.pop(it.filename, files[it.filename]))
    for k, v in new.items(): zo.writestr(k, v)
    zo.close()
    return ref


def read_tv(boq):
    wb = openpyxl.load_workbook(boq, read_only=True)
    ws = wb[SHEET]; ws.reset_dimensions()
    it = ws.iter_rows(values_only=True); H = [str(x) for x in next(it)]
    return [dict(zip(H, r)) for r in it if any(v is not None for v in r)]


def cmd_check(libdir, boq, code):
    """Đối chiếu bảng TV trong BOQ với thư viện nguồn: thiếu / thừa / lệch giá từng dòng."""
    src, _, _ = load(libdir, code, {})
    tv = read_tv(boq)
    key = lambda r: (r.get('Column1'), str(r.get('Mã hiệu') or ''), str(r.get('Tên công việc') or ''), str(r.get('ĐVT') or ''))
    S = collections.Counter(key(r) for r in src); T = collections.Counter(key(r) for r in tv)
    miss = S - T; extra = T - S
    price = collections.defaultdict(list)
    for r in src: price[key(r)].append((r.get('Giá VL gốc') or 0, r.get('Giá NC gốc') or 0))
    wrong = []
    for r in tv:
        k = key(r)
        if k in price and (float(r.get('Giá VL gốc') or 0), float(r.get('Giá NC gốc') or 0)) not in [(float(a), float(b)) for a, b in price[k]]:
            wrong.append((k, r.get('Giá VL gốc'), r.get('Giá NC gốc'), price[k][0]))
    print('Nguồn %d dòng | bảng TV %d dòng | THIẾU %d | THỪA %d | LỆCH GIÁ %d' % (len(src), len(tv), sum(miss.values()), sum(extra.values()), len(wrong)))
    for k in list(miss)[:30]: print('  THIẾU', k)
    for k in list(extra)[:30]: print('  THỪA ', k)
    for w in wrong[:30]: print('  LỆCH ', w)
    by = collections.Counter(k[0] for k in miss)
    if by: print('  Thiếu theo nhóm:', dict(by))
    sys.exit(0 if not (miss or extra or wrong) else 2)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('cmd', choices=['load', 'check'])
    ap.add_argument('libdir'); ap.add_argument('boq'); ap.add_argument('out', nargs='?')
    ap.add_argument('--code', required=True); ap.add_argument('--mtimes'); ap.add_argument('--replace', action='store_true')
    a = ap.parse_args()
    if a.cmd == 'check': cmd_check(a.libdir, a.boq, a.code)
    mt = json.load(open(a.mtimes)) if a.mtimes else {}
    rows, stat, warn = load(a.libdir, a.code, mt)
    for w in warn: print('LƯU Ý:', w)
    for s in stat:
        print('%-36s %-26s nhóm %-12s %5d dòng  cập nhật %s  bỏ: %s' % (s['file'][:36], s['sheet'][:26], s['group'], s['rows'], s['day'], s['skipped'] or '-'))
    bad = verify(rows, a.libdir, a.code)
    if bad: raise SystemExit('DỪNG: số dòng/tổng giá đã nạp lệch file nguồn ở %s - không ghi file.' % bad)
    dup_report(rows)
    for r in rows: r.pop('_row', None)
    ref = write(a.boq, a.out, rows, a.replace)
    groups = sorted({r['Column1'] for r in rows})
    print('Đã nạp %d nhóm, %d dòng (%d dòng từ sheet ẩn, %d dòng không mã, %d dòng thiếu ĐVT) vào sheet "%s", bảng TV %s' % (
        len(groups), len(rows), sum('(sheet ẩn)' in r['Sheet nguồn'] for r in rows),
        sum(r.get('Mã hiệu') in (None, '0') for r in rows), sum('Thiếu ĐVT' in str(r.get('Ghi chú TV')) for r in rows), SHEET, ref))
