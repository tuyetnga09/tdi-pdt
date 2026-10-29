#!/usr/bin/env python3
"""boq_dvt_tools.py - Bước 6f (Quy tắc 35, thêm 21/10/2026): cột "Đơn Vị" của phần TDI hiển thị ĐVT của thư viện giá (bảng TV).

Đặt CÙNG THƯ MỤC với boq_xml_tools.py (dùng lại Pkg, detect_cols, setc... của file đó, không sửa file đó).

Lệnh:
  dvt   FILE OUT [SHEET...]   ô Đơn vị TDI của dòng `Dòng` / `Con` = công thức tra TV[ĐVT] theo Nhóm VT + Mã tra cứu + Tên tra cứu
                              của chính dòng (cùng khóa với công thức giá). Chưa có mã, hoặc TV không có ĐVT cho khóa đó
                              -> giữ ĐVT cũ (tham chiếu ĐVT CĐT, hoặc chữ đang có ở dòng Con). Chọn mã từ droplist là ĐVT đổi theo.
                              Ô KL (điều kiện m/m2/m3 nhân HSKL) đang đọc ô Đơn vị TDI được trỏ sang ĐVT CĐT (giá trị cũ),
                              nên thành tiền KHÔNG đổi. Dòng `Cụm`, tiêu đề, tổng: không động vào.
  kiem  FILE [--json]         tính lại bản sao: QT35 (mọi dòng Dòng / Con đã có công thức ĐVT, 0 ô lỗi ở cột Đơn vị TDI)
                              + danh sách dòng ĐVT thư viện khác ĐVT CĐT (để người duyệt quy đổi hoặc chiết tính). Thoát 1 nếu Lệch.

Chạy sau `cho` (Bước 6e), trước `fillcdt`; chạy lại được (kết quả giống nhau), chạy lại sau mỗi lần apply / write / insert / cho / migrate2.
"""
import os, sys, re, json, collections, unicodedata
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import boq_xml_tools as T
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as CL

KINDS = ('Dòng', 'Con')
LINEAR = {'m', 'm2', 'm3', 'md', 'mét', 'met', 'm²', 'm³'}          # như LINEAR_UNITS của skill chuyển template (Bước 1)
REF1 = re.compile(r'^\+?(\$?[A-Z]{1,3}\$?\d+)$')                    # =+G12 / =G12 / =$G12 (ĐVT CĐT, Bước 1)
OURS = re.compile(r'^IF\(OR\(\$[A-Z]+\d+="",AND\(\$[A-Z]+\d+="",\$[A-Z]+\d+=""\)\),(.+?),IFERROR\(INDEX\(TV\[')
KLOLD = re.compile(r'^IF\(OR\((\$?[A-Z]+\$?\d+)="m",\1="m2",\1="m3"\),')


def nu(s): return unicodedata.normalize('NFC', str(s if s is not None else '')).strip().lower().replace('.', '')


def tv_unit_col(pkg):
    """Tên cột ĐVT thật trong bảng TV (so không phân biệt hoa thường / dạng Unicode)."""
    sh, tp, tb = T._tv_table(pkg)
    for tc in tb.find(T.q('tableColumns')):
        if nu(tc.get('name')) == 'đvt': return tc.get('name')
    raise SystemExit('DỪNG: bảng TV không có cột "ĐVT" - nạp lại thư viện (tv_loader load --replace)')


def tv_rows(pkg, ucol):
    """[(Loại VL, Mã hiệu, Tên công việc, ĐVT)] theo thứ tự bảng TV, chuẩn hóa giống công thức (mã/tên so như chữ)."""
    from openpyxl.utils.cell import range_boundaries
    sh, tp, tb = T._tv_table(pkg)
    c1, r1, c2, r2 = range_boundaries(tb.get('ref'))
    names = [tc.get('name') for tc in tb.find(T.q('tableColumns'))]
    want = {c1 + names.index(n): k for n, k in ((names[0], 'g'), ('Mã hiệu', 'm'), ('Tên công việc', 'n'), (ucol, 'u'))}
    out = []
    for row in pkg.root(sh).find(T.q('sheetData')):
        rn = int(row.get('r'))
        if rn <= r1 or rn > r2: continue
        v = {}
        for c in row:
            k = want.get(T.colnum(c.get('r')))
            if k: v[k] = T._cellv(pkg, c)
        s = lambda k: '' if v.get(k) is None else str(v[k])
        m = s('m').strip()
        if re.fullmatch(r'-?\d+\.0', m): m = m[:-2]
        out.append((s('g').strip(), m, s('n'), s('u').strip()))
    return out


def tv_unit(TVR, g, code_text):
    """ĐVT thư viện mà công thức sẽ ra (dòng TV đầu tiên khớp khóa và có ĐVT), None nếu giữ ĐVT cũ."""
    mc, nm = T.split_code(code_text)
    if not str(code_text or '').strip() or (not mc and not nm): return None
    gf = str(g or '').strip() not in ('', 'Chưa xác định')
    for tg, tm, tn, tu in TVR:
        if gf and tg.lower() != str(g).strip().lower(): continue
        if mc and tm.lower() != mc.lower(): continue
        if nm and tn.lower() != nm.lower(): continue
        if tu: return tu
    return None


def unit_formula(r, c, fb, ucol):
    L = lambda k: '$%s%d' % (CL(c[k]), r)
    G, M, N, C = L('Nhóm VT'), L('Mã tra cứu'), L('Tên tra cứu'), L('code')
    unesc = lambda x: 'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(%s,"~*","*"),"~?","?"),"~~","~")' % x   # cột phụ đã escape cho SUMIFS
    U = 'TV[%s]' % ucol
    cond = ('((%s="")+(%s="Chưa xác định")+(TV[%s]=%s)>0)*((%s="")+(TV[Mã hiệu]&""=%s)>0)*((%s="")+(TV[Tên công việc]&""=%s)>0)*(%s&""<>"")'
            % (G, G, T.GCOL, G, M, unesc(M), N, unesc(N), U))
    return 'IF(OR(%s="",AND(%s="",%s="")),%s,IFERROR(INDEX(%s,SUMPRODUCT(MATCH(1,%s,0))),%s))' % (C, M, N, fb, U, cond, fb)


def unshare_cols(sd, cols):
    """Gỡ shared formula của mọi nhóm có ô nằm trong các cột sẽ ghi (ghi đè ô gốc của nhóm làm hỏng các ô còn lại)."""
    groups = collections.defaultdict(list); master = {}
    for row in sd:
        for c in row:
            fx = c.find(T.q('f'))
            if fx is not None and fx.get('t') == 'shared':
                groups[fx.get('si')].append(c)
                if fx.text: master[fx.get('si')] = (fx.text, c.get('r'))
    n = 0
    for si, cells in groups.items():
        if not any(T.colnum(c.get('r')) in cols for c in cells) or si not in master: continue
        txt, org = master[si]
        for c in cells:
            fx = c.find(T.q('f'))
            if not fx.text: fx.text = Translator('=' + txt, origin=org).translate_formula(c.get('r'))[1:]
            for a in ('t', 'si', 'ref'):
                if a in fx.attrib: del fx.attrib[a]
            n += 1
    return n


def cmd_dvt(f, out, sheets=None):
    pkg = T.Pkg(f); ucol = tv_unit_col(pkg); TVR = tv_rows(pkg, ucol)
    report = {}
    for sh, c in T._work_sheets(pkg, sheets):
        sd = pkg.root(sh).find(T.q('sheetData')); un = c['unit_tdi']
        stat = collections.Counter(); skip, diff, kl_skip = [], [], []
        work = []
        for row in sd:
            r = int(row.get('r'))
            if r > c['hdr_row'] and T._is_work_row(pkg, row, c, KINDS): work.append(row)
        if not work: continue
        UL = CL(un)
        kl_cols = set()
        for row in work:           # cột chứa công thức KL đọc ô Đơn vị TDI
            for x in row:
                fx = x.find(T.q('f'))
                if fx is not None and fx.text and KLOLD.match(fx.text): kl_cols.add(T.colnum(x.get('r')))
        nun = unshare_cols(sd, {un} | kl_cols)
        for row in work:
            r = int(row.get('r')); kind = pkg.val(T._cell(row, c['Loại dòng']))
            cell = T._cell(row, un); fx = cell.find(T.q('f'))
            if fx is not None and fx.text:
                t = fx.text.strip(); m = REF1.match(t); m2 = OURS.match(t)
                if m and re.sub(r'\D', '', m.group(1)) == str(r): fb, how = m.group(1), 'ghi'
                elif m2: fb, how = m2.group(1), 'đã có (ghi lại)'
                else: stat['bỏ qua'] += 1; skip.append('%s%d (công thức khác: %s)' % (UL, r, t[:40])); continue
            else:
                v = pkg.val(cell)
                if v in (None, '') or not str(v).strip(): continue
                fb, how = '"%s"' % str(v).replace('"', '""'), 'ghi (giữ chữ cũ làm dự phòng)'
            T.setc(row, un, formula=unit_formula(r, c, fb, ucol)); stat[how] += 1
            # KL: điều kiện HSKL đang đọc ô Đơn vị TDI -> đọc ĐVT cũ (fb) để thành tiền không đổi
            for x in row:
                kf = x.find(T.q('f'))
                if kf is None or not kf.text: continue
                mk = KLOLD.match(kf.text)
                if mk and mk.group(1).replace('$', '') == '%s%d' % (UL, r):
                    kf.text = 'IF(OR(%s="m",%s="m2",%s="m3"),' % (fb, fb, fb) + kf.text[mk.end():]; stat['KL trỏ về ĐVT cũ'] += 1
                elif mk and mk.group(1) != fb and kind == 'Dòng':
                    kl_skip.append('%s (%s)' % (x.get('r'), kf.text[:40]))
            # báo cáo lệch ĐVT (tĩnh, theo bảng TV hiện có)
            code = T._cellv(pkg, T._cell(row, c['code'])) if T._cell(row, c['code']) is not None else ''
            g = T._cellv(pkg, T._cell(row, c['Nhóm VT'])) if T._cell(row, c['Nhóm VT']) is not None else ''
            tu = tv_unit(TVR, g, code)
            if fb.startswith('"'): old = fb[1:-1].replace('""', '"')
            else:
                rc = T._cell(row, T.colnum(fb.replace('$', '')))
                old = T._cellv(pkg, rc) if rc is not None else ''
            if tu is not None and nu(tu) != nu(old):
                diff.append({'dong': r, 'loai': kind, 'dvt_cdt_cu': old, 'dvt_thu_vien': tu, 'ma': str(code or '')[:60],
                             'doi_nhom_m_m2_m3': (nu(old) in LINEAR) != (nu(tu) in LINEAR)})
        pkg.touch(sh)
        report[sh] = {'thong_ke': dict(stat), 'go_shared': nun, 'bo_qua': skip, 'kl_khong_doi': kl_skip, 'lech_dvt': diff}
        print('%s: %s%s' % (sh, dict(stat), ' | gỡ shared %d ô' % nun if nun else ''))
        if skip: print('  BỎ QUA (ô Đơn vị TDI không phải dạng chuẩn, giữ nguyên):', '; '.join(skip[:15]))
        if diff:
            print('  ĐVT thư viện KHÁC ĐVT CĐT: %d dòng (độ tin cậy phải là Thấp, cần quy đổi `coef` hoặc chiết tính)' % len(diff))
            for d in diff[:20]:
                print('    dòng %d %-4s CĐT "%s" -> TV "%s"%s  %s' % (d['dong'], d['loai'], d['dvt_cdt_cu'], d['dvt_thu_vien'],
                      '  [khác nhóm m/m2/m3]' if d['doi_nhom_m_m2_m3'] else '', d['ma']))
    T._drop_calcchain(pkg); pkg.save(out, drop=('xl/calcChain.xml',))
    json.dump(report, open(out + '.dvt_report.json', 'w'), ensure_ascii=False, indent=1, default=str)
    print('Đã ghi %s; báo cáo %s.dvt_report.json' % (out, out))


def cmd_kiem(f, as_json=False):
    pkg = T.Pkg(f); wbv = T._recalc_wb(f); res = []
    ok = lambda sh, muc, dat, ct: res.append({'sheet': sh, 'muc': muc, 'dat': bool(dat), 'chi_tiet': ct})
    for sh, c in T._work_sheets(pkg, None):
        ws = wbv[sh]; miss, err, lech, n = [], [], [], 0
        for row in pkg.root(sh).find(T.q('sheetData')):
            r = int(row.get('r'))
            if r <= c['hdr_row'] or not T._is_work_row(pkg, row, c, KINDS): continue
            n += 1
            fx = T._cell(row, c['unit_tdi']).find(T.q('f'))
            m = OURS.match(fx.text) if fx is not None and fx.text else None
            if not m: miss.append(r); continue
            v = ws.cell(r, c['unit_tdi']).value
            if v is None or (isinstance(v, str) and v.startswith('#')): err.append(r); continue
            fb = m.group(1)
            old = fb[1:-1].replace('""', '"') if fb.startswith('"') else ws[fb.replace('$', '')].value
            if nu(v) != nu(old): lech.append('%d: %s -> %s' % (r, old, v))
        if not n: continue
        ok(sh, 'QT35 dòng Dòng/Con có công thức ĐVT thư viện', not miss, '%d/%d dòng chưa có %s' % (len(miss), n, miss[:10]))
        ok(sh, 'QT35 cột Đơn vị TDI không ra lỗi / trống', not err, '%d dòng %s' % (len(err), err[:10]))
        res.append({'sheet': sh, 'muc': 'QT35 dòng ĐVT thư viện khác ĐVT CĐT (liệt kê, cần duyệt)', 'dat': True,
                    'chi_tiet': '%d dòng %s' % (len(lech), lech[:15])})
    if as_json: print(json.dumps(res, ensure_ascii=False, indent=1))
    else:
        for x in res: print('%-5s %-24s %-58s %s' % ('ĐẠT' if x['dat'] else 'LỆCH', x['sheet'][:24], x['muc'], x['chi_tiet']))
    return 0 if all(x['dat'] for x in res) else 1


if __name__ == '__main__':
    a = sys.argv[1:]
    if not a: print(__doc__); sys.exit()
    if a[0] == 'dvt': cmd_dvt(a[1], a[2], a[3:] or None)
    elif a[0] == 'kiem': sys.exit(cmd_kiem(a[1], '--json' in a))
    else: print(__doc__)
