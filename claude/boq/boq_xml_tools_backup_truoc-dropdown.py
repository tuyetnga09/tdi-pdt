#!/usr/bin/env python3
"""boq_xml_tools.py - công cụ ghi trực tiếp XML vào file BOQ template TDI (không dùng openpyxl để ghi).

Lệnh:
  sheets   FILE                         : tên sheet -> file XML (luôn tra lại, Excel hay đổi số)
  cols     FILE SHEET                   : nhận diện cột TDI + cột phụ theo tiêu đề (in JSON)
  helpers  FILE OUT SHEET...            : tạo 5 cột phụ (Nhóm VT, Độ tin cậy, Ghi chú gán mã, Hệ số ĐV, Loại dòng) nếu chưa có
  insert   FILE OUT PLAN.json           : bóc tách - chèn dòng con dưới dòng cụm, tô màu phần TDI
  fillcdt  FILE OUT CFG.json            : đưa đơn giá TDI sang cột giá BOQ CĐT bằng công thức
  recalc   FILE [--sheets 'A|B']          : xóa cache <v>, tính lại bằng LibreOffice trên BẢN SAO, báo lỗi/tổng/có mã không giá
  cleanup  FILE OUT                     : xóa sheet tạm _TDI_JOB / _TDI_TMP_*

PLAN.json (insert): [[sheet, parent_row, expected_desc_prefix, children[], parent_note, parent_conf], ...]
  child = {g, key ('Mã hiệu'|'Tên công việc'|null), val, ma, name, unit, coef, conf, note}
  parent_row theo số dòng HIỆN TẠI của FILE. Chèn là bước cuối (số dòng đổi sau đó).
CFG.json (fillcdt): {"<sheet>": {"kl":["D","E","F"], "zones":["K","N","Q"], "h":"H","i":"I","j":"J",
                     "unit":"C", "start":8}}  - cột TDI (VL/NC sau hệ số, Loại dòng) tự nhận theo tiêu đề.
  CĐT chỉ 1 cột đơn giá: {"<sheet>": {"kl":["E"], "unit":"D", "start":6, "single":{"f":"F","g":"G"}}}
detect_cols: cột "Nội Dung Công Việc" lấy ô NGOÀI CÙNG BÊN PHẢI (CĐT có thể dùng cùng tiêu đề ở cột B).
Công thức giá (tv_sumifs, bản 14/10/2026): SUMIFS tham chiếu ô Nhóm VT / Mã hiệu của chính dòng, ví dụ
  SUMIFS(TV[Giá NC gốc],TV[Column1],$T265,TV[Mã hiệu],$N265). Chỉ viết cứng khi tra theo Tên công việc (mã 0),
  số hạng ghép khác nhóm/mã của dòng, hoặc giá trị có ký tự đại diện ~ * ?.
"""
import zipfile, re, json, copy, sys, os, subprocess, tempfile, shutil, collections
from lxml import etree
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
q = lambda t: '{%s}%s' % (NS, t)
HELPERS = ['Nhóm VT', 'Độ tin cậy', 'Ghi chú gán mã', 'Hệ số ĐV', 'Loại dòng']
COLOR_NAME, COLOR_CODE, FILL_CHILD = 'FF0070C0', 'FFFF0000', 'FFFFF2CC'
CONF_ORDER = ['Cao', 'TB', 'Thấp', 'Chưa khớp']


# ---------------------------------------------------------------- package I/O
class Pkg:
    def __init__(self, path):
        z = zipfile.ZipFile(path)
        self.infos = z.infolist()
        self.files = {i.filename: z.read(i.filename) for i in self.infos}
        self.new = {}
        self.wb = etree.fromstring(self.files['xl/workbook.xml'])
        rels = etree.fromstring(self.files['xl/_rels/workbook.xml.rels'])
        rid = {r.get('Id'): r.get('Target') for r in rels}
        self.sheets = []
        for s in self.wb.find(q('sheets')):
            t = rid[s.get('{%s}id' % RNS)]
            t = t.lstrip('/')
            t = t if t.startswith('xl/') else 'xl/' + t
            self.sheets.append((s.get('name'), t))
        self.path = dict(self.sheets)
        self.ss = []
        if 'xl/sharedStrings.xml' in self.files:
            sst = etree.fromstring(self.files['xl/sharedStrings.xml'])
            self.ss = [''.join(t.text or '' for t in si.iter(q('t'))) for si in sst]
        self._roots = {}

    def root(self, sheet):
        p = self.path[sheet]
        if p not in self._roots:
            self._roots[p] = etree.fromstring(self.new.get(p, self.files[p]))
        return self._roots[p]

    def touch(self, sheet):
        self.new[self.path[sheet]] = None  # serialise at save

    def val(self, c):
        t = c.get('t'); v = c.find(q('v'))
        if t == 's' and v is not None: return self.ss[int(v.text)]
        if t == 'inlineStr': return ''.join(x.text or '' for x in c.iter(q('t')))
        f = c.find(q('f'))
        if f is not None and f.text: return '=' + f.text
        return v.text if v is not None else None

    def grid(self, sheet, r1, r2):
        out = {}
        for row in self.root(sheet).find(q('sheetData')):
            r = int(row.get('r'))
            if r1 <= r <= r2:
                for c in row:
                    out[c.get('r')] = self.val(c)
        return out

    def save(self, out, drop=()):
        for p, rt in self._roots.items():
            if p in self.new:
                self.new[p] = etree.tostring(rt, xml_declaration=True, encoding='UTF-8', standalone=True)
        cp = self.wb.find(q('calcPr'))
        if cp is None: cp = etree.SubElement(self.wb, q('calcPr'))
        cp.set('fullCalcOnLoad', '1')
        self.new['xl/workbook.xml'] = etree.tostring(self.wb, xml_declaration=True, encoding='UTF-8', standalone=True)
        z = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
        for it in self.infos:
            if it.filename in drop: continue
            z.writestr(it, self.new.get(it.filename) or self.files[it.filename])
        z.close()


def colnum(ref): return CI(re.match(r'[A-Z]+', ref).group(0))


def setc(row, col, val=None, formula=None, num=None, style=None, clear=False):
    """Ghi 1 ô vào <row>, giữ thứ tự cột và style sẵn có."""
    ref = CL(col) + row.get('r'); cell = None; pos = len(row)
    for i, c in enumerate(row):
        cc = colnum(c.get('r'))
        if cc == col: cell = c; break
        if cc > col: pos = i; break
    if cell is None:
        if clear: return None
        cell = etree.Element(q('c')); cell.set('r', ref); row.insert(pos, cell)
    s = cell.get('s') if style is None else str(style)
    for ch in list(cell): cell.remove(ch)
    for a in list(cell.attrib):
        if a != 'r': del cell.attrib[a]
    if s: cell.set('s', s)
    if formula is not None:
        etree.SubElement(cell, q('f')).text = formula
    elif num is not None:
        etree.SubElement(cell, q('v')).text = repr(num)
    elif val not in (None, ''):
        cell.set('t', 'inlineStr')
        etree.SubElement(etree.SubElement(cell, q('is')), q('t')).text = str(val)
    return cell


def esc(s): return str(s).replace('~', '~~').replace('*', '~*').replace('?', '~?').replace('"', '""')


WILD = re.compile(r'[~*?]')


def tv_sumifs(lc, g, key, val, r=None, gcol=None, row_g=None, ccol=None, row_code=None):
    """Một số hạng SUMIFS tra bảng TV.
    Tham chiếu ô trên chính dòng đó thay vì viết cứng chuỗi (để người duyệt nhìn dòng là biết tra theo gì,
    sửa mã ở ô Mã hiệu thì giá tự đổi):
      - Nhóm  -> $<cột Nhóm VT><r> khi nhóm của số hạng = Nhóm VT của dòng
      - Khóa  -> $<cột Mã hiệu><r> khi tra theo Mã hiệu và mã = ô Mã hiệu của dòng
    Viết cứng (có escape) khi: tra theo Tên công việc (mã 0), số hạng ghép khác nhóm/khác mã của dòng,
    hoặc giá trị chứa ký tự đại diện ~ * ? (ô tham chiếu không escape được)."""
    sg, sv = str(g if g is not None else ''), str(val if val is not None else '')
    if r and gcol and row_g is not None and sg == str(row_g) and not WILD.search(sg): gs = '$%s%d' % (CL(gcol), r)
    else: gs = '"%s"' % esc(sg)
    if r and ccol and key == 'Mã hiệu' and row_code is not None and sv == str(row_code) and sv and not WILD.search(sv):
        vs = '$%s%d' % (CL(ccol), r)
    else: vs = '"%s"' % esc(sv)
    return 'SUMIFS(TV[%s],TV[Column1],%s,TV[%s],%s)' % (lc, gs, key, vs)


# ---------------------------------------------------------------- column detection
def detect_cols(pkg, sheet, hdr_rows=(1, 12)):
    g = pkg.grid(sheet, *hdr_rows)
    norm = lambda s: re.sub(r'\s+', ' ', str(s or '')).strip().lower()
    by = collections.defaultdict(list)
    for ref, v in g.items():
        if v is None or str(v).startswith('='): continue
        by[norm(v)].append((int(re.search(r'\d+', ref).group(0)), colnum(ref)))
    def first(name, after=0):
        cands = sorted(c for r, c in by.get(norm(name), []) if c > after)
        return cands[0] if cands else None
    _c = sorted(c for r, c in by.get(norm('Nội Dung Công Việc'), []))
    name = _c[-1] if _c else None   # phần TDI nằm bên phải; CĐT có thể trùng tiêu đề
    if not name: raise SystemExit('Không thấy cột "Nội Dung Công Việc" (phần TDI) trong %s' % sheet)
    res = {'name': name, 'code': name + 1, 'brand': name + 2, 'unit_tdi': name + 3,
           'vl0': name + 4, 'nc0': name + 5, 'chung_loai': name + 6}
    dg = first('Đơn giá', name + 6)          # nhóm đơn giá sau hệ số
    if dg: res.update(vl=dg, nc=dg + 1, dg=dg + 2)
    for h in HELPERS:
        res[h] = first(h, name)
    # dòng tiêu đề chính (để tạo cột phụ)
    rows = [r for r, c in by.get(norm('Nội Dung Công Việc'), []) if c == name]
    res['hdr_row'] = rows[0] if rows else 5
    res['letters'] = {k: CL(v) for k, v in res.items() if isinstance(v, int) and k != 'hdr_row'}
    return res


def cmd_cols(f, sheet):
    print(json.dumps(detect_cols(Pkg(f), sheet)['letters'], ensure_ascii=False, indent=1))


def cmd_helpers(f, out, sheets):
    pkg = Pkg(f)
    for sh in sheets:
        c = detect_cols(pkg, sh)
        if all(c[h] for h in HELPERS): print(sh, 'đã có cột phụ', [CL(c[h]) for h in HELPERS]); continue
        # đặt sau cột có dữ liệu cuối cùng của dòng tiêu đề
        sd = pkg.root(sh).find(q('sheetData'))
        rows = {int(r.get('r')): r for r in sd}
        last = max(colnum(x.get('r')) for x in rows[c['hdr_row']])
        start = last + 1
        for i, h in enumerate(HELPERS):
            setc(rows[c['hdr_row']], start + i, h)
        pkg.touch(sh)
        print(sh, 'tạo cột phụ', CL(start), '->', CL(start + 4))
    pkg.save(out)


# ---------------------------------------------------------------- insert (bóc tách)
def cmd_insert(f, out, planfile):
    P = json.load(open(planfile))
    pkg = Pkg(f)
    plan = collections.defaultdict(list)
    for sh, rn, exp, ch, note, conf in P: plan[sh].append((rn, len(ch)))
    for sh in plan: plan[sh].sort()
    fmap = lambda sh, n: n + sum(k for p, k in plan.get(sh, []) if n > p)
    CELL = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
    ROWR = re.compile(r"^(\$?)(\d+):(\$?)(\d+)$")

    def shift_ref(ref, sh):
        if '[' in ref: return ref
        m = ROWR.match(ref)
        if m: return '%s%d:%s%d' % (m.group(1), fmap(sh, int(m.group(2))), m.group(3), fmap(sh, int(m.group(4))))
        return CELL.sub(lambda m: '%s%s%s%d' % (m.group(1), m.group(2), m.group(3), fmap(sh, int(m.group(4)))), ref)

    def shift_formula(fx, cur):
        tok = Tokenizer('=' + fx); outv = []; changed = False
        for t in tok.items:
            v = t.value
            if t.type == 'OPERAND' and t.subtype == 'RANGE':
                if '!' in v:
                    pre, ref = v.rsplit('!', 1); sn = pre.strip("'").replace("''", "'")
                else:
                    pre, ref, sn = None, v, cur
                if sn in plan:
                    nr = shift_ref(ref, sn)
                    if nr != ref: v = (pre + '!' if pre is not None else '') + nr; changed = True
            outv.append(v)
        return ''.join(outv) if changed else fx

    def rt_ok(fx):
        try: return ''.join(t.value for t in Tokenizer('=' + fx).items) == fx
        except Exception: return False
    # styles
    st = etree.fromstring(pkg.files['xl/styles.xml'])
    fonts, fills, xfs = st.find(q('fonts')), st.find(q('fills')), st.find(q('cellXfs'))
    fl = etree.SubElement(fills, q('fill')); pf = etree.SubElement(fl, q('patternFill'), patternType='solid')
    etree.SubElement(pf, q('fgColor'), rgb=FILL_CHILD); etree.SubElement(pf, q('bgColor'), indexed='64')
    fills.set('count', str(len(fills))); FILL = len(fills) - 1
    fcache, xcache = {}, {}

    def font_with(fid, rgb):
        if (fid, rgb) in fcache: return fcache[(fid, rgb)]
        fo = copy.deepcopy(fonts[fid]); c = fo.find(q('color'))
        if c is None:
            c = etree.Element(q('color')); order = ['b', 'i', 'strike', 'condense', 'extend', 'outline', 'shadow', 'u', 'vertAlign', 'sz']
            pos = 0
            for i, ch in enumerate(fo):
                if etree.QName(ch).localname in order: pos = i + 1
            fo.insert(pos, c)
        for a in list(c.attrib): del c.attrib[a]
        c.set('rgb', rgb); fonts.append(fo); fonts.set('count', str(len(fonts)))
        fcache[(fid, rgb)] = len(fonts) - 1; return fcache[(fid, rgb)]

    def xf_child(s, rgb=None):
        if (s, rgb) in xcache: return xcache[(s, rgb)]
        x = copy.deepcopy(xfs[int(s or 0)]); x.set('fillId', str(FILL)); x.set('applyFill', '1')
        if rgb: x.set('fontId', str(font_with(int(x.get('fontId', '0')), rgb))); x.set('applyFont', '1')
        xfs.append(x); xfs.set('count', str(len(xfs))); xcache[(s, rgb)] = len(xfs) - 1; return xcache[(s, rgb)]

    def unshare(sd):
        masters = {}
        for row in sd:
            for c in row:
                fx = c.find(q('f'))
                if fx is not None and fx.get('t') == 'shared' and fx.text: masters[fx.get('si')] = (fx.text, c.get('r'))
        for row in sd:
            for c in row:
                fx = c.find(q('f'))
                if fx is not None and fx.get('t') == 'shared':
                    txt, org = masters[fx.get('si')]
                    if not fx.text: fx.text = Translator('=' + txt, origin=org).translate_formula(c.get('r'))[1:]
                    for a in ('t', 'si', 'ref'):
                        if a in fx.attrib: del fx.attrib[a]

    report = []
    for name, p in pkg.sheets:
        root = pkg.root(name); sd = root.find(q('sheetData')); changed = name in plan
        if changed: unshare(sd)
        bad = 0
        for row in sd:
            for c in row:
                fx = c.find(q('f'))
                if fx is not None and fx.text:
                    if not rt_ok(fx.text): bad += 1; continue
                    nf = shift_formula(fx.text, name)
                    if nf != fx.text: fx.text = nf; changed = True
        if bad: print('CẢNH BÁO: %s có %d công thức tokenizer không tái tạo được, không dời' % (name, bad))
        if name in plan:
            C = detect_cols(pkg, name)
            miss = [h for h in HELPERS if not C[h]]
            if miss: raise SystemExit('%s thiếu cột phụ %s - chạy lệnh helpers trước' % (name, miss))
            nm, cd, un, vl0, nc0 = C['name'], C['code'], C['unit_tdi'], C['vl0'], C['nc0']
            hg, hc, hn, hk, hl = (C[h] for h in HELPERS)
            KL, VL, NC = CL(hk), CL(vl0), CL(nc0)
            rows = {int(r.get('r')): r for r in sd}
            for sh, rn, exp, ch, note, conf in P:
                if sh != name: continue
                desc = next((pkg.val(c) for c in rows[rn] if c.get('r') == 'B%d' % rn), None)
                if exp and not str(desc or '').strip().startswith(exp):
                    raise SystemExit('Dòng cha lệch %s!%d: cần "%s", thấy "%s"' % (sh, rn, exp, desc))
            for r in sorted(rows, reverse=True):
                nr = fmap(name, r)
                if nr != r:
                    rows[r].set('r', str(nr))
                    for c in rows[r]: c.set('r', re.sub(r'\d+$', str(nr), c.get('r')))
            for sh, rn, exp, ch, note, conf in P:
                if sh != name: continue
                prow = rows[rn]; pr = int(prow.get('r')); idx = list(sd).index(prow)
                first, last = pr + 1, pr + len(ch)
                for j, c in enumerate(ch):
                    nr = pr + 1 + j
                    row = etree.Element(q('row')); row.set('r', str(nr))
                    if prow.get('spans'): row.set('spans', prow.get('spans'))
                    for pc in prow:
                        col = colnum(pc.get('r'))
                        ncell = etree.SubElement(row, q('c')); ncell.set('r', CL(col) + str(nr))
                        if col >= nm:   # chỉ tô phần TDI
                            ncell.set('s', str(xf_child(pc.get('s'), COLOR_NAME if col == nm else COLOR_CODE if col == cd else None)))
                        elif pc.get('s'):
                            ncell.set('s', pc.get('s'))
                    if prow.get('s'): row.set('s', prow.get('s')); row.set('customFormat', '1')
                    setc(row, nm, c['name'])
                    ctext = c.get('ma') or ('(mã 0) ' + c['val'] if c.get('key') == 'Tên công việc' else '')
                    setc(row, cd, ctext)
                    setc(row, un, c.get('unit'))
                    if c.get('key'):
                        for col, lc in ((vl0, 'Giá VL gốc'), (nc0, 'Giá NC gốc')):
                            setc(row, col, formula=tv_sumifs(lc, c['g'], c['key'], c['val'], nr, hg, c['g'], cd, ctext))
                    setc(row, hg, c['g']); setc(row, hc, c['conf'])
                    if c.get('note'): setc(row, hn, c['note'])
                    setc(row, hk, num=c.get('coef', 1)); setc(row, hl, 'Con')
                    sd.insert(idx + 1 + j, row)
                setc(prow, cd, '')
                for col, L in ((vl0, VL), (nc0, NC)):
                    setc(prow, col, formula='SUMPRODUCT($%s$%d:$%s$%d,%s%d:%s%d)' % (KL, first, KL, last, L, first, L, last))
                confs = [c['conf'] for c in ch]
                pconf = max(confs, key=lambda x: CONF_ORDER.index(x) if x in CONF_ORDER else 3)
                if 'Chưa khớp' in confs and len(set(confs)) > 1:
                    pconf = 'Thấp'; note += '; còn thành phần chưa có giá, đơn giá cụm chưa đủ'
                setc(prow, hg, ' + '.join(dict.fromkeys(c['g'] for c in ch))); setc(prow, hc, pconf)
                setc(prow, hn, note + '. Đơn giá = SUMPRODUCT(Hệ số ĐV, đơn giá) các dòng con bên dưới')
                setc(prow, hk, ''); setc(prow, hl, 'Cụm')
                report.append((name, rn, pr, len(ch)))
            for mc in root.iter(q('mergeCell')): mc.set('ref', shift_ref(mc.get('ref'), name))
            for cf in root.iter(q('conditionalFormatting')):
                cf.set('sqref', ' '.join(shift_ref(x, name) for x in cf.get('sqref').split()))
                for fo in cf.iter(q('formula')): fo.text = shift_formula(fo.text, name)
            for tag in ('autoFilter', 'dimension'):
                for e in root.iter(q(tag)): e.set('ref', shift_ref(e.get('ref'), name))
            for e in root.iter(q('selection')):
                if e.get('sqref'): e.set('sqref', ' '.join(shift_ref(x, name) for x in e.get('sqref').split()))
                if e.get('activeCell'): e.set('activeCell', shift_ref(e.get('activeCell'), name))
        if changed: pkg.touch(name)
    for dn in pkg.wb.iter(q('definedName')):
        if dn.text and '!' in dn.text and not dn.text.startswith('{'):
            try:
                if rt_ok(dn.text): dn.text = shift_formula(dn.text, '')
                else: print('CẢNH BÁO: không dời được defined name', dn.get('name'))
            except Exception as e:
                print('CẢNH BÁO: defined name %s: %s' % (dn.get('name'), e))
    pkg.new['xl/styles.xml'] = etree.tostring(st, xml_declaration=True, encoding='UTF-8', standalone=True)
    _drop_calcchain(pkg)
    pkg.save(out, drop=('xl/calcChain.xml',))
    json.dump(report, open(out + '.insert_report.json', 'w'), ensure_ascii=False)
    print('Đã bóc tách %d cụm, %d dòng con. Bản đồ dòng: %s.insert_report.json' % (len(report), sum(r[3] for r in report), out))


def _drop_calcchain(pkg):
    rels = etree.fromstring(pkg.files['xl/_rels/workbook.xml.rels'])
    for r in list(rels):
        if 'calcChain' in r.get('Target'): rels.remove(r)
    pkg.new['xl/_rels/workbook.xml.rels'] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
    ct = pkg.files['[Content_Types].xml'].decode('utf8')
    pkg.new['[Content_Types].xml'] = re.sub(r'<Override[^>]*calcChain[^>]*/>', '', ct).encode('utf8')


# ---------------------------------------------------------------- fill CĐT
def cmd_fillcdt(f, out, cfgfile):
    CFG = json.load(open(cfgfile)); pkg = Pkg(f)
    for sh, c in CFG.items():
        C = detect_cols(pkg, sh)
        VL, NC, LO = CL(C['vl']), CL(C['nc']), CL(C['Loại dòng']) if C['Loại dòng'] else None
        Hc, Ic, Jc, U = c.get('h', 'H'), c.get('i', 'I'), c.get('j', 'J'), c.get('unit', 'C')
        sd = pkg.root(sh).find(q('sheetData')); rows = {int(r.get('r')): r for r in sd}
        get = lambda r, col: next((pkg.val(x) for x in rows.get(r, []) if x.get('r') == '%s%d' % (col, r)), None)
        n, over, tot = 0, [], None
        if c.get('single'):   # CĐT chỉ có 1 cột đơn giá + 1 cột thành tiền (vd. MTT: E=KL, F=Đơn giá, G=Thành tiền)
            Fc, Gc, KLc = c['single'].get('f', 'F'), c['single'].get('g', 'G'), c['kl'][0]
            DG = CL(C['dg']); gadd = []
            for r in sorted(rows):
                if r < c.get('start', 8): continue
                if str(get(r, U) or '').strip() in ('', '\xa0'): continue
                if LO and get(r, LO) == 'Con': continue
                v = get(r, Fc)
                if v not in (None, '') and not str(v).startswith('=N('): over.append('%s%d=%s' % (Fc, r, str(v)[:30]))
                setc(rows[r], CI(Fc), formula='N(%s%d)' % (DG, r))   # số (0 khi chưa có giá) để G = E*F không ra #VALUE!
                gc = next((x for x in rows[r] if x.get('r') == '%s%d' % (Gc, r)), None)
                if gc is None or (gc.find(q('f')) is None and gc.find(q('v')) is None and gc.find(q('is')) is None):
                    setc(rows[r], CI(Gc), formula='%s%d*%s%d' % (KLc, r, Fc, r)); gadd.append(r)
                n += 1
            pkg.touch(sh)
            print('%s: %d dòng, ô CĐT bị ghi đè %d %s, thêm công thức %s ở %s' % (sh, n, len(over), over[:15], Gc, gadd))
            continue
        for r in sorted(rows):
            if r < c.get('start', 8): continue
            if str(get(r, CL(C['name'])) or '').startswith('Tổng cộng trước thuế') or str(get(r, CL(C['unit_tdi'])) or '').startswith('Tổng cộng trước thuế'):
                tot = r; continue
            if str(get(r, U) or '').strip() == '': continue
            if LO and get(r, LO) == 'Con': continue
            for col in (Hc, Ic):
                v = get(r, col)
                if v not in (None, '') and not str(v).startswith('=IF(N('): over.append('%s%d=%s' % (col, r, str(v)[:30]))
            row = rows[r]
            setc(row, CI(Hc), formula='IF(N(%s%d)=0,"",%s%d)' % (VL, r, VL, r))
            setc(row, CI(Ic), formula='IF(N(%s%d)=0,"",%s%d)' % (NC, r, NC, r))
            setc(row, CI(Jc), formula='IF(COUNT(%s%d:%s%d)=0,"",SUM(%s%d:%s%d))' % (Hc, r, Ic, r, Hc, r, Ic, r))
            for kl, z in zip(c['kl'], c['zones']):
                a = CI(z); A, B, S = CL(a), CL(a + 1), CL(a + 2)
                setc(row, a, formula='IF(OR(%s%d="",N(%s%d)=0),"",%s%d*%s%d)' % (Hc, r, kl, r, Hc, r, kl, r))
                setc(row, a + 1, formula='IF(OR(%s%d="",N(%s%d)=0),"",%s%d*%s%d)' % (Ic, r, kl, r, Ic, r, kl, r))
                setc(row, a + 2, formula='IF(COUNT(%s%d:%s%d)=0,"",SUM(%s%d:%s%d))' % (A, r, B, r, A, r, B, r))
            n += 1
        if tot:
            for z in c['zones']:
                for k in range(3):
                    L = CL(CI(z) + k); setc(rows[tot], CI(L), formula='SUBTOTAL(9,%s%d:%s%d)' % (L, c.get('start', 8), L, tot - 1))
        pkg.touch(sh)
        print('%s: %d dòng, dòng tổng %s, ô CĐT bị ghi đè %d %s' % (sh, n, tot, len(over), over[:15]))
        if len(c['zones']) != len(c['kl']): print('  LƯU Ý: số khu giá trị khác số cột KL - kiểm tra lại cấu hình')
    pkg.save(out)


# ---------------------------------------------------------------- recalc check
def cmd_recalc(f, sheets=None):
    import openpyxl
    tmp = tempfile.mkdtemp()
    nc = os.path.join(tmp, 'nocache.xlsx')
    zin = zipfile.ZipFile(f); zo = zipfile.ZipFile(nc, 'w', zipfile.ZIP_DEFLATED)
    pat = re.compile(r'(<f>[^<]*</f>|<f [^>]*>[^<]*</f>|<f [^>]*/>)<v>[^<]*</v>')
    for it in zin.infolist():
        d = zin.read(it.filename)
        if it.filename.startswith('xl/worksheets/sheet'): d = pat.sub(r'\1', d.decode('utf8')).encode('utf8')
        zo.writestr(it, d)
    zo.close()
    od = os.path.join(tmp, 'out'); os.makedirs(od)
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', od, nc], capture_output=True, timeout=600)
    wbv = openpyxl.load_workbook(os.path.join(od, 'nocache.xlsx'), data_only=True)
    wbf = openpyxl.load_workbook(f, read_only=False)
    pkg = Pkg(f)
    for sh in (sheets or [s for s in wbv.sheetnames if s in wbf.sheetnames]):
        ws, wf = wbv[sh], wbf[sh]
        errs = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('#')]
        msg = '%s: ô lỗi %d %s' % (sh, len(errs), errs[:8])
        try:
            C = detect_cols(pkg, sh)
            no = []
            for r in range(C['hdr_row'] + 3, ws.max_row + 1):
                code = wf.cell(r, C['code']).value
                if code and str(code).strip() and not str(code).startswith('='):
                    a, b = ws.cell(r, C['vl0']).value, ws.cell(r, C['nc0']).value
                    if not (isinstance(a, (int, float)) and a) and not (isinstance(b, (int, float)) and b): no.append(r)
            msg += ' | có mã không giá %d %s' % (len(no), no[:15])
            for r in range(ws.max_row, 0, -1):
                if str(ws.cell(r, C['name']).value or ws.cell(r, C['unit_tdi']).value or '').startswith('Tổng cộng trước thuế'):
                    vals = [(CL(cc), ws.cell(r, cc).value) for cc in range(1, ws.max_column + 1) if isinstance(ws.cell(r, cc).value, (int, float)) and ws.cell(r, cc).value]
                    msg += ' | dòng tổng %d: %s' % (r, vals[:20]); break
        except SystemExit:
            pass
        print(msg)
    shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- cleanup temp sheets
def cmd_cleanup(f, out):
    pkg = Pkg(f)
    shs = pkg.wb.find(q('sheets'))
    rels = etree.fromstring(pkg.files['xl/_rels/workbook.xml.rels'])
    drop = set(); removed = []
    for s in list(shs):
        n = s.get('name')
        if n == '_TDI_JOB' or n.startswith('_TDI_TMP'):
            rid = s.get('{%s}id' % RNS); shs.remove(s); removed.append(n)
            for r in list(rels):
                if r.get('Id') == rid:
                    t = r.get('Target').lstrip('/'); t = t if t.startswith('xl/') else 'xl/' + t
                    drop.add(t); rels.remove(r)
                    rp = t.replace('worksheets/', 'worksheets/_rels/') + '.rels'
                    drop.add(rp)
    if not removed: print('Không có sheet tạm'); shutil.copy(f, out); return
    # defined names gắn localSheetId cần đánh lại -> bỏ các name của sheet tạm, đánh lại chỉ số
    names = [s.get('name') for s in shs]
    old = [n for n, _ in pkg.sheets]
    dns = pkg.wb.find(q('definedNames'))
    if dns is not None:
        for dn in list(dns):
            lid = dn.get('localSheetId')
            if lid is not None:
                on = old[int(lid)]
                if on in removed: dns.remove(dn)
                else: dn.set('localSheetId', str(names.index(on)))
            elif dn.text and any(r in dn.text for r in removed): dns.remove(dn)
    bv = pkg.wb.find(q('bookViews'))
    if bv is not None:
        for v in bv:
            for a in ('activeTab', 'firstSheet'):
                if v.get(a) and int(v.get(a)) >= len(names): v.set(a, '0')
    pkg.new['xl/_rels/workbook.xml.rels'] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
    ct = pkg.files['[Content_Types].xml'].decode('utf8')
    for d in drop: ct = re.sub(r'<Override[^>]*PartName="/%s"[^>]*/>' % re.escape(d), '', ct)
    pkg.new['[Content_Types].xml'] = ct.encode('utf8')
    pkg.save(out, drop=drop)
    print('Đã xóa sheet tạm:', removed)


if __name__ == '__main__':
    a = sys.argv[1:]
    if not a: print(__doc__); sys.exit()
    cmd = a[0]
    if cmd == 'sheets':
        for n, p in Pkg(a[1]).sheets: print('%-40s %s' % (n, p))
    elif cmd == 'cols': cmd_cols(a[1], a[2])
    elif cmd == 'helpers': cmd_helpers(a[1], a[2], a[3:])
    elif cmd == 'insert': cmd_insert(a[1], a[2], a[3])
    elif cmd == 'fillcdt': cmd_fillcdt(a[1], a[2], a[3])
    elif cmd == 'recalc':
        sh = a[a.index('--sheets') + 1].split('|') if '--sheets' in a else None
        cmd_recalc(a[1], sh)
    elif cmd == 'cleanup': cmd_cleanup(a[1], a[2])
    else: print(__doc__)
