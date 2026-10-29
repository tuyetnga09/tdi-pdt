#!/usr/bin/env python3
"""boq_memory.py - Bảng khớp đã duyệt: lưu kết quả gán mã/giá đã duyệt và dùng lại cho lần chạy sau / dự án sau.
Cần boq_xml_tools.py cùng thư mục.

  export FILE MEM.xlsx --project CODE [--fix FIX.json]   : xuất kết quả đã duyệt của FILE vào MEM (thay dòng cũ cùng dự án)
  apply  FILE OUT MEM.xlsx [--project CODE]              : áp bảng khớp vào FILE -> OUT, kèm OUT.todo.json, OUT.plan.json
  write  FILE OUT ASSIGN.json                            : ghi các dòng Claude tự gán (cùng định dạng, công thức chuẩn)
  diff   FILE MEM.xlsx                                   : so FILE với bảng khớp, liệt kê dòng mất mã/giá hoặc đổi kết quả

Khóa khớp: (mô tả CĐT, ĐVT, tiêu đề phụ gần nhất, Chủng loại) -> nới dần: bỏ Chủng loại -> bỏ tiêu đề phụ (chỉ khi kết quả duy nhất).
ASSIGN.json: [{"sheet","row","g","terms":[[g,"Mã hiệu"|"Tên công việc",val],...] hoặc "ma"+"keycol"+"keyval",
               "code_text","conf","note","coef","loai"}]
"""
import sys, os, re, json, unicodedata, collections, argparse
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import boq_xml_tools as T
from openpyxl.utils import get_column_letter as CL

MEM_COLS = ['Dự án', 'Sheet', 'Dòng', 'STT', 'Mô tả CĐT', 'ĐVT', 'Tiêu đề phụ', 'Chủng loại', 'Ghi chú CĐT',
            'Nhóm VT', 'Mã hiệu', 'Giá VL (terms)', 'Giá NC (terms)', 'Độ tin cậy', 'Ghi chú gán mã', 'Hệ số ĐV', 'Loại dòng']
CUM_COLS = ['Dự án', 'Sheet', 'Dòng cha', 'Mô tả cha', 'ĐVT cha', 'Tiêu đề phụ', 'Chủng loại', 'Ghi chú cha', 'Độ tin cậy cha',
            'Thứ tự', 'g', 'key', 'val', 'ma', 'name', 'unit', 'coef', 'conf', 'note']
SUMIFS = re.compile(r'SUMIFS\(TV\[(Giá [VN][LC][^\]]*)\],TV\[Column1\],(\$?[A-Z]+\$?\d+|"(?:[^"]|"")*"),TV\[(Mã hiệu|Tên công việc)\],"((?:[^"]|"")*)"\)')


def norm(s):
    s = unicodedata.normalize('NFC', str(s or '')).lower()
    return re.sub(r'\s+', ' ', s).strip(' .;:-')


def unq(s): return s.replace('""', '"')


# ------------------------------------------------------------ đọc BOQ
def read_sheet(path, sheet):
    """Trả về (cols, rows): rows = list dict cho mọi dòng công việc (kể cả Con), giữ thứ tự."""
    pkg = T.Pkg(path); C = T.detect_cols(pkg, sheet)
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[sheet]; ws.reset_dimensions()   # read_only cắt theo <dimension>, cột phụ mới thêm sẽ bị mất nếu không reset
    hidx = {h: C[h] for h in T.HELPERS}
    remark = None; ucol = 3; scol = None
    for r in ws.iter_rows(min_row=1, max_row=C['hdr_row'] + 2, max_col=C['name'] - 1):
        for c in r:
            if isinstance(c.value, str) and re.search(r'ghi ch[uú]|remark', c.value, re.I): remark = c.column
            if isinstance(c.value, str) and re.fullmatch(r'(?i)\s*(đơn vị|đvt|unit)\s*', c.value): ucol = c.column
            if isinstance(c.value, str) and re.fullmatch(r'(?i)\s*(thông số|quy cách|specification)\s*', c.value): scol = c.column
    out, heading = [], ''
    for i, r in enumerate(ws.iter_rows(min_row=C['hdr_row'] + 3, values_only=True), C['hdr_row'] + 3):
        g = lambda col: r[col - 1] if col and len(r) >= col else None
        desc, unit = g(2), g(ucol)
        is_con = C['Loại dòng'] and g(C['Loại dòng']) == 'Con'
        if (unit in (None, '') or str(unit).strip() == '') and not is_con:
            if desc not in (None, '') and not str(desc).startswith('=') and not re.match(r'(?i)phụ kiện', str(desc).strip()):
                heading = str(desc).strip()
            continue
        out.append(dict(row=i, stt=g(1), desc=str(desc or '').strip(), unit=str(unit or '').strip(), heading=heading,
                        cl=g(C['chung_loai']), remark=g(remark) if remark else None, spec=g(scol) if scol else None,
                        code=g(C['code']), name=g(C['name']), unit_tdi=g(C['unit_tdi']), vl=g(C['vl0']), nc=g(C['nc0']),
                        **{h: g(c) for h, c in hidx.items()}))
    return C, out


def terms_of(f, group_of_row):
    if not isinstance(f, str) or not f.startswith('='): return []
    out = []
    for m in SUMIFS.finditer(f):
        g = m.group(2)
        g = unq(g[1:-1]) if g.startswith('"') else group_of_row
        out.append([g, m.group(3), re.sub(r'~([~*?])', r'\1', unq(m.group(4)))])
    return out


def keys(d):
    a, u, h, c = norm(d['desc']), norm(d['unit']), norm(d['heading']), norm(d.get('cl'))
    return [('k4', a, u, h, c), ('k3', a, u, h), ('k2', a, u)]


# ------------------------------------------------------------ export
def cmd_export(f, mem, project, fixfile=None):
    fixes = json.load(open(fixfile)) if fixfile else {}
    wb = openpyxl.load_workbook(f, read_only=True)
    M, CU = [], []
    for sh in wb.sheetnames:
        try: C, rows = read_sheet(f, sh)
        except SystemExit: continue
        if not any(r.get('Nhóm VT') or r.get('Độ tin cậy') for r in rows): continue
        cur = None; k = 0
        for d in rows:
            fx = fixes.get('%s!%d' % (sh, d['row']))
            if d.get('Loại dòng') == 'Con':
                if cur is None: continue
                k += 1
                tv = terms_of(d['vl'], d.get('Nhóm VT')) or terms_of(d['nc'], d.get('Nhóm VT'))
                g, key, val = (tv[0] if tv else [d.get('Nhóm VT'), None, None])
                CU.append([project, sh, cur['row'], cur['desc'], cur['unit'], cur['heading'], cur['cl'], cur.get('Ghi chú gán mã'),
                           cur.get('Độ tin cậy'), k, g, key, val, d['code'] if key == 'Mã hiệu' else '', d['name'], d['unit_tdi'],
                           d.get('Hệ số ĐV') or 1, d.get('Độ tin cậy'), d.get('Ghi chú gán mã')])
                continue
            cur, k = (d if d.get('Loại dòng') == 'Cụm' else None), 0
            vt, nt = terms_of(d['vl'], d.get('Nhóm VT')), terms_of(d['nc'], d.get('Nhóm VT'))
            rec = dict(g=d.get('Nhóm VT'), code=d['code'], vt=vt, nt=nt, conf=d.get('Độ tin cậy'), note=d.get('Ghi chú gán mã'),
                       coef=d.get('Hệ số ĐV'), loai=d.get('Loại dòng'))
            if fx: rec.update(fx)
            if not (rec['g'] or rec['conf']): continue
            M.append([project, sh, d['row'], d['stt'], d['desc'], d['unit'], d['heading'], d['cl'], d['remark'],
                      rec['g'], rec['code'], json.dumps(rec['vt'], ensure_ascii=False) if rec['vt'] else '',
                      json.dumps(rec['nt'], ensure_ascii=False) if rec['nt'] else '', rec['conf'], rec['note'], rec['coef'], rec['loai']])
    # gộp với bảng cũ: bỏ dòng cùng dự án
    oldM, oldC = [], []
    if os.path.exists(mem):
        w0 = openpyxl.load_workbook(mem, read_only=True)
        oldM = [list(r) for r in w0['Khop'].iter_rows(min_row=2, values_only=True) if r[0] != project]
        oldC = [list(r) for r in w0['Cum'].iter_rows(min_row=2, values_only=True) if r[0] != project]
    out = openpyxl.Workbook(); a = out.active; a.title = 'Khop'; a.append(MEM_COLS)
    for r in oldM + M: a.append(r)
    b = out.create_sheet('Cum'); b.append(CUM_COLS)
    for r in oldC + CU: b.append(r)
    for ws in (a, b):
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for j, w in enumerate([10, 16, 7, 7, 50, 7, 35, 10, 30, 14, 26, 30, 30, 11, 50, 8, 8, 8, 40], 1): ws.column_dimensions[CL(j)].width = w
    out.save(mem)
    print('Xuất %s: %d dòng khớp, %d dòng con của %d cụm -> %s (tổng %d dòng)' %
          (project, len(M), len(CU), len({(r[1], r[2]) for r in CU}), mem, len(oldM) + len(M)))


# ------------------------------------------------------------ load memory
def load_mem(mem, project=None):
    w = openpyxl.load_workbook(mem, read_only=True)
    idx = collections.defaultdict(list)
    for r in w['Khop'].iter_rows(min_row=2, values_only=True):
        d = dict(zip(MEM_COLS, r))
        d.update(desc=d['Mô tả CĐT'], unit=d['ĐVT'], heading=d['Tiêu đề phụ'], cl=d['Chủng loại'])
        for k in keys(d): idx[k].append(d)
    cum = collections.defaultdict(list)
    for r in w['Cum'].iter_rows(min_row=2, values_only=True):
        d = dict(zip(CUM_COLS, r)); cum[(d['Dự án'], d['Sheet'], d['Dòng cha'])].append(d)
    return idx, cum


def sig(d): return (d['Nhóm VT'], d['Mã hiệu'], d['Giá VL (terms)'], d['Giá NC (terms)'], d['Độ tin cậy'], d['Loại dòng'])


def lookup(idx, d, project=None):
    for k in keys(d):
        hits = idx.get(k, [])
        if not hits: continue
        if project:
            same = [h for h in hits if h['Dự án'] == project]
            if same: hits = same
        if len({sig(h) for h in hits}) == 1: return hits[0], k[0]
        if k[0] == 'k2': return None, 'trùng nhiều kết quả'
    return None, None


# ------------------------------------------------------------ writer
def tv_keys(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = next((wb[s] for s in wb.sheetnames if s.strip().lower() == '2 thu vien gia'), None)
    if ws is None: raise SystemExit('DỪNG: file chưa có sheet "2 Thu vien gia" (bảng TV). Nạp thư viện trước (skill chuyển template).')
    rows = ws.iter_rows(values_only=True); H = [str(x) for x in next(rows)]
    ig, im, it = H.index('Column1'), H.index('Mã hiệu'), H.index('Tên công việc')
    S = collections.Counter()
    for r in rows:
        S[(r[ig], 'Mã hiệu', str(r[im]))] += 1; S[(r[ig], 'Tên công việc', str(r[it]))] += 1
    return S


def formula(terms, col):
    return '+'.join('SUMIFS(TV[%s],TV[Column1],"%s",TV[%s],"%s")' % (col, T.esc(g), k, T.esc(v)) for g, k, v in terms)


def write_rows(f, out, assigns):
    pkg = T.Pkg(f); by = collections.defaultdict(list)
    for a in assigns: by[a['sheet']].append(a)
    for sh, lst in by.items():
        C = T.detect_cols(pkg, sh)
        miss = [h for h in T.HELPERS if not C[h]]
        if miss: raise SystemExit('%s thiếu cột phụ %s - chạy boq_xml_tools.py helpers trước' % (sh, miss))
        sd = pkg.root(sh).find(T.q('sheetData')); rows = {int(r.get('r')): r for r in sd}
        for a in lst:
            row = rows[a['row']]
            T.setc(row, C['code'], a.get('code_text') or '')
            for col, lc, key in ((C['vl0'], 'Giá VL gốc', 'vt'), (C['nc0'], 'Giá NC gốc', 'nt')):
                tm = a.get(key)
                if tm:
                    fx = formula(tm, lc)
                    if a.get('coef') not in (None, '', 1, 1.0) and a.get('loai') != 'Cụm': fx = '(%s)*%s' % (fx, a['coef'])
                    T.setc(row, col, formula=fx)
                else:
                    cell = next((c for c in row if T.colnum(c.get('r')) == col), None)
                    if cell is not None and (cell.findtext(T.q('f')) or '').startswith('SUMIFS(TV['): T.setc(row, col, '')
            T.setc(row, C['Nhóm VT'], a.get('g') or '')
            T.setc(row, C['Độ tin cậy'], a.get('conf') or '')
            T.setc(row, C['Ghi chú gán mã'], a.get('note') or '')
            if a.get('coef') not in (None, ''): T.setc(row, C['Hệ số ĐV'], num=float(a['coef']))
            else: T.setc(row, C['Hệ số ĐV'], '')
            T.setc(row, C['Loại dòng'], a.get('loai') or 'Dòng')
        pkg.touch(sh)
    pkg.save(out)


# ------------------------------------------------------------ apply
def cmd_apply(f, out, mem, project=None):
    idx, cum = load_mem(mem, project)
    TVK = tv_keys(f)
    wb = openpyxl.load_workbook(f, read_only=True)
    assigns, todo, plan, stat = [], [], [], collections.Counter()
    for sh in wb.sheetnames:
        try: C, rows = read_sheet(f, sh)
        except SystemExit: continue
        for d in rows:
            if d.get('Loại dòng') == 'Con': continue
            hit, how = lookup(idx, d, project)
            ctx = dict(sheet=sh, row=d['row'], stt=d['stt'], desc=d['desc'], unit=d['unit'], heading=d['heading'], cl=d['cl'], remark=d['remark'], spec=d.get('spec'))
            if not hit:
                todo.append(dict(ctx, reason=how or 'chưa có trong bảng khớp')); stat['todo'] += 1; continue
            vt = json.loads(hit['Giá VL (terms)']) if hit['Giá VL (terms)'] else []
            nt = json.loads(hit['Giá NC (terms)']) if hit['Giá NC (terms)'] else []
            bad = [t for t in vt + nt if TVK[tuple(t)] == 0]
            if bad:
                todo.append(dict(ctx, reason='mã trong bảng khớp không có ở thư viện dự án này: %s' % bad, suggest=hit['Mã hiệu'])); stat['todo'] += 1; continue
            dup = [t for t in vt + nt if TVK[tuple(t)] > 1]
            note = hit['Ghi chú gán mã'] or ''
            if hit['Loại dòng'] == 'Cụm':
                kids = cum.get((hit['Dự án'], hit['Sheet'], hit['Dòng']), [])
                kids.sort(key=lambda k: k['Thứ tự'])
                if not kids:   # cụm chưa có cấu hình: giữ nhãn, không giá
                    assigns.append(dict(sheet=sh, row=d['row'], g=hit['Nhóm VT'], code_text='', vt=[], nt=[], conf=hit['Độ tin cậy'],
                                        note=note, coef='', loai='Cụm'))
                    stat['cụm chưa cấu hình'] += 1; continue
                ch = [dict(g=k['g'], key=k['key'], val=k['val'], ma=k['ma'], name=k['name'], unit=k['unit'], coef=k['coef'] or 1,
                           conf=k['conf'], note=k['note'] or '') for k in kids]
                pnote = re.sub(r'\. Đơn giá = SUMPRODUCT.*$', '', note)
                plan.append([sh, d['row'], d['desc'][:15], ch, pnote, hit['Độ tin cậy']])
                assigns.append(dict(sheet=sh, row=d['row'], g=hit['Nhóm VT'], code_text='', vt=[], nt=[], conf=hit['Độ tin cậy'],
                                    note=pnote, coef='', loai='Cụm'))
                stat['cụm'] += 1; continue
            if dup: note = (note + '; ' if note else '') + 'CẢNH BÁO: khóa trùng trong thư viện %s' % dup
            assigns.append(dict(sheet=sh, row=d['row'], g=hit['Nhóm VT'], code_text=hit['Mã hiệu'], vt=vt, nt=nt, conf=hit['Độ tin cậy'],
                                note=note, coef=hit['Hệ số ĐV'], loai=hit['Loại dòng'] or 'Dòng'))
            stat[how] += 1
    write_rows(f, out, assigns)
    json.dump(todo, open(out + '.todo.json', 'w'), ensure_ascii=False, indent=0, default=str)
    json.dump(plan, open(out + '.plan.json', 'w'), ensure_ascii=False, indent=0, default=str)
    print('Cụm chưa có cấu hình (giữ nhãn, không giá): %d' % stat['cụm chưa cấu hình'])
    print('Áp bảng khớp: %d dòng (khớp đủ khóa %d, bỏ Chủng loại %d, chỉ mô tả+ĐVT %d), cụm %d; còn %d dòng cần Claude gán -> %s.todo.json; plan bóc tách -> %s.plan.json'
          % (len(assigns), stat['k4'], stat['k3'], stat['k2'], stat['cụm'], stat['todo'], out, out))


def cmd_write(f, out, assign_file):
    A = json.load(open(assign_file)); TVK = tv_keys(f)
    for a in A:
        if 'terms' in a: a['vt'] = a['nt'] = a.pop('terms')
        elif a.get('keycol'):
            a['vt'] = a['nt'] = [[a['g'], a['keycol'], a['keyval']]]
        for t in a.get('vt', []):
            if TVK[tuple(t)] == 0: raise SystemExit('DỪNG: %s!%d khóa %s không có trong TV' % (a['sheet'], a['row'], t))
            if TVK[tuple(t)] > 1: print('CẢNH BÁO: %s!%d khóa %s trùng %d dòng trong TV (SUMIFS cộng dồn)' % (a['sheet'], a['row'], t, TVK[tuple(t)]))
        a.setdefault('code_text', a.get('ma') or (('(mã 0) ' + a['keyval']) if a.get('keycol') == 'Tên công việc' else ''))
    write_rows(f, out, A); print('Đã ghi %d dòng' % len(A))


# ------------------------------------------------------------ diff
def cmd_diff(f, mem, project=None):
    idx, _ = load_mem(mem, project)
    wb = openpyxl.load_workbook(f, read_only=True)
    lost, changed, new = [], [], []
    for sh in wb.sheetnames:
        try: C, rows = read_sheet(f, sh)
        except SystemExit: continue
        for d in rows:
            if d.get('Loại dòng') == 'Con': continue
            hit, how = lookup(idx, d, project)
            if not hit: continue
            had = bool(hit['Mã hiệu']) or hit['Loại dòng'] == 'Cụm'
            has = bool(d['code']) or d.get('Loại dòng') == 'Cụm'
            line = '%s!%d %s | bảng khớp: %s %s | file: %s %s | %s' % (sh, d['row'], d['desc'][:45], hit['Mã hiệu'] or hit['Loại dòng'], hit['Độ tin cậy'],
                                                                   d['code'] or d.get('Loại dòng'), d.get('Độ tin cậy'), str(d.get('Ghi chú gán mã') or '')[:60])
            if had and not has: lost.append(line)
            elif str(hit['Mã hiệu'] or '') != str(d['code'] or '') or hit['Độ tin cậy'] != d.get('Độ tin cậy'): changed.append(line)
            elif not had and has: new.append(line)
    print('MẤT mã/giá so với bảng khớp: %d' % len(lost)); print('\n'.join(lost[:300]))
    print('\nĐỔI mã/độ tin cậy: %d' % len(changed)); print('\n'.join(changed[:300]))
    print('\nTHÊM mới: %d' % len(new)); print('\n'.join(new[:100]))


if __name__ == '__main__':
    ap = argparse.ArgumentParser(); ap.add_argument('cmd'); ap.add_argument('args', nargs='*')
    ap.add_argument('--project'); ap.add_argument('--fix')
    a = ap.parse_args()
    if a.cmd == 'export': cmd_export(a.args[0], a.args[1], a.project, a.fix)
    elif a.cmd == 'apply': cmd_apply(a.args[0], a.args[1], a.args[2], a.project)
    elif a.cmd == 'write': cmd_write(*a.args[:3])
    elif a.cmd == 'diff': cmd_diff(a.args[0], a.args[1], a.project)
    else: print(__doc__)
