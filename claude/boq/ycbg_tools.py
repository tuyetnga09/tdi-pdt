#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ycbg_tools.py - Tao YCBG (De nghi bao gia) theo he tu file tong hop dau muc chua co trong thu vien gia.

Dau vao  : file <ten BOQ>_Dau_muc_chua_co_trong_thu_vien_gia_tham_khao.xlsx (dau ra Buoc 6d cua skill
           boq-gan-ma-gia-thu-vien) + mau YCBG trang cua TDI + YCBG_CONFIG.json.
Dau ra   : moi he 1 file YCBG_<ma>_<He>.xlsx (1 sheet, cac nhom ngan bang dong tieu de vang)
           + YCBG_<ma>_TongHop.xlsx (theo doi, chi tiet dong, dong khong dua vao YCBG, canh bao).
Khong gui gi ra ngoai. Khong ghi file BOQ. Gia tham khao / ly do noi bo KHONG dua vao file gui NCC.

Lenh:
  read  SRC.xlsx ITEMS.json [--ma TDIxxxxx]      doc file tong hop -> ITEMS.json (+ ITEMS.json.extra.json)
  plan  ITEMS.json CONFIG.json                   bang xep he x nhom de nguoi dung duyet (diem dung 1)
  build MAU.xlsx ITEMS.json CONFIG.json OUTDIR [--cot-gia]
  check OUTDIR ITEMS.json CONFIG.json            phai ra KET QUA: DAT
"""
import sys, os, re, json, copy, math, unicodedata, datetime
from collections import OrderedDict, defaultdict
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter

YELLOW = PatternFill("solid", fgColor="FFFFFF00")
POS_RE = re.compile(r"(.+?)!(\d+)(?:,\s*|$)")      # ten sheet BOQ co the chua dau phay


# ----------------------------------------------------------------- tien ich
def norm(s):
    s = unicodedata.normalize("NFC", str(s if s is not None else "")).lower()
    s = re.sub(r"[\r\n\t•]+", " ", s)
    s = s.replace("×", "x").replace("mm²", "mm2").replace("ø", "d")
    s = re.sub(r"\b(dn|phi)\s*(\d)", r"d\2", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_unit(u):
    u = re.sub(r"[　-鿿＀-￯/]+", " ", str(u or ""))   # bo chu Han va dau /
    return re.sub(r"\s+", " ", u).strip().lower()


def clean(s):
    return re.sub(r"\s+", " ", str(s if s is not None else "")).strip()


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except Exception:
        return None


def safe_name(s):
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").replace("đ", "d").replace("Đ", "D")
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")


def parse_pos(text):
    return [f"{a.strip()}!{b}" for a, b in POS_RE.findall(str(text or ""))]


def find_header(ws, must, maxrow=12):
    for r in range(1, maxrow + 1):
        vals = [norm(c.value) for c in ws[r]]
        if all(any(m in v for v in vals) for m in must):
            return r, {v: i for i, v in enumerate(vals) if v}
    return None, None


def col(h, *keys):
    for k in keys:
        for name, i in h.items():
            if name.startswith(k):
                return i
    return None


def sheet_like(wb, key):
    for n in wb.sheetnames:
        if norm(n).startswith(norm(key)):
            return wb[n]
    return None


# ----------------------------------------------------------------- doc nguon
def read_source(path, ma_expect=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    items, excluded, warn = [], [], []
    meta = {"ma": None, "du_an": None}

    ws = sheet_like(wb, "Chưa có trong TV")
    if ws is None:
        raise SystemExit("Không thấy sheet 'Chưa có trong TV'. Đây không phải file tổng hợp Bước 6d của boq-gan-ma-gia-thu-vien.")
    for r in range(1, 4):
        t = str(ws.cell(r, 1).value or "")
        m = re.match(r"^(?:BOQ\s+)?(.+?)\s+[–-]\s+(TDI\d+)", t)
        if m:
            meta["du_an"], meta["ma"] = m.group(1).strip(), m.group(2)
    if ma_expect and meta["ma"] and ma_expect != meta["ma"]:
        raise SystemExit(f"Mã dự án không khớp: file nguồn là {meta['ma']}, bạn nói {ma_expect}. Dừng, hỏi lại người dùng.")

    hr, h = find_header(ws, ["id", "nhóm", "mô tả", "đvt"])
    cI, cG, cD, cU = col(h, "id"), col(h, "nhóm"), col(h, "mô tả"), col(h, "đvt")
    cK, cP, cR = col(h, "tổng kl", "kl"), col(h, "vị trí"), col(h, "lý do")
    cGia, cGT, cHang = col(h, "giá tham khảo"), col(h, "thành tiền"), col(h, "hãng")
    cTT, cLink = col(h, "tình trạng"), col(h, "link")
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        if row[cD] in (None, ""):
            continue
        pos = parse_pos(row[cP])
        if not pos:
            warn.append(f"Không đọc được vị trí BOQ của ID {row[cI]}: {clean(row[cD])[:60]}")
        items.append(dict(src="Chưa có trong TV", id=row[cI], nhom=clean(row[cG]) or None, desc=clean(row[cD]),
                          unit=clean(row[cU]), kl=num(row[cK]), pos=pos, quycach="",
                          ly_do=clean(row[cR]) if cR is not None else "",
                          tinh_trang=clean(row[cTT]) if cTT is not None else "",
                          gia_tk=num(row[cGia]) if cGia is not None else None,
                          gt=num(row[cGT]) if cGT is not None else None,
                          hang_tk=row[cHang] if cHang is not None else None,
                          link=row[cLink] if cLink is not None else None))
    main_keys = {(norm(i["desc"]), norm_unit(i["unit"])) for i in items}
    main_names = {norm(i["desc"]) for i in items}

    ws = sheet_like(wb, "Có mã")
    if ws is not None:
        hr, h = find_header(ws, ["mô tả", "đvt", "dòng"])
        cS, cL, cD, cU = col(h, "sheet"), col(h, "dòng"), col(h, "mô tả"), col(h, "đvt")
        cK, cM, cN = col(h, "tổng kl", "kl"), col(h, "mã"), col(h, "ghi chú")
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            if row[cD] in (None, ""):
                continue
            items.append(dict(src="Có mã – TV thiếu giá VL", id=None, nhom=None, desc=clean(row[cD]),
                              unit=clean(row[cU]), kl=num(row[cK]), pos=[f"{clean(row[cS])}!{int(num(row[cL]) or 0)}"],
                              quycach="", ly_do="Có mã TV nhưng thiếu giá VL: " + clean(row[cM]),
                              tinh_trang="Có mã TV, thiếu giá VL", gia_tk=None, gt=None, hang_tk=None, link=None))

    ws = sheet_like(wb, "Thành phần bóc tách")
    if ws is not None:
        hr, h = find_header(ws, ["sheet", "thành phần", "tình trạng"])
        cSh, cDC, cDCh, cTP = col(h, "sheet"), col(h, "dòng cha"), col(h, "đầu mục cha"), col(h, "thành phần")
        cU, cHS, cKLC, cKL, cTT = col(h, "đvt"), col(h, "hệ số"), col(h, "kl cha"), col(h, "kl thành phần"), col(h, "tình trạng")
        cCon = col(h, "dòng con")
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            if not row[cSh] or not str(row[cTT] or "").startswith("Chưa khớp"):
                continue
            name = clean(row[cTP])
            if norm(name) in main_names:
                excluded.append(dict(ly_do="Thành phần cụm chưa khớp, đã có trong danh mục chính", pos=f"{row[cSh]}!{row[cCon]}",
                                     desc=name, unit=row[cU], kl=num(row[cKL]), note=clean(row[cDCh])[:100]))
                continue
            kl = num(row[cKL])
            if kl is None and num(row[cHS]) is not None and num(row[cKLC]) is not None:
                kl = num(row[cHS]) * num(row[cKLC])
            items.append(dict(src="Thành phần bóc tách", id=None, nhom=None, desc=name, unit=clean(row[cU]), kl=kl,
                              pos=[f"{row[cSh]}!{row[cCon]}"], quycach="Thành phần của: " + clean(row[cDCh])[:80],
                              ly_do="Thành phần cụm chưa khớp", tinh_trang="Chưa khớp – cần báo giá",
                              gia_tk=None, gt=None, hang_tk=None, link=None))

    ws = sheet_like(wb, "Cần chiết tính")
    if ws is not None:
        hr, h = find_header(ws, ["sheet", "dòng", "mô tả"], 3)
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            if row[0]:
                excluded.append(dict(ly_do="Cần chiết tính (lô/gói/đơn vị đặc biệt), không xin giá NCC theo dòng",
                                     pos=f"{row[0]}!{row[1]}", desc=clean(row[2]), unit=clean(row[3]), kl=num(row[4]),
                                     note=clean(row[5])[:160]))
    if sheet_like(wb, "TĐ – linh kiện") or sheet_like(wb, "TĐ - linh kiện"):
        warn.append("File có sheet 'TĐ – linh kiện': skill này CHƯA xuất phụ lục cấu hình từng tủ. Linh kiện tủ cần gửi theo bản vẽ / sơ đồ nguyên lý riêng.")

    for it in items:
        if it["kl"] in (None, 0):
            warn.append(f"Thiếu KL hoặc KL = 0: {it['desc'][:70]} ({', '.join(it['pos'][:2])})")
    return items, excluded, warn, meta


def merge_dups(rows):
    """Gop dong trung (dau muc + quy cach + DVT). KL chi cong phan vi tri BOQ moi."""
    out = OrderedDict()
    for it in rows:
        k = (norm(it["desc"]), norm(it["quycach"]), norm_unit(it["unit"]))
        if k not in out:
            o = dict(it); o["pos"] = list(it["pos"]); o["ids"] = [it["id"]] if it["id"] is not None else []
            o["srcs"] = [it["src"]]; o["gt"] = it.get("gt") or 0
            out[k] = o
            continue
        o = out[k]
        new_pos = [p for p in it["pos"] if p not in o["pos"]]
        if new_pos or not it["pos"]:
            if o["kl"] is not None or it["kl"] is not None:
                o["kl"] = (o["kl"] or 0) + (it["kl"] or 0)
        o["pos"] += new_pos
        if it["id"] is not None:
            o["ids"].append(it["id"])
        if it["src"] not in o["srcs"]:
            o["srcs"].append(it["src"])
        o["gt"] = (o["gt"] or 0) + (it.get("gt") or 0)
    return list(out.values())


# ----------------------------------------------------------------- xep he
def pick_he(it, cfg):
    if it.get("he"):
        return it["he"]
    nhom, mo_ta = norm(it.get("nhom") or ""), norm(it.get("desc") or "")
    for rule in cfg.get("quy_tac_he", []):
        ok = True
        if rule.get("nhom") and not re.search(norm(rule["nhom"]), nhom):
            ok = False
        if ok and rule.get("mo_ta") and not re.search(norm(rule["mo_ta"]), mo_ta):
            ok = False
        if ok and (rule.get("nhom") or rule.get("mo_ta")):
            return rule["he"]
    return cfg.get("he_mac_dinh", "Khác")


def label_nhom(it, cfg):
    g = it.get("nhom") or cfg.get("nhom_mac_dinh", "Hạng mục khác")
    g = re.sub(r"^XD\s*[–-]\s*", "", g)
    return cfg.get("doi_ten_nhom", {}).get(g, g)


def group_all(items, cfg):
    """he -> nhom -> [dong da gop]; giu thu tu xuat hien."""
    by = OrderedDict()
    for it in items:
        it["_he"] = pick_he(it, cfg)
        by.setdefault(it["_he"], OrderedDict()).setdefault(label_nhom(it, cfg), []).append(it)
    out = OrderedDict()
    order = cfg.get("thu_tu_he", [])
    for he in sorted(by, key=lambda x: order.index(x) if x in order else 99):
        out[he] = OrderedDict((g, merge_dups(rows)) for g, rows in by[he].items())
    return out


# ----------------------------------------------------------------- mau YCBG
def locate(ws):
    HR = next((r for r in range(1, 60) if clean(ws.cell(r, 1).value).upper() == "STT"), None)
    if not HR:
        raise SystemExit("Mẫu YCBG không có dòng tiêu đề bảng (ô cột A = 'STT').")
    CY = next((r for r in range(HR + 1, ws.max_row + 1) if clean(ws.cell(r, 1).value).lower().startswith("chú ý")), None)
    HN = next((r for r in range(HR + 1, ws.max_row + 1) if clean(ws.cell(r, 1).value).upper().startswith("THÔNG TIN NGƯỜI NHẬN")), None)
    HD = next((r for r in range(HR + 1, ws.max_row + 1) if clean(ws.cell(r, 1).value).upper().startswith("THÔNG TIN XUẤT HÓA ĐƠN")), None)
    if not (CY and HN and HD):
        raise SystemExit("Mẫu YCBG thiếu 'Chú ý' / 'THÔNG TIN NGƯỜI NHẬN HÀNG' / 'THÔNG TIN XUẤT HÓA ĐƠN'.")
    FE = next((r for r in range(HD, min(HD + 12, ws.max_row) + 1) if clean(ws.cell(r, 2).value).lower().startswith("điện thoại")), HD + 4)
    return HR, CY, HN, HD, FE


def snapshot(ws, r1, r2, ncol):
    rows = []
    for r in range(r1, r2 + 1):
        rows.append(([(ws.cell(r, c).value, copy.copy(ws.cell(r, c)._style)) for c in range(1, ncol + 1)],
                     ws.row_dimensions[r].height))
    merges = [(m.min_row - r1, m.min_col, m.max_row - r1, m.max_col) for m in ws.merged_cells.ranges
              if m.min_row >= r1 and m.max_row <= r2]
    return rows, merges


def insert_blank(snap, at, n):
    rows, merges = snap
    cells, h = rows[at]
    blank = ([(None, st) for (_, st) in cells], h)
    rows = rows[:at] + [blank] * n + rows[at:]
    merges = [(a + (n if a >= at else 0), c1, b + (n if b >= at else 0), c2) for (a, c1, b, c2) in merges]
    return rows, merges


def clear_rows(ws, r1, r2, ncol):
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= r1:
            ws.unmerge_cells(str(m))
    for r in range(r1, r2 + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell._style = StyleArray()
        ws.row_dimensions[r].height = None


def paste(ws, snap, r0):
    rows, merges = snap
    for i, (cells, h) in enumerate(rows):
        for c, (v, st) in enumerate(cells, start=1):
            ws.cell(r0 + i, c).value = v
            ws.cell(r0 + i, c)._style = copy.copy(st)
        ws.row_dimensions[r0 + i].height = h
    for (a, c1, b, c2) in merges:
        ws.merge_cells(start_row=r0 + a, start_column=c1, end_row=r0 + b, end_column=c2)


def fill_labels(ws, r1, r2, mapping):
    """mapping: [(regex nhan o cot B, gia tri)] -> ghi ': gia tri' vao cot C."""
    for rx, val in mapping:
        if not val:
            continue
        for r in range(r1, r2 + 1):
            if re.search(rx, clean(ws.cell(r, 2).value), re.I):
                ws.cell(r, 3).value = f": {val}"
                break


def est_height(pairs, base=30.0, line=14.5):
    lines = 1
    for text, w in pairs:
        if text:
            lines = max(lines, sum(max(1, math.ceil(len(seg) / max(w * 1.05, 1))) for seg in str(text).split("\n")))
    return max(base, lines * line + 3)


def check_config(cfg):
    bad = []
    def walk(x, path):
        if isinstance(x, dict):
            for k, v in x.items(): walk(v, f"{path}.{k}")
        elif isinstance(x, list):
            for i, v in enumerate(x): walk(v, f"{path}[{i}]")
        elif isinstance(x, str) and re.search(r"<[^<>]+>", x) and not path.endswith(".mo_ta") and not path.endswith(".nhom"):
            bad.append(f"{path} = {x!r}")
    walk(cfg, "cfg")
    for k in ("du_an", "ma_du_an", "dia_chi_giao_hang"):
        if not clean(cfg.get(k)):
            bad.append(f"cfg.{k} đang trống")
    nn = cfg.get("nguoi_nhan", {})
    if not clean(nn.get("_mac_dinh")):
        bad.append("cfg.nguoi_nhan._mac_dinh đang trống")
    if bad:
        raise SystemExit("Cấu hình chưa điền đủ, dừng để không in dấu <...> ra file YCBG:\n  " + "\n  ".join(bad))


# ----------------------------------------------------------------- dung 1 file he
def build_one(template, he, groups, cfg, outpath, so, cot_gia):
    wb = openpyxl.load_workbook(template)
    ws = wb.worksheets[0]
    ws.title = (cfg.get("ten_sheet", {}).get(he) or he)[:31]
    NC = 8
    HR, CY, HN, HD, FE = locate(ws)
    FR = HR + 1
    item_style = [copy.copy(ws.cell(FR, c)._style) for c in range(1, NC + 1)]
    foot = snapshot(ws, CY, FE, NC)
    chu_y = cfg.get("chu_y", [])
    free = (HN - CY) - 2                      # dong trong that su giua 'Chu y' va khoi nguoi nhan (chua 1 dong cach)
    if len(chu_y) > free:
        foot = insert_blank(foot, 1, len(chu_y) - free)
    idx_hn = HN - CY + max(0, len(chu_y) - free)
    idx_hd = HD - CY + max(0, len(chu_y) - free)
    clear_rows(ws, FR, max(ws.max_row, FE), NC + 6)

    # phan dau: So, Ngay, Vv
    for r in range(1, HR):
        a = None
        for c in range(1, NC + 1):
            v = clean(ws.cell(r, c).value)
            if v == "Số":
                tgt = ws.cell(r, c + 1); tgt.value = so
                tgt.font = Font(name="Times New Roman", size=11, bold=True)
                if c + 2 <= NC and ws.cell(r, c + 2).value is None:
                    ws.merge_cells(start_row=r, start_column=c + 1, end_row=r, end_column=c + 2)
            elif v.startswith("Ngày gửi"):
                if cfg.get("ngay_gui"):
                    ws.cell(r, c + 1).value = cfg["ngay_gui"]
            elif v == "Vv:":
                txt = (cfg.get("vv") or "Đề nghị báo giá {he} - Dự án {du_an}, quy cách như bảng dưới đây.").format(
                    he=cfg.get("ten_he", {}).get(he, he), du_an=cfg.get("du_an", ""))
                b = ws.cell(r, c + 1); b.value = txt
                ws.merge_cells(start_row=r, start_column=c + 1, end_row=r, end_column=NC)
                b.alignment = Alignment(wrap_text=True, vertical="top")
                ws.cell(r, c).alignment = Alignment(vertical="top")
                ws.row_dimensions[r].height = 31.5 if len(txt) > 90 else 18
    for k, v in (cfg.get("company_header") or {}).items():
        ws[k].value = v

    # cot gia NCC (tuy chon) va cot Ref
    extra = list(cfg.get("cot_gia", ["Đơn giá VL (chưa VAT)", "Đơn giá NC (chưa VAT)", "Thành tiền (chưa VAT)",
                                      "Thời gian giao hàng"])) if cot_gia else []
    for i, t in enumerate(extra):
        c = NC + 1 + i
        ws.cell(HR, c).value = t
        ws.cell(HR, c)._style = copy.copy(ws.cell(HR, NC)._style)
        ws.column_dimensions[get_column_letter(c)].width = 16
    ref_col = NC + 1 + len(extra)
    ws.cell(HR, ref_col).value = "Ref BOQ"
    ws.cell(HR, ref_col).font = Font(name="Times New Roman", sz=9, i=True, color="FF808080")
    ws.column_dimensions[get_column_letter(ref_col)].hidden = True
    wB = ws.column_dimensions["B"].width or 40
    wC = ws.column_dimensions["C"].width or 30

    # bang dau muc
    r, stt, gi, n_rows = FR, 0, 0, 0
    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV", "XVI",
             "XVII", "XVIII", "XIX", "XX", "XXI", "XXII", "XXIII", "XXIV", "XXV"]
    show_head = len(groups) > 1 or cfg.get("luon_ghi_nhom")
    for g, rows in groups.items():
        if show_head:
            for c in range(1, NC + len(extra) + 1):
                ws.cell(r, c)._style = copy.copy(item_style[min(c, NC) - 1])
                ws.cell(r, c).fill = YELLOW
                ws.cell(r, c).font = Font(name="Times New Roman", sz=11, b=True)
            ws.cell(r, 1).value = roman[gi] if gi < len(roman) else gi + 1
            ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(r, 2).value = g
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[r].height = est_height([(g, wB)], base=18)
            gi += 1; r += 1
        for it in rows:
            stt += 1
            vals = [stt, it["desc"], it["quycach"] or None, None, None, it["unit"] or None,
                    round(it["kl"], 3) if it["kl"] is not None else None, None]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(r, c)
                cell._style = copy.copy(item_style[c - 1])
                cell.value = v
                cell.alignment = Alignment(wrap_text=True, vertical="center",
                                           horizontal="center" if c in (1, 6, 7) else "left")
            ws.cell(r, 7).number_format = "#,##0.###"
            for i in range(len(extra)):
                cell = ws.cell(r, NC + 1 + i)
                cell._style = copy.copy(item_style[6]); cell.number_format = "#,##0"
            if cot_gia and len(extra) >= 3:
                L = get_column_letter
                ws.cell(r, NC + 3).value = (f'=IF(COUNT({L(NC+1)}{r}:{L(NC+2)}{r})=0,"",G{r}*({L(NC+1)}{r}+{L(NC+2)}{r}))')
            ws.cell(r, ref_col).value = "; ".join(it["pos"])
            ws.cell(r, ref_col).font = Font(name="Times New Roman", sz=9, i=True, color="FF808080")
            ws.row_dimensions[r].height = est_height([(it["desc"], wB), (it["quycach"], wC)])
            r += 1; n_rows += 1
    # chan duoi bang giu net mong cua mau
    r_foot = r + 1
    paste(ws, foot, r_foot)
    last = r_foot + len(foot[0]) - 1

    # chu y
    for i, line in enumerate(chu_y):
        c = ws.cell(r_foot + 1 + i, 1)
        c.value = line
        c.font = Font(name="Times New Roman", size=11, italic=True)
        c.alignment = Alignment(wrap_text=False, vertical="center")
        ws.row_dimensions[r_foot + 1 + i].height = 15.75
    # nguoi nhan / hoa don
    nn = cfg.get("nguoi_nhan", {})
    fill_labels(ws, r_foot + idx_hn + 1, r_foot + idx_hd - 1, [
        (r"^Tên người nhận", nn.get(he) or nn.get("_mac_dinh")),
        (r"^Địa chỉ giao hàng", cfg.get("dia_chi_giao_hang")),
        (r"^Ngày yêu cầu báo giá", cfg.get("han_bao_gia"))])
    hd = cfg.get("hoa_don", {})
    fill_labels(ws, r_foot + idx_hd + 1, last, [
        (r"^Công ty$", hd.get("cong_ty")), (r"^Địa chỉ$", hd.get("dia_chi")),
        (r"^Mã số thuế", hd.get("mst")), (r"^Điện thoại", hd.get("dien_thoai"))])

    endcol = get_column_letter(NC + len(extra))
    ws.print_area = f"A1:{endcol}{last}"
    ws.print_title_rows = f"{HR}:{HR}"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if cot_gia:
        ws.page_setup.orientation = "landscape"
    for n in list(wb.defined_names.keys()):
        if re.fullmatch(r"_+a1", n):
            del wb.defined_names[n]
    wb.save(outpath)
    return n_rows


def unique_path(p):
    if not os.path.exists(p):
        return p, 1
    b, e = os.path.splitext(p)
    i = 2
    while os.path.exists(f"{b}_dot{i}{e}"):
        i += 1
    return f"{b}_dot{i}{e}", i


# ----------------------------------------------------------------- lenh
def cmd_read(src, out, ma=None):
    items, excluded, warn, meta = read_source(src, ma)
    json.dump(items, open(out, "w", encoding="utf-8"), ensure_ascii=False)
    json.dump(dict(excluded=excluded, warn=warn, meta=meta), open(out + ".extra.json", "w", encoding="utf-8"), ensure_ascii=False)
    n_src = defaultdict(int)
    for it in items:
        n_src[it["src"]] += 1
    print(json.dumps(dict(du_an=meta["du_an"], ma=meta["ma"], so_dong_doc=len(items), theo_sheet=n_src,
                          khong_dua_vao=len(excluded), canh_bao=len(warn)), ensure_ascii=False, indent=1))
    for w in warn[:15]:
        print("CẢNH BÁO:", w)
    if len(warn) > 15:
        print(f"... còn {len(warn)-15} cảnh báo (nằm trong {out}.extra.json)")


def cmd_plan(items_p, cfg_p):
    items = json.load(open(items_p, encoding="utf-8")); cfg = json.load(open(cfg_p, encoding="utf-8"))
    data = group_all(items, cfg)
    merged = sum(len(r) for g in data.values() for r in g.values())
    print(f"{len(items)} dòng đọc -> {merged} dòng YCBG sau khi gộp trùng\n")
    for he, groups in data.items():
        n = sum(len(r) for r in groups.values())
        print(f"## {he}: {n} dòng, {len(groups)} nhóm")
        for g, rows in groups.items():
            print(f"   - {g}: {len(rows)} dòng | vd: " + " / ".join(x["desc"][:38] for x in rows[:2]))
    kh = [it for it in items if it["_he"] == cfg.get("he_mac_dinh", "Khác")]
    if kh:
        print(f"\nCHƯA XẾP ĐƯỢC HỆ ({len(kh)} dòng) – cần quy tắc hoặc chỉ định `he`:")
        for it in kh[:25]:
            print("   *", (it.get("nhom") or "-"), "|", it["desc"][:70])


def cmd_build(template, items_p, cfg_p, outdir, cot_gia=False):
    items = json.load(open(items_p, encoding="utf-8")); cfg = json.load(open(cfg_p, encoding="utf-8"))
    extra = json.load(open(items_p + ".extra.json", encoding="utf-8")) if os.path.exists(items_p + ".extra.json") else {}
    check_config(cfg)
    if extra.get("meta", {}).get("ma") and cfg["ma_du_an"] != extra["meta"]["ma"]:
        raise SystemExit(f"Mã dự án không khớp: file nguồn {extra['meta']['ma']} ≠ cấu hình {cfg['ma_du_an']}. Dừng, hỏi lại người dùng.")
    os.makedirs(outdir, exist_ok=True)
    ma = cfg["ma_du_an"]
    data = group_all(items, cfg)
    warn = list(extra.get("warn", []))
    sheets, detail, files = [], [], []
    prio_rx = re.compile("|".join(cfg.get("uu_tien", [])), re.I) if cfg.get("uu_tien") else None
    hang_rx = re.compile(r"\b(" + "|".join(re.escape(h) for h in cfg.get("hang_khoa", [])) + r")\b", re.I) if cfg.get("hang_khoa") else None
    for he, groups in data.items():
        base = os.path.join(outdir, f"YCBG_{ma}_{safe_name(he)}.xlsx")
        path, dot = unique_path(base)
        so = f"{ma}-YCBG-{safe_name(he).upper()}-{dot:02d}"
        flat = [it for rows in groups.values() for it in rows]
        n = build_one(template, he, groups, cfg, path, so, cot_gia)
        files.append((he, os.path.basename(path), so, n))
        stt = 0
        for g, rows in groups.items():
            hk = sorted({m.group(1) for it in rows if hang_rx for m in [hang_rx.search(it["desc"])] if m})
            pr = "Cao" if (hk or (prio_rx and any(prio_rx.search(g + " " + it["desc"]) for it in rows))) else "Thường"
            sheets.append([pr, he, os.path.basename(path), so, g, len(rows), sum(it.get("gt") or 0 for it in rows) or None,
                           ("Hãng khóa: " + ", ".join(hk)) if hk else ("Thiết bị chính / hàng đặt dài ngày" if pr == "Cao" else ""),
                           None, None, None, "Chưa gửi", None])
            for it in rows:
                stt += 1
                detail.append([so, stt, he, g, it["desc"], it["quycach"] or None, it["unit"], it["kl"], "; ".join(it["pos"]),
                               ", ".join(str(x) for x in it["ids"]) or None, " + ".join(it["srcs"]), it["tinh_trang"],
                               it["ly_do"] or None, it.get("gia_tk"), it.get("hang_tk"), it.get("link"), None, None, None])
        print(f"[OK] {he}: {n} dòng, {len(groups)} nhóm -> {os.path.basename(path)}")
    for it in items:
        if not clean(it.get("unit")):
            warn.append(f"Thiếu ĐVT trong file nguồn, NCC không chào giá được nếu không có ĐVT: {it['desc'][:70]} ({', '.join(it['pos'][:2])})")
    if any(f[0] == cfg.get("he_mac_dinh", "Khác") for f in files):
        warn.append(f"Có dòng ở hệ '{cfg.get('he_mac_dinh','Khác')}' (chưa xếp được hệ). Duyệt trước khi gửi.")
    tpath, _ = unique_path(os.path.join(outdir, f"YCBG_{ma}_TongHop.xlsx"))
    write_tracking(tpath, sheets, detail, extra.get("excluded", []), warn)
    print(f"[OK] Tổng hợp -> {os.path.basename(tpath)} | {len(detail)} dòng YCBG | không đưa vào: {len(extra.get('excluded', []))} | cảnh báo: {len(warn)}")


def write_tracking(path, sheets, detail, excluded, warn):
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    hf, hfill = Font(bold=True, color="FFFFFFFF"), PatternFill("solid", fgColor="FF1F4E78")

    def table(ws, headers, data, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font = hf; c.fill = hfill; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in data:
            ws.append(row)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active; ws.title = "Theo dõi YCBG"
    table(ws, ["Ưu tiên", "Hệ", "File YCBG", "Số YCBG", "Nhóm vật tư", "Số dòng", "Giá trị tham khảo (VNĐ, nội bộ)",
               "Lý do ưu tiên", "NCC gửi", "Ngày gửi", "Hạn phản hồi", "Trạng thái", "Ghi chú"], sheets,
          [8, 14, 30, 26, 38, 8, 18, 30, 22, 12, 12, 14, 30])
    dv = DataValidation(type="list", formula1='"Chưa gửi,Đã gửi,Đã nhận BG,Không phản hồi,Hủy"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"L2:L{max(ws.max_row, 2)}")
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 7).number_format = "#,##0"
        if ws.cell(r, 1).value == "Cao":
            ws.cell(r, 1).fill = PatternFill("solid", fgColor="FFF8CBAD")
    ws = wb.create_sheet("Chi tiết dòng")
    table(ws, ["Số YCBG", "STT", "Hệ", "Nhóm", "Diễn giải (BOQ CĐT)", "Quy cách", "ĐVT", "KL", "Ref BOQ (sheet!dòng)", "ID nguồn",
               "Sheet nguồn", "Tình trạng giá", "Lý do chưa khớp (nội bộ)", "Giá TK VL (nội bộ)", "Hãng/model TK (nội bộ)",
               "Link (nội bộ)", "Đơn giá NCC", "NCC", "Ngày BG"], detail,
          [24, 6, 14, 24, 50, 28, 9, 10, 40, 10, 22, 24, 40, 14, 18, 18, 14, 18, 12])
    ws = wb.create_sheet("Không đưa vào YCBG")
    table(ws, ["Lý do", "Vị trí", "Mô tả", "ĐVT", "KL", "Ghi chú"],
          [[e["ly_do"], e["pos"], e["desc"], e["unit"], e["kl"], e.get("note")] for e in excluded], [40, 28, 60, 9, 10, 50])
    ws = wb.create_sheet("Cảnh báo")
    table(ws, ["Cảnh báo cần người rà soát"], [[w] for w in warn], [140])
    wb.save(path)


def cmd_check(outdir, items_p, cfg_p):
    items = json.load(open(items_p, encoding="utf-8")); cfg = json.load(open(cfg_p, encoding="utf-8"))
    ok = True
    th = sorted(f for f in os.listdir(outdir) if "TongHop" in f)[-1]
    wb = openpyxl.load_workbook(os.path.join(outdir, th), data_only=True)
    refs_th, kl_th = set(), 0.0
    for row in wb["Chi tiết dòng"].iter_rows(min_row=2, values_only=True):
        refs_th.update(x for x in str(row[8]).split("; ") if x)
        kl_th += float(row[7] or 0)
    refs_in = {p for it in items for p in it["pos"]}
    miss = refs_in - refs_th
    print(f"Ref BOQ vào: {len(refs_in)} | có trong tổng hợp: {len(refs_in & refs_th)} | thiếu: {len(miss)}")
    if miss:
        ok = False; print("  THIẾU:", sorted(miss)[:15])
    src_nounit = {norm(it["desc"]) for it in items if not clean(it.get("unit"))}
    if src_nounit:
        print(f"Lưu ý: {len(src_nounit)} đầu mục nguồn không có ĐVT (đã báo trong sheet Cảnh báo).")
    kl_out, n_out = 0.0, 0
    for f in sorted(os.listdir(outdir)):
        if not f.startswith("YCBG_") or "TongHop" in f or not f.endswith(".xlsx"):
            continue
        s = openpyxl.load_workbook(os.path.join(outdir, f)).worksheets[0]
        HR, CY, HN, HD, FE = locate(s)
        n = kl = 0; bad_unit = []; ph = []
        for r in range(HR + 1, s.max_row + 1):
            a = s.cell(r, 1).value
            if isinstance(a, (int, float)) and s.cell(r, 2).value:
                n += 1; kl += float(s.cell(r, 7).value or 0)
                if not s.cell(r, 6).value and norm(s.cell(r, 2).value) not in src_nounit: bad_unit.append(r)
            if clean(a).lower().startswith("chú ý"):
                cy = r; break
        # chu y khong de len tieu de khoi nguoi nhan
        hn = next(r for r in range(cy, s.max_row + 1) if clean(s.cell(r, 1).value).upper().startswith("THÔNG TIN NGƯỜI NHẬN"))
        overlap = [r for r in range(cy + 1, hn) if clean(s.cell(r, 2).value)]
        for row in s.iter_rows():
            for c in row:
                if isinstance(c.value, str) and re.search(r"<[^<>]+>", c.value):
                    ph.append(c.coordinate)
        logo = len(s._images)
        kl_out += kl; n_out += n
        flag = "" if (logo and not bad_unit and not overlap and not ph) else "  <-- LỖI"
        print(f"{f}: {n} dòng, KL {kl:,.3f}, logo {logo}, thiếu ĐVT {bad_unit[:5]}, chú ý đè {overlap[:3]}, dấu <> {ph[:3]}{flag}")
        if flag: ok = False
    print(f"Tổng dòng ra {n_out} | tổng KL tổng hợp {kl_th:,.3f} / KL các file YCBG {kl_out:,.3f}")
    if abs(kl_th - kl_out) > 0.001 * max(n_out, 1): ok = False
    print("KẾT QUẢ:", "ĐẠT" if ok else "KHÔNG ĐẠT")


if __name__ == "__main__":
    a = sys.argv[1:]
    try:
        if a and a[0] == "read":
            cmd_read(a[1], a[2], a[a.index("--ma") + 1] if "--ma" in a else None)
        elif a and a[0] == "plan":
            cmd_plan(a[1], a[2])
        elif a and a[0] == "build":
            cmd_build(a[1], a[2], a[3], a[4], "--cot-gia" in a)
        elif a and a[0] == "check":
            cmd_check(a[1], a[2], a[3])
        else:
            print(__doc__)
    except SystemExit as e:
        print(e); sys.exit(1)
