#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""boq_xd_tools.py — công cụ hệ Xây dựng (XD) cho BOQ TDI.

Bản v1 16/10/2026 (viết mới theo đặc tả skill covert-boq-cdt-into-boq-tdi mục "Hệ Xây dựng"
và skill boq-gan-ma-gia-thu-vien "Phần XD"). Ghi file bằng XML (lxml), không dùng openpyxl
ghi lại file BOQ, nên phần CĐT (giá trị, font, màu, merge, độ rộng, hình) giữ nguyên.

Lệnh:
  checklib LIB.xlsx                                   Kiểm bố cục thư viện XD (sheet TVGIA, cột B/G/J)
  auto     FILE xdcfg.json [MA_DA]                    Nhận diện sheet XD, hàng tiêu đề, cột, block_start
  build    FILE OUT xdcfg.json LIB.xlsx [--nguon TXT] Dựng khối XD bên phải, nạp TVGIA, thêm Summary XD
  todo     FILE todo.json [SHEET...]                  Liệt kê dòng công việc của khối XD + trạng thái
  recalc   FILE rep.json [SHEET...]                   Tính lại bằng LibreOffice (bản sao), đếm ô lỗi, tổng
  cmp      ORIG OUT [SHEET...]                        So ô CĐT giữa file gốc và file ra (giá trị, đậm, nền, viền)
  tv       FILE tv.json                               Danh mục TVGIA (mục, cờ trùng tên, dòng con)
  apply    FILE BANG_KHOP_XD.xlsx todo.json A.json MA_DA   Áp Bảng khớp XD cho dòng chưa gán
  write    FILE OUT ASSIGN.json                       Ghi Công tác + đơn giá VL/NC (một lần)
  export   FILE BANG_KHOP_XD.xlsx MA_DA [SHEET...]    Lưu kết quả đã duyệt vào Bảng khớp XD

ASSIGN.json: list (hoặc {"assign":[...]}) các mục
  {"sheet","row","tv","plus_vl","tv_nc","coef_vl","coef_nc","vl","nc","web_vl","web_src","conf","note"}
  tv / plus_vl / tv_nc = số dòng TVGIA của bản đang dùng; vl/nc=false: chỉ NC / chỉ VL;
  web_vl = giá tham khảo web (số cứng, tô E2EFDA, Nguồn giá = "Giá tham khảo", không gán Công tác).
"""
import copy
import datetime as _dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
from pathlib import Path

from lxml import etree
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

VERSION = "v1 16/10/2026"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NSR = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSP = "http://schemas.openxmlformats.org/package/2006/relationships"
NSCT = "http://schemas.openxmlformats.org/package/2006/content-types"
WS_TYPE = NSR + "/worksheet"
WS_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def q(tag, ns=NS):
    return "{%s}%s" % (ns, tag)


# ----------------------------------------------------------------------------
# Chuẩn hóa chữ
# ----------------------------------------------------------------------------
_CJK = re.compile(r"[\u2E80-\u9FFF\uF900-\uFAFF\uFF00-\uFFEF\u3000-\u303F]+")


def detone(s):
    """Bỏ dấu tiếng Việt, chữ thường, gộp khoảng trắng."""
    if s is None:
        return ""
    s = str(s).replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return " ".join(s.lower().split())


def norm_key(s):
    """Khóa so khớp: bỏ chữ Trung/Hán, bỏ dấu, bỏ ký tự đặc biệt."""
    s = _CJK.sub(" ", str(s or ""))
    s = detone(s)
    s = s.replace("m³", "m3").replace("m²", "m2")
    s = re.sub(r"[^0-9a-z.,x/%<>=+-]+", " ", s)
    return " ".join(s.split())


def norm_unit(u):
    u = norm_key(u)
    u = u.replace("m 3", "m3").replace("m 2", "m2")
    al = {"md": "m", "met": "m", "m dai": "m", "lit": "lit", "lít": "lit", "cai": "cai", "chiec": "cai",
          "tan": "tan", "ton": "tan", "tron goi": "goi", "goi": "goi", "bo": "bo", "vi tri": "vi tri"}
    first = u.split(" ")[0] if u else ""
    for k, v in al.items():
        if u == k or u.startswith(k + " "):
            return v
    return {"m3": "m3", "m2": "m2", "kg": "kg", "m": "m"}.get(first, u)


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def to_num(v):
    if is_num(v):
        return float(v)
    if isinstance(v, str):
        t = v.strip().replace(",", "")
        if re.fullmatch(r"-?\d+(\.\d+)?", t):
            return float(t)
    return None


ROMAN = re.compile(r"^(?=[IVXLC]+$)M{0,3}(CM|CD|D?C{0,3})(XC|XL|L?X{0,3})(IX|IV|V?I{0,3})$")


def stt_class(stt):
    """Lớp STT dùng để nhận dòng tiêu đề: roman / letter / star / dotted / num / ''."""
    s = str(stt or "").strip().rstrip(".")
    if not s:
        return ""
    if s in ("*", "-", "+", "•"):
        return "star"
    if ROMAN.match(s.upper()) and s.isupper():
        return "roman"
    if re.fullmatch(r"[A-Z]", s):
        return "letter"
    if re.fullmatch(r"[A-Za-z0-9]+(\.[A-Za-z0-9]+)+", s):
        return "dotted"
    if re.fullmatch(r"\d+[a-z]?", s):
        return "num"
    return "other"


def seq_class(stt, prev_letter):
    """Như stt_class nhưng I, V, X, L, C, D, M đứng ngay sau chữ cái liền trước (H→I, K→L…) là chữ cái, không phải số La Mã."""
    cls = stt_class(stt)
    s = str(stt or "").strip().rstrip(".")
    if cls == "roman" and len(s) == 1 and prev_letter and ord(s) == ord(prev_letter) + 1:
        return "letter"
    return cls


EXCL = re.compile(r"khong bao gia|khong chao gia|da tinh o|da tinh tai|da bao gia o|cdt cung cap|chu dau tu cung cap|"
                  r"khong thuoc pham vi|nha thau khac|do cdt|tinh o hang muc")

# ----------------------------------------------------------------------------
# Bố cục khối XD (c0 = block_start)
# ----------------------------------------------------------------------------
NCOL = 30
C_STT, C_DESC, C_CL, C_UNIT = 0, 1, 2, 3
C_SL, C_DAI, C_RONG, C_CAO, C_HS = 4, 5, 6, 7, 8
C_KLTDI, C_KL, C_DGVL, C_DGNC, C_DGT, C_TT, C_NOTE, C_NHAN = 9, 10, 11, 12, 13, 14, 15, 16
C_VCT, C_VDV, C_VKL, C_VDG, C_VTT = 17, 18, 19, 20, 21
C_NCT, C_NDV, C_NKL, C_NDG, C_NTT = 22, 23, 24, 25, 26
C_SRC, C_CONF, C_GNOTE = 27, 28, 29
WIDTHS = [6, 45, 18, 8, 7, 7, 7, 7, 7, 10, 11, 12, 12, 13, 15, 20, 9, 38, 7, 10, 12, 14, 38, 7, 10, 12, 14, 14, 11, 32]
MARK = "NỘI DUNG (TDI)"
SRC_LIB, SRC_WEB, SRC_NONE = "Giá thư viện", "Giá tham khảo", "Chưa có giá"


# ----------------------------------------------------------------------------
# Gói xlsx dạng XML
# ----------------------------------------------------------------------------
class Package:
    def __init__(self, path):
        self.path = path
        with zipfile.ZipFile(path) as z:
            self.order = z.namelist()
            self.parts = {n: z.read(n) for n in self.order}
        self.wb = etree.fromstring(self.parts["xl/workbook.xml"])
        self.rels = etree.fromstring(self.parts["xl/_rels/workbook.xml.rels"])
        self.ct = etree.fromstring(self.parts["[Content_Types].xml"])

    def sheets(self):
        rid2t = {r.get("Id"): r.get("Target") for r in self.rels}
        out = []
        for s in self.wb.find(q("sheets")):
            t = rid2t[s.get(q("id", NSR))]
            t = t.lstrip("/")
            if not t.startswith("xl/"):
                t = "xl/" + t
            out.append((s.get("name"), t, s))
        return out

    def sheet_part(self, name):
        for n, t, _ in self.sheets():
            if n == name:
                return t
        raise KeyError(f"Không có sheet '{name}'")

    def add_sheet(self, name, xml_bytes):
        k = 1
        while f"xl/worksheets/sheet{k}.xml" in self.parts:
            k += 1
        part = f"xl/worksheets/sheet{k}.xml"
        self.parts[part] = xml_bytes
        self.order.append(part)
        ids = [int(re.sub(r"\D", "", r.get("Id")) or 0) for r in self.rels]
        rid = f"rId{max(ids + [0]) + 1}"
        while any(r.get("Id") == rid for r in self.rels):
            rid = rid + "x"
        rel = etree.SubElement(self.rels, q("Relationship", NSP))
        rel.set("Id", rid)
        rel.set("Type", WS_TYPE)
        rel.set("Target", f"worksheets/sheet{k}.xml")
        sh = self.wb.find(q("sheets"))
        el = etree.SubElement(sh, q("sheet"))
        el.set("name", name)
        el.set("sheetId", str(max([int(s.get("sheetId")) for s in sh if s is not el] + [0]) + 1))
        el.set(q("id", NSR), rid)
        ov = etree.SubElement(self.ct, q("Override", NSCT))
        ov.set("PartName", "/" + part)
        ov.set("ContentType", WS_CT)
        return part

    def drop_calcchain(self):
        for n in list(self.parts):
            if n.endswith("calcChain.xml"):
                del self.parts[n]
                self.order.remove(n)
        for r in list(self.rels):
            if r.get("Type", "").endswith("/calcChain"):
                self.rels.remove(r)
        for o in list(self.ct):
            if o.get("PartName", "").endswith("calcChain.xml"):
                self.ct.remove(o)
        cp = self.wb.find(q("calcPr"))
        if cp is None:
            cp = etree.Element(q("calcPr"))
            # calcPr đứng sau definedNames / sheets
            anchor = self.wb.find(q("definedNames"))
            if anchor is None:
                anchor = self.wb.find(q("sheets"))
            anchor.addnext(cp)
        cp.set("fullCalcOnLoad", "1")

    def add_defined_name(self, name, value, local_id=None, hidden=False):
        dns = self.wb.find(q("definedNames"))
        if dns is None:
            dns = etree.Element(q("definedNames"))
            self.wb.find(q("sheets")).addnext(dns)
        for d in dns:
            if d.get("name") == name and (d.get("localSheetId") == (None if local_id is None else str(local_id))):
                d.text = value
                return
        d = etree.SubElement(dns, q("definedName"))
        d.set("name", name)
        if local_id is not None:
            d.set("localSheetId", str(local_id))
        if hidden:
            d.set("hidden", "1")
        d.text = value

    def save(self, out):
        self.parts["xl/workbook.xml"] = etree.tostring(self.wb, xml_declaration=True, encoding="UTF-8", standalone=True)
        self.parts["xl/_rels/workbook.xml.rels"] = etree.tostring(self.rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        self.parts["[Content_Types].xml"] = etree.tostring(self.ct, xml_declaration=True, encoding="UTF-8", standalone=True)
        tmp = str(out) + ".tmp"
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
            for n in self.order:
                if n in self.parts:
                    z.writestr(n, self.parts[n])
        os.replace(tmp, out)


# thứ tự phần tử con của <worksheet> theo schema
WS_ORDER = ["sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData", "sheetCalcPr",
            "sheetProtection", "protectedRanges", "scenarios", "autoFilter", "sortState", "dataConsolidate",
            "customSheetViews", "mergeCells", "phoneticPr", "conditionalFormatting", "dataValidations",
            "hyperlinks", "printOptions", "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks",
            "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing", "legacyDrawing",
            "legacyDrawingHF", "picture", "oleObjects", "controls", "webPublishItems", "tableParts", "extLst"]


def _local(el):
    return etree.QName(el).localname


def ensure_child(root, name):
    el = root.find(q(name))
    if el is not None:
        return el
    el = etree.Element(q(name))
    pos = WS_ORDER.index(name)
    for i, ch in enumerate(root):
        ln = _local(ch)
        if ln in WS_ORDER and WS_ORDER.index(ln) > pos:
            ch.addprevious(el)
            return el
    root.append(el)
    return el


def split_ref(ref):
    m = re.match(r"\$?([A-Z]+)\$?(\d+)", ref)
    return CI(m.group(1)), int(m.group(2))


class Sheet:
    """Ghi ô vào sheet XML, giữ nguyên ô khác."""

    def __init__(self, xml):
        self.root = etree.fromstring(xml)
        self.sd = self.root.find(q("sheetData"))
        self.rows = {}
        last = 0
        for r in self.sd.findall(q("row")):
            if r.get("r") is None:
                r.set("r", str(last + 1))
            last = int(r.get("r"))
            self.rows[last] = r
            ci = 0
            for c in r.findall(q("c")):
                if c.get("r") is None:
                    c.set("r", f"{L(ci + 1)}{last}")
                ci = split_ref(c.get("r"))[0]
        self.touched = set()
        self.maxc = 0
        self.maxr = 0

    def row(self, r):
        el = self.rows.get(r)
        if el is not None:
            return el
        el = etree.Element(q("row"))
        el.set("r", str(r))
        after = [k for k in self.rows if k > r]
        if after:
            self.rows[min(after)].addprevious(el)
        else:
            self.sd.append(el)
        self.rows[r] = el
        return el

    def cell(self, r, c, create=True):
        row = self.row(r) if create else self.rows.get(r)
        if row is None:
            return None
        ref = f"{L(c)}{r}"
        nxt = None
        for ce in row.findall(q("c")):
            cc = split_ref(ce.get("r"))[0]
            if cc == c:
                return ce
            if cc > c and nxt is None:
                nxt = ce
        if not create:
            return None
        ce = etree.Element(q("c"))
        ce.set("r", ref)
        if nxt is not None:
            nxt.addprevious(ce)
        else:
            row.append(ce)
        if "spans" in row.attrib:
            del row.attrib["spans"]
        return ce

    def set(self, r, c, value=None, f=None, s=None, cache=None):
        ce = self.cell(r, c)
        keep_s = ce.get("s")
        for k in list(ce.attrib):
            if k != "r":
                del ce.attrib[k]
        for ch in list(ce):
            ce.remove(ch)
        if s is not None:
            ce.set("s", str(s))
        elif keep_s is not None:
            ce.set("s", keep_s)
        if f is not None:
            fe = etree.SubElement(ce, q("f"))
            fe.text = f[1:] if f.startswith("=") else f
            cv = cache
            if isinstance(cv, bool):
                cv = int(cv)
            if is_num(cv):
                etree.SubElement(ce, q("v")).text = repr(float(cv)) if isinstance(cv, float) else str(cv)
            elif cv not in (None, ""):
                ce.set("t", "str")
                etree.SubElement(ce, q("v")).text = str(cv)
            elif cv == "":
                ce.set("t", "str")
                etree.SubElement(ce, q("v")).text = ""
        elif value is not None and value != "":
            if is_num(value):
                etree.SubElement(ce, q("v")).text = repr(float(value)) if isinstance(value, float) else str(value)
            else:
                ce.set("t", "inlineStr")
                is_ = etree.SubElement(ce, q("is"))
                t = etree.SubElement(is_, q("t"))
                t.text = str(value)
                t.set(XML_SPACE, "preserve")
        self.maxc = max(self.maxc, c)
        self.maxr = max(self.maxr, r)
        return ce

    def style_of(self, r, c):
        ce = self.cell(r, c, create=False)
        return None if ce is None else ce.get("s")

    def formula_of(self, r, c):
        ce = self.cell(r, c, create=False)
        if ce is None:
            return None
        fe = ce.find(q("f"))
        return None if fe is None else ("=" + (fe.text or ""))

    def value_of(self, r, c):
        ce = self.cell(r, c, create=False)
        if ce is None:
            return None
        t = ce.get("t")
        if t == "inlineStr":
            return "".join(ce.find(q("is")).itertext())
        v = ce.find(q("v"))
        if v is None:
            return None
        if t in ("str", "s", "e"):
            return v.text  # 's' cần sharedStrings: xử lý ở reader
        try:
            return float(v.text)
        except (TypeError, ValueError):
            return v.text

    def merge(self, ref):
        mc = ensure_child(self.root, "mergeCells")
        for m in mc:
            if m.get("ref") == ref:
                return
        etree.SubElement(mc, q("mergeCell")).set("ref", ref)
        mc.set("count", str(len(mc)))

    def merges(self):
        mc = self.root.find(q("mergeCells"))
        return [] if mc is None else [m.get("ref") for m in mc]

    def set_widths(self, widths):
        """widths: {col_index: width}. Tách các <col> đang phủ cột cần đặt."""
        cols = ensure_child(self.root, "cols")
        items = [(int(c.get("min")), int(c.get("max")), dict(c.attrib)) for c in cols]
        targets = sorted(widths)
        out = []
        for mn, mx, at in items:
            cur = mn
            for t in [x for x in targets if mn <= x <= mx]:
                if cur < t:
                    out.append((cur, t - 1, at))
                cur = t + 1
            if cur <= mx:
                out.append((cur, mx, at))
        base = {}
        for mn, mx, at in items:
            for t in targets:
                if mn <= t <= mx:
                    base[t] = at
        for t in targets:
            at = dict(base.get(t, {}))
            at.pop("hidden", None)
            at["width"] = str(widths[t])
            at["customWidth"] = "1"
            out.append((t, t, at))
        out.sort(key=lambda x: x[0])
        for c in list(cols):
            cols.remove(c)
        for mn, mx, at in out:
            el = etree.SubElement(cols, q("col"))
            at = dict(at)
            at["min"], at["max"] = str(mn), str(mx)
            for k in ("min", "max", "width", "style", "hidden", "bestFit", "customWidth", "phonetic", "outlineLevel", "collapsed"):
                if k in at:
                    el.set(k, at[k])
            for k, v in at.items():
                if k not in el.attrib:
                    el.set(k, v)

    def fix_dimension(self):
        maxr = max(self.rows) if self.rows else 1
        maxc = 1
        minr = min(self.rows) if self.rows else 1
        for r in self.rows.values():
            for ce in r.findall(q("c")):
                maxc = max(maxc, split_ref(ce.get("r"))[0])
        d = ensure_child(self.root, "dimension")
        d.set("ref", f"A{minr}:{L(maxc)}{maxr}")

    def tobytes(self):
        self.fix_dimension()
        return etree.tostring(self.root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ----------------------------------------------------------------------------
# styles.xml: thêm định dạng cho khối XD
# ----------------------------------------------------------------------------
class Styles:
    SPEC = {
        # key: (bold, color, fill, numFmtId, halign, wrap, sz_delta)
        "hdr": (True, None, "DDEBF7", 0, "center", True, 0),
        "hdr_coef": (True, "C00000", "FFF2CC", 4, "center", False, 0),
        "lbl_coef": (True, None, "FFF2CC", 0, "center", False, 0),
        "txt": (False, None, None, 0, "left", True, 0),
        "txt_c": (False, None, None, 0, "center", True, 0),
        "num0": (False, None, None, 3, "right", False, 0),
        "num2": (False, None, None, 4, "right", False, 0),
        "in_txt": (False, "0070C0", None, 0, "left", True, 0),
        "in_num": (False, "0070C0", None, 4, "right", False, 0),
        "in_num0": (False, "0070C0", None, 3, "right", False, 0),
        "ttl_txt": (True, None, "FFFF00", 0, "left", True, 0),
        "ttl_c": (True, None, "FFFF00", 0, "center", True, 0),
        "ttl_num": (True, None, "FFFF00", 3, "right", False, 0),
        "web_num": (False, None, "E2EFDA", 3, "right", False, 0),
        "tot_txt": (True, None, "FCE4D6", 0, "left", False, 0),
        "tot_num": (True, None, "FCE4D6", 3, "right", False, 0),
        "big": (True, None, None, 0, "left", False, 4),
        "bold": (True, None, None, 0, "left", False, 0),
    }

    def __init__(self, pkg):
        self.pkg = pkg
        self.root = etree.fromstring(pkg.parts["xl/styles.xml"])
        self.ids = {}

    def _list(self, name, child):
        el = self.root.find(q(name))
        if el is None:
            el = etree.SubElement(self.root, q(name))
        return el

    def get(self, key):
        if key in self.ids:
            return self.ids[key]
        bold, color, fill, nf, ha, wrap, dsz = self.SPEC[key]
        fonts = self.root.find(q("fonts"))
        f0 = copy.deepcopy(fonts[0])
        for ch in list(f0):
            if _local(ch) in ("b", "color"):
                f0.remove(ch)
        if bold:
            f0.insert(0, etree.Element(q("b")))
        if color:
            c = etree.Element(q("color"))
            c.set("rgb", "FF" + color)
            sz = f0.find(q("sz"))
            (sz.addnext(c) if sz is not None else f0.insert(1 if bold else 0, c))
        if dsz:
            sz = f0.find(q("sz"))
            if sz is not None:
                sz.set("val", str(float(sz.get("val", "11")) + dsz).rstrip("0").rstrip("."))
        fonts.append(f0)
        fonts.set("count", str(len(fonts)))
        font_id = len(fonts) - 1
        fill_id = 0
        if fill:
            fills = self.root.find(q("fills"))
            fe = etree.SubElement(fills, q("fill"))
            pf = etree.SubElement(fe, q("patternFill"))
            pf.set("patternType", "solid")
            etree.SubElement(pf, q("fgColor")).set("rgb", "FF" + fill)
            etree.SubElement(pf, q("bgColor")).set("indexed", "64")
            fills.set("count", str(len(fills)))
            fill_id = len(fills) - 1
        borders = self.root.find(q("borders"))
        be = etree.SubElement(borders, q("border"))
        for side in ("left", "right", "top", "bottom"):
            s = etree.SubElement(be, q(side))
            s.set("style", "thin")
            etree.SubElement(s, q("color")).set("indexed", "64")
        etree.SubElement(be, q("diagonal"))
        borders.set("count", str(len(borders)))
        border_id = len(borders) - 1
        xfs = self.root.find(q("cellXfs"))
        xf = etree.SubElement(xfs, q("xf"))
        for k, v in (("numFmtId", nf), ("fontId", font_id), ("fillId", fill_id), ("borderId", border_id), ("xfId", 0)):
            xf.set(k, str(v))
        xf.set("applyFont", "1")
        xf.set("applyBorder", "1")
        xf.set("applyAlignment", "1")
        if nf:
            xf.set("applyNumberFormat", "1")
        if fill:
            xf.set("applyFill", "1")
        al = etree.SubElement(xf, q("alignment"))
        al.set("horizontal", ha)
        al.set("vertical", "center")
        if wrap:
            al.set("wrapText", "1")
        xfs.set("count", str(len(xfs)))
        self.ids[key] = len(xfs) - 1
        return self.ids[key]

    def save(self):
        self.pkg.parts["xl/styles.xml"] = etree.tostring(self.root, xml_declaration=True, encoding="UTF-8", standalone=True)


def qs(name):
    """Tên sheet trong công thức."""
    return "'" + name.replace("'", "''") + "'"


# ----------------------------------------------------------------------------
# auto
# ----------------------------------------------------------------------------
HK = {
    "stt": re.compile(r"^(stt|tt|so tt|so thu tu|item|no\.?)\b"),
    "desc": re.compile(r"noi dung|ten cong viec|hang muc|mo ta|dien giai cong viec|description|cong viec"),
    "spec": re.compile(r"chung loai|quy cach|thong so|dac tinh|specification|\bspec\b"),
    "unit": re.compile(r"^(don vi|dvt|dv|unit)\b|don vi tinh"),
    "kl": re.compile(r"khoi luong|so luong|\bkl\b|\bqty\b|quantity"),
    "note": re.compile(r"ghi chu|remark|\bnote"),
    "price": re.compile(r"don gia|thanh tien|\bgia\b|amount|price|unit rate|vat lieu|nhan cong"),
}
SKIP_SHEET = re.compile(r"summary|tong hop|bia|chi dan|huong dan|cover|^tvgia$|thu vien|ty gia|thong tin")
XD_WORDS = re.compile(r"be tong|\bcoc\b|van khuon|cot thep|\bxay\b|\btrat\b|\bson\b|op lat|\blat\b|dao dat|lap dat dam|"
                      r"\bmong\b|\btuong\b|ket cau thep|\bton\b|chong tham|\bnen\b|\bduong\b|\bcong btct|\branh\b|ho ga|"
                      r"\bbase\b|\bcat\b|da dam|xa go|\bkeo\b|\bmai\b|tran thach cao|lan can|cau thang|\bgieng\b|bo via|"
                      r"\bbe\b|giao|hoan thien|nilon|\bbtct\b|thep hop|\bgach\b|\bvua\b|dam chat|lu len|cua di|cua so|cua cuon")
MEP_WORDS = re.compile(r"cap dien|tu dien|bong den|den led|den chieu sang|den downlight|den exit|o cam|cong tac den|"
                       r"ong nuoc cap|may bom|dieu hoa|hvac|pccc|bao chay|sprinkler|camera|\belv\b|ong gio|\bquat\b|"
                       r"\bmcb\b|\bmccb\b|\bacb\b|aptomat|thang mang cap|mang cap|day cap|cap dong|cap quang|ong luon day|"
                       r"\bppr\b|chua chay|dau bao|\bloa\b|tu bom|van chan|dong ho nuoc|fcu|ahu|chiller")


def _cell_text(ws, r, c, mmap):
    key = mmap.get((r, c))
    v = ws.cell(*key).value if key else ws.cell(r, c).value
    return detone(_CJK.sub(" ", str(v))) if v not in (None, "") else ""


def _merge_map(ws):
    mm = {}
    for mr in ws.merged_cells.ranges:
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                mm[(r, c)] = (mr.min_row, mr.min_col)
    return mm


def last_used_col(ws, maxr):
    last = 0
    for row in ws.iter_rows(min_row=1, max_row=maxr):
        for ce in row:
            if ce.value not in (None, "") and ce.column > last:
                last = ce.column
    for mr in ws.merged_cells.ranges:
        last = max(last, mr.max_col)
    return last


def auto_sheet(ws):
    mm = _merge_map(ws)
    maxc = min(ws.max_column, 60)
    best = (0, None)
    for r in range(1, min(ws.max_row, 30) + 1):
        hits = set()
        for c in range(1, maxc + 1):
            t = _cell_text(ws, r, c, mm)
            if not t or len(t) > 60:
                continue
            for k, rx in HK.items():
                if rx.search(t):
                    hits.add(k)
        score = len(hits & {"stt", "desc", "unit", "kl"}) * 2 + len(hits)
        if score > best[0]:
            best = (score, r)
    if best[0] < 6:
        return None, "không thấy hàng tiêu đề (STT / Nội dung / Đơn vị / Khối lượng)"
    h1 = best[1]
    h2 = h1
    vertical = any(mr.min_row == h1 and mr.max_row > h1 for mr in ws.merged_cells.ranges)
    below = [_cell_text(ws, h1 + 1, c, mm) for c in range(1, maxc + 1)]
    nums_below = any(is_num(ws.cell(h1 + 1, c).value) for c in range(1, maxc + 1))
    hspan = any(mr.min_row == h1 and mr.max_col > mr.min_col and
                any(ws.cell(h1 + 1, c).value not in (None, "") for c in range(mr.min_col, mr.max_col + 1))
                for mr in ws.merged_cells.ranges)
    texts_below = sum(1 for t in below if t and not re.fullmatch(r"[\d\s.,()-]+", t))
    if vertical or hspan or (not nums_below and (sum(1 for t in below if t and HK["price"].search(t)) >= 1 or texts_below >= 2)):
        h2 = h1 + 1
    ds = h2 + 1
    vals = [ws.cell(ds, c).value for c in range(1, maxc + 1)]
    if sum(1 for v in vals if is_num(v)) >= 4 and all(is_num(v) or v in (None, "") for v in vals):
        ds += 1  # hàng đánh số cột
    while ds <= ws.max_row and all(ws.cell(ds, c).value in (None, "") for c in range(1, maxc + 1)):
        ds += 1
    cols, kls = {}, []
    for c in range(1, maxc + 1):
        top = _cell_text(ws, h1, c, mm)
        low = _cell_text(ws, h2, c, mm) if h2 != h1 else ""
        both = (top + " | " + low)
        if HK["kl"].search(both) and not HK["price"].search(low or top):
            kls.append((c, both))
            continue
        if HK["price"].search(both):
            continue
        for k in ("stt", "desc", "spec", "unit", "note"):
            if k not in cols and (HK[k].search(top) or (low and HK[k].search(low))):
                cols[k] = c
                break
    warn = []
    kl = None
    if len(kls) == 1:
        kl = kls[0][0]
    elif len(kls) > 1:
        tot = [c for c, t in kls if re.search(r"tong|total|cong\b|sum", t)]
        if tot:
            kl = tot[-1]
        else:
            warn.append("nhiều cột khối lượng, không có cột tổng: " + ", ".join(L(c) for c, _ in kls) + " -> hỏi người dùng")
    if "desc" not in cols:
        return None, "không thấy cột Nội dung"
    if "unit" not in cols:
        return None, "không thấy cột Đơn vị"
    # dòng cuối dữ liệu: bỏ dòng trống cuối và dòng "Tổng cộng" cuối bảng
    dc = cols["desc"]
    end = ws.max_row
    while end > ds and all(ws.cell(end, c).value in (None, "") for c in range(1, maxc + 1)):
        end -= 1
    tot_rx = re.compile(r"^(tong cong|tong gia tri|cong\s*(:|$)|total\b|grand total)")
    for r in range(end, max(ds, end - 15), -1):
        if tot_rx.match(_cell_text(ws, r, dc, mm)):
            end = r - 1
            while end > ds and all(ws.cell(end, c).value in (None, "") for c in range(1, maxc + 1)):
                end -= 1
            break
    last = last_used_col(ws, ws.max_row)
    hidden_right = [c for c, d in ws.column_dimensions.items() if d.hidden and CI(c) > last]
    g = {"header_rows": [h1, h2], "data_start_row": ds, "data_end_row": end,
         "stt": L(cols["stt"]) if "stt" in cols else None, "desc": L(dc), "unit": L(cols["unit"]),
         "kl": L(kl) if kl else None, "note": L(cols["note"]) if "note" in cols else None,
         "spec": L(cols["spec"]) if "spec" in cols else None,
         "block_start": L(last + 2), "warn": warn}
    if hidden_right:
        g["warn"].append("có cột ẩn bên phải: " + ", ".join(hidden_right) + " -> cân nhắc dời block_start")
    if h2 - 2 < 1:
        g["warn"].append("tiêu đề sát đầu sheet: khối XD cần 2 hàng phía trên hàng tiêu đề cuối -> báo người dùng")
    return g, None


def classify_sheet(ws, g):
    dc = CI(g["desc"])
    xd = mep = 0
    for r in range(g["data_start_row"], g["data_end_row"] + 1):
        t = detone(_CJK.sub(" ", str(ws.cell(r, dc).value or "")))
        if not t:
            continue
        if XD_WORDS.search(t):
            xd += 1
        if MEP_WORDS.search(t):
            mep += 1
    if xd >= 5 and xd >= 1.5 * mep:
        return "XD", xd, mep
    if mep >= 5 and mep >= 1.5 * xd:
        return "MEP", xd, mep
    return "?", xd, mep


def has_block(ws, upto_row):
    for row in ws.iter_rows(min_row=1, max_row=upto_row):
        for ce in row:
            if isinstance(ce.value, str) and ce.value.strip().upper() == MARK:
                return ce.column
    return None


def cmd_auto(args):
    f, out = args[0], args[1]
    ma = args[2] if len(args) > 2 else ""
    wb = openpyxl.load_workbook(f, data_only=True)
    sheets, other = [], []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            other.append({"sheet": ws.title, "ly_do": "sheet ẩn"})
            continue
        if SKIP_SHEET.search(detone(_CJK.sub(" ", ws.title))):
            other.append({"sheet": ws.title, "ly_do": "tên giống bìa / tổng hợp / thư viện"})
            print(f"- bỏ qua '{ws.title}' (tên giống bìa / tổng hợp / thư viện)")
            continue
        g, err = auto_sheet(ws)
        if not g:
            other.append({"sheet": ws.title, "ly_do": err})
            print(f"- bỏ qua '{ws.title}' ({err})")
            continue
        hb = has_block(ws, g["header_rows"][1])
        if hb:
            other.append({"sheet": ws.title, "ly_do": f"đã có khối XD từ cột {L(hb - 1)}"})
            print(f"- bỏ qua '{ws.title}' (đã có khối XD từ cột {L(hb - 1)})")
            continue
        kind, xd, mep = classify_sheet(ws, g)
        g = {"sheet": ws.title, "he": kind, "diem_xd": xd, "diem_mep": mep, **g}
        n = g["data_end_row"] - g["data_start_row"] + 1
        msg = (f"'{ws.title}': hệ {kind} (XD {xd} / MEP {mep}), tiêu đề {g['header_rows']}, dữ liệu {g['data_start_row']}–"
               f"{g['data_end_row']} ({n} dòng), STT {g['stt']} Nội dung {g['desc']} ĐVT {g['unit']} KL {g['kl']} "
               f"Ghi chú {g['note']}, khối XD từ {g['block_start']}")
        if kind == "XD":
            sheets.append(g)
            print("+ " + msg)
        else:
            other.append({"sheet": ws.title, "ly_do": f"hệ {kind} (XD {xd} / MEP {mep})", "cfg": g})
            print(("? " if kind == "?" else "- ") + msg + ("  -> chưa chắc hệ, hỏi người dùng" if kind == "?" else "  -> sheet MEP, bỏ"))
        for w in g["warn"]:
            print("   ! " + w)
    cfg = {"tool": "boq_xd_tools " + VERSION, "input": str(Path(f).resolve()), "ma_da": ma, "sheets": sheets,
           "khong_dung": other}
    Path(out).write_text(json.dumps(cfg, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Đã ghi {out}: {len(sheets)} sheet XD. Sửa file nếu nhận sai; chuyển sheet từ 'khong_dung' sang 'sheets' nếu cần.")


# ----------------------------------------------------------------------------
# Thư viện XD (TVGIA)
# ----------------------------------------------------------------------------
def check_lib(path):
    """Kiểm bố cục file thư viện XD: sheet TVGIA, B = Tên công việc, G = Thành tiền VL, J = Thành tiền NC."""
    res = {"file": str(path), "ok": False, "loi": []}
    try:
        res["sua_ngay"] = _dt.datetime.fromtimestamp(os.path.getmtime(path)).strftime("%d/%m/%Y %H:%M")
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # noqa
        res["loi"].append(f"không mở được file: {e}")
        return res
    res["sheets"] = wb.sheetnames
    if "TVGIA" not in wb.sheetnames:
        res["loi"].append("không có sheet tên đúng 'TVGIA' (sheet đang có: " + ", ".join(wb.sheetnames) + ")")
        return res
    ws = wb["TVGIA"]
    h1 = [detone(ws.cell(1, c).value) for c in range(1, 12)]
    h2 = [detone(ws.cell(2, c).value) for c in range(1, 12)]
    res["tieu_de"] = {L(c + 1): (h1[c] + " / " + h2[c]).strip(" /") for c in range(11)}
    if "ten cong viec" not in h1[1]:
        res["loi"].append(f"cột B hàng 1 không phải 'TÊN CÔNG VIỆC' (đang là '{ws['B1'].value}')")
    if "vat tu" not in " ".join(h1[3:7]):
        res["loi"].append("hàng 1 cột D–G không có nhóm 'VẬT TƯ'")
    if "nhan cong" not in " ".join(h1[7:10]):
        res["loi"].append("hàng 1 cột H–J không có nhóm 'NHÂN CÔNG'")
    if "thanh tien" not in h2[6]:
        res["loi"].append(f"cột G hàng 2 không phải 'THÀNH TIỀN' VL (đang là '{ws['G2'].value}')")
    if "thanh tien" not in h2[9]:
        res["loi"].append(f"cột J hàng 2 không phải 'THÀNH TIỀN' NC (đang là '{ws['J2'].value}')")
    n = sg = sj = 0
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 2).value not in (None, ""):
            n += 1
            sg += to_num(ws.cell(r, 7).value) or 0
            sj += to_num(ws.cell(r, 10).value) or 0
    res.update({"so_dong": ws.max_row, "dong_co_ten": n, "sum_G": round(sg), "sum_J": round(sj)})
    if n < 10:
        res["loi"].append(f"chỉ có {n} dòng có tên ở cột B")
    if sg == 0 and sj == 0:
        res["loi"].append("ΣG = ΣJ = 0: file chưa có giá trị đã tính (mở bằng Excel, lưu lại rồi stage lại)")
    res["ok"] = not res["loi"]
    return res


def cmd_checklib(args):
    r = check_lib(args[0])
    print(json.dumps(r, ensure_ascii=False, indent=1))
    print("KẾT QUẢ:", "ĐẠT" if r["ok"] else "KHÔNG ĐẠT -> dừng, không chạy build")
    return 0 if r["ok"] else 1


def tvgia_xml(lib_path, note):
    """Sheet TVGIA mới: cột A–K, giữ công thức nội bộ sheet và giá trị đã tính, bỏ định dạng, ghi chú ở L1."""
    wf = openpyxl.load_workbook(lib_path)["TVGIA"]
    wv = openpyxl.load_workbook(lib_path, data_only=True)["TVGIA"]
    root = etree.Element(q("worksheet"), nsmap={None: NS, "r": NSR})
    sh = Sheet(etree.tostring(etree.Element(q("worksheet"), nsmap={None: NS})).replace(b"/>", b"><sheetData/></worksheet>"))
    nf = 0
    for r in range(1, wf.max_row + 1):
        for c in range(1, 12):
            f = wf.cell(r, c).value
            v = wv.cell(r, c).value
            if hasattr(f, "text"):  # ArrayFormula
                f = "=" + f.text.lstrip("=")
            if isinstance(v, (_dt.datetime, _dt.date)):
                v = v.strftime("%d/%m/%Y")
            if isinstance(f, str) and f.startswith("="):
                if "!" in f or "[" in f:
                    if v not in (None, ""):
                        sh.set(r, c, value=v)
                else:
                    sh.set(r, c, f=f, cache=v if v is not None else 0)
                    nf += 1
            elif f not in (None, ""):
                if isinstance(f, (_dt.datetime, _dt.date)):
                    f = f.strftime("%d/%m/%Y")
                sh.set(r, c, value=f)
    sh.set(1, 12, value=note)
    sh.set_widths({1: 6, 2: 60, 3: 8, 4: 9, 5: 8, 6: 12, 7: 13, 8: 9, 9: 12, 10: 13, 11: 25, 12: 50})
    return sh.tobytes(), wf.max_row, nf


def tv_rows_from_ws(ws):
    """Danh mục TVGIA từ worksheet openpyxl (data_only)."""
    out, sec, parent = [], "", None
    names = {}
    prev_letter = None
    for r in range(3, ws.max_row + 1):
        stt, name = ws.cell(r, 1).value, ws.cell(r, 2).value
        if name in (None, ""):
            continue
        unit = ws.cell(r, 3).value
        g, j = to_num(ws.cell(r, 7).value), to_num(ws.cell(r, 10).value)
        cls = seq_class(stt, prev_letter)
        if re.fullmatch(r"[A-Z]", str(stt or "").strip()) and unit in (None, ""):  # TVGIA không dùng số La Mã
            sec = f"{stt} {str(name).strip()}"
            prev_letter = str(stt).strip()
        row = {"row": r, "stt": stt, "ten": str(name).strip(), "dvt": unit, "vl": g, "nc": j, "muc": sec,
               "tieu_de": unit in (None, "") and g is None and j is None,
               "dong_con": stt in (None, "") and parent is not None and unit not in (None, "")}
        if stt not in (None, ""):
            parent = r
        row["cha"] = parent if row["dong_con"] else None
        out.append(row)
        names.setdefault(norm_key(name), []).append(r)
    for x in out:
        x["trung_ten"] = [y for y in names[norm_key(x["ten"])] if y != x["row"]]
    return out


def cmd_tv(args):
    wb = openpyxl.load_workbook(args[0], data_only=True)
    if "TVGIA" not in wb.sheetnames:
        print("File chưa có sheet TVGIA -> dừng (chép sheet từ file thư viện XD hoặc chạy build).")
        return 1
    rows = tv_rows_from_ws(wb["TVGIA"])
    Path(args[1]).write_text(json.dumps(rows, ensure_ascii=False, indent=0), encoding="utf-8")
    dup = sum(1 for x in rows if x["trung_ten"])
    child = sum(1 for x in rows if x["dong_con"])
    muc = []
    for x in rows:
        if x["muc"] and (not muc or muc[-1] != x["muc"]):
            muc.append(x["muc"])
    print(f"TVGIA: {len(rows)} dòng có tên, {dup} dòng trùng tên (dùng tham chiếu trực tiếp), {child} dòng con định mức.")
    print("Mục:", " | ".join(muc))
    print("Đã ghi", args[1])


# ----------------------------------------------------------------------------
# build
# ----------------------------------------------------------------------------
def classify_rows(ws, sc):
    """Phân loại dòng dữ liệu CĐT: work / title / note / blank."""
    col = lambda k: CI(sc[k]) if sc.get(k) else None  # noqa
    cs, cd, cu, ck, cn = col("stt"), col("desc"), col("unit"), col("kl"), col("note")
    rows = []
    ranks = {}
    prev_letter = None
    for r in range(sc["data_start_row"], sc["data_end_row"] + 1):
        stt = ws.cell(r, cs).value if cs else None
        desc = ws.cell(r, cd).value
        unit = ws.cell(r, cu).value
        kl = ws.cell(r, ck).value if ck else None
        note = ws.cell(r, cn).value if cn else None
        hasu = unit not in (None, "") and str(unit).strip() != ""
        klv = to_num(kl)
        cls = seq_class(stt, prev_letter)
        if desc in (None, "") and not hasu and stt in (None, ""):
            kind = "blank"
        elif cls in ("roman", "letter", "star") and (not hasu or klv is None):
            kind = "title"
        elif cls in ("dotted", "num", "other") and not hasu and klv is None and desc not in (None, ""):
            kind = "title"
        elif hasu:
            kind = "work"
        else:
            kind = "note"
        x = {"row": r, "kind": kind, "stt": stt, "desc": desc, "unit": unit, "kl": kl, "note": note, "cls": cls}
        if kind == "title":
            ranks.setdefault(cls, len(ranks))
            x["level"] = ranks[cls]
            if cls == "letter":
                prev_letter = str(stt).strip().rstrip(".")
        rows.append(x)
    # phạm vi mục
    for i, x in enumerate(rows):
        if x["kind"] != "title":
            continue
        end = sc["data_end_row"]
        for y in rows[i + 1:]:
            if y["kind"] == "title" and y["level"] <= x["level"]:
                end = y["row"] - 1
                break
        x["end"] = end
        x["children"] = sum(1 for y in rows if x["row"] < y["row"] <= end and y["kind"] == "work")
    return rows


def build_sheet(pkg, st, ws, sc, sum_name):
    part = pkg.sheet_part(sc["sheet"])
    sh = Sheet(pkg.parts[part])
    c0 = CI(sc["block_start"])
    C = lambda k: L(c0 + k)  # noqa
    hdr = sc["header_rows"][1]
    top, coef = hdr - 1, hdr - 2
    if coef < 1:
        raise SystemExit(f"[{sc['sheet']}] tiêu đề sát đầu sheet: cần 2 hàng phía trên hàng tiêu đề {hdr}. Dừng, báo người dùng.")
    if not sc.get("kl"):
        raise SystemExit(f"[{sc['sheet']}] chưa có cột KL tổng (kl = null). Hỏi người dùng cột tổng KL rồi sửa xdcfg.json.")
    for r in (coef, top, hdr):
        for k in range(NCOL):
            if ws.cell(r, c0 + k).value not in (None, ""):
                raise SystemExit(f"[{sc['sheet']}] ô {C(k)}{r} đã có dữ liệu CĐT: dời block_start sang phải.")
    rows = classify_rows(ws, sc)
    S = st.get
    # --- hàng hệ số, tiêu đề
    sh.set(coef, c0 + C_NHAN, value="HS chung", s=S("lbl_coef"))
    sh.set(coef, c0 + C_VCT, f=f"={qs(sum_name)}!$C$6", cache=1, s=S("hdr_coef"))
    sh.set(coef, c0 + C_NCT, f=f"={qs(sum_name)}!$C$7", cache=1, s=S("hdr_coef"))
    single = {C_STT: "STT", C_DESC: MARK, C_CL: "CHỦNG LOẠI VẬT TƯ", C_UNIT: "ĐƠN VỊ", C_KLTDI: "Khối lượng TDI",
              C_KL: "KHỐI LƯỢNG", C_TT: "TỔNG", C_NOTE: "GHI CHÚ (CĐT)", C_NHAN: "NHÃN", C_SRC: "Nguồn giá",
              C_CONF: "Độ tin cậy", C_GNOTE: "Ghi chú gán giá"}
    groups = {(C_SL, C_HS): ("DIỄN GIẢI KHỐI LƯỢNG", ["SL", "Dài", "Rộng", "Cao", "Hệ số"]),
              (C_DGVL, C_DGT): ("ĐƠN GIÁ", ["VẬT LIỆU", "NHÂN CÔNG", "TỔNG"]),
              (C_VCT, C_VTT): ("VẬT LIỆU", ["Công tác", "Đơn vị", "Khối lượng", "Đơn giá", "Thành tiền"]),
              (C_NCT, C_NTT): ("NHÂN CÔNG", ["Công tác", "Đơn vị", "Khối lượng", "Đơn giá", "Thành tiền"])}
    hs = S("hdr")
    for k, t in single.items():
        sh.set(top, c0 + k, value=t, s=hs)
        sh.set(hdr, c0 + k, s=hs)
        sh.merge(f"{C(k)}{top}:{C(k)}{hdr}")
    for (a, b), (t, subs) in groups.items():
        sh.set(top, c0 + a, value=t, s=hs)
        for k in range(a + 1, b + 1):
            sh.set(top, c0 + k, s=hs)
        sh.merge(f"{C(a)}{top}:{C(b)}{top}")
        for i, t2 in enumerate(subs):
            sh.set(hdr, c0 + a + i, value=t2, s=hs)
    # --- dòng dữ liệu
    col = {k: sc.get(k) for k in ("stt", "desc", "unit", "kl", "note")}
    cv = lambda r, k: ws[f"{col[k]}{r}"].value if col[k] else None  # noqa
    hsvl, hsnc = f"${C(C_VCT)}${coef}", f"${C(C_NCT)}${coef}"
    n_work = n_title = n_note = 0
    for x in rows:
        r = x["row"]
        if x["kind"] == "blank":
            continue
        ref = lambda k: f'=IF(ISBLANK({col[k]}{r}),"",{col[k]}{r})'  # noqa
        cache = lambda k: "" if cv(r, k) is None else cv(r, k)  # noqa
        if x["kind"] == "title":
            n_title += 1
            sty_t, sty_n = S("ttl_txt"), S("ttl_num")
            if col["stt"]:
                sh.set(r, c0 + C_STT, f=ref("stt"), cache=cache("stt"), s=S("ttl_c"))
            sh.set(r, c0 + C_DESC, f=ref("desc"), cache=cache("desc"), s=sty_t)
            for k in range(C_CL, NCOL):
                if k in (C_STT, C_DESC):
                    continue
                sh.set(r, c0 + k, s=sty_n if k in (C_TT, C_VTT, C_NTT) else sty_t)
            if x["children"]:
                sh.set(r, c0 + C_TT, f=f"=SUBTOTAL(9,{C(C_TT)}{r + 1}:{C(C_TT)}{x['end']})", cache=0, s=sty_n)
            continue
        if x["kind"] == "note":
            n_note += 1
            if col["stt"]:
                sh.set(r, c0 + C_STT, f=ref("stt"), cache=cache("stt"), s=S("txt_c"))
            sh.set(r, c0 + C_DESC, f=ref("desc"), cache=cache("desc"), s=S("txt"))
            continue
        n_work += 1
        t, n0, n2 = S("txt"), S("num0"), S("num2")
        if col["stt"]:
            sh.set(r, c0 + C_STT, f=ref("stt"), cache=cache("stt"), s=S("txt_c"))
        sh.set(r, c0 + C_DESC, f=ref("desc"), cache=cache("desc"), s=t)
        sh.set(r, c0 + C_CL, s=S("in_txt"))
        sh.set(r, c0 + C_UNIT, f=ref("unit"), cache=cache("unit"), s=S("txt_c"))
        for k in (C_SL, C_DAI, C_RONG, C_CAO, C_HS, C_KLTDI):
            sh.set(r, c0 + k, s=S("in_num"))
        kl_cache = to_num(cv(r, "kl"))
        sh.set(r, c0 + C_KL, f=ref("kl"), cache=cache("kl"), s=n2)
        KL = f"{C(C_KL)}{r}"
        sh.set(r, c0 + C_DGVL, f=f"=IF(N({KL})=0,0,ROUND({C(C_VTT)}{r}*{hsvl}/N({KL}),0))", cache=0, s=n0)
        sh.set(r, c0 + C_DGNC, f=f"=IF(N({KL})=0,0,ROUND({C(C_NTT)}{r}*{hsnc}/N({KL}),0))", cache=0, s=n0)
        sh.set(r, c0 + C_DGT, f=f"={C(C_DGVL)}{r}+{C(C_DGNC)}{r}", cache=0, s=n0)
        sh.set(r, c0 + C_TT, f=f"=ROUND(N({KL})*{C(C_DGT)}{r},0)", cache=0, s=n0)
        if col["note"]:
            sh.set(r, c0 + C_NOTE, f=ref("note"), cache=cache("note"), s=t)
        else:
            sh.set(r, c0 + C_NOTE, s=t)
        sh.set(r, c0 + C_NHAN, s=t)
        for ct, dv, kk, dg, tt in ((C_VCT, C_VDV, C_VKL, C_VDG, C_VTT), (C_NCT, C_NDV, C_NKL, C_NDG, C_NTT)):
            sh.set(r, c0 + ct, s=t)
            sh.set(r, c0 + dv, f=f"={C(C_UNIT)}{r}", cache=cache("unit"), s=S("txt_c"))
            sh.set(r, c0 + kk, f=f"={KL}", cache=kl_cache if kl_cache is not None else "", s=n2)
            sh.set(r, c0 + dg, s=n0)
            sh.set(r, c0 + tt, f=f"=N({C(kk)}{r})*N({C(dg)}{r})", cache=0, s=n0)
        sh.set(r, c0 + C_SRC, value=SRC_NONE, s=S("txt_c"))
        sh.set(r, c0 + C_CONF, s=S("txt_c"))
        sh.set(r, c0 + C_GNOTE, s=t)
    # --- tổng cộng
    ds, de = sc["data_start_row"], sc["data_end_row"]
    tr = de + 2
    for k in range(NCOL):
        if ws.cell(tr, c0 + k).value not in (None, ""):
            raise SystemExit(f"[{sc['sheet']}] ô {C(k)}{tr} (dòng tổng TDI) đã có dữ liệu.")
    sh.set(tr, c0 + C_DESC, value="TỔNG CỘNG (TDI)", s=S("tot_txt"))
    for k in range(NCOL):
        if k == C_DESC:
            continue
        if k in (C_TT, C_VTT, C_NTT):
            sh.set(tr, c0 + k, f=f"=SUBTOTAL(9,{C(k)}{ds}:{C(k)}{de})", cache=0, s=S("tot_num"))
        else:
            sh.set(tr, c0 + k, s=S("tot_txt"))
    # --- độ rộng, bộ lọc
    sh.set_widths({c0 + k: w for k, w in enumerate(WIDTHS)} | {c0 - 1: 2})
    af = sh.root.find(q("autoFilter"))
    filt = None
    if af is None:
        af = ensure_child(sh.root, "autoFilter")
        filt = f"{C(0)}{hdr}:{C(NCOL - 1)}{de}"
        af.set("ref", filt)
    info = {"sheet": sc["sheet"], "block": f"{C(0)}:{C(NCOL - 1)}", "hang_he_so": coef, "hang_tieu_de": [top, hdr],
            "dong_tong": tr, "cong_viec": n_work, "tieu_de": n_title, "ghi_chu": n_note, "loc": filt,
            "tt": f"{qs(sc['sheet'])}!{C(C_TT)}{tr}", "vl": f"{qs(sc['sheet'])}!{C(C_VTT)}{tr}",
            "nc": f"{qs(sc['sheet'])}!{C(C_NTT)}{tr}", "c0": C(0)}
    if filt is None:
        info["canh_bao"] = f"sheet đã có bộ lọc ({af.get('ref')}), chưa thêm lọc cho cột Nguồn giá {C(C_SRC)}"
    pkg.parts[part] = sh.tobytes()
    return info


def summary_xml(st, project, infos, old_coef=None):
    base = etree.tostring(etree.Element(q("worksheet"), nsmap={None: NS})).replace(b"/>", b"><sheetData/></worksheet>")
    sh = Sheet(base)
    S = st.get
    sh.set(1, 1, value="BẢNG TỔNG HỢP PHẦN XÂY DỰNG (TDI)", s=S("big"))
    sh.set(2, 1, value="Dự án: " + (project or ""), s=S("bold"))
    sh.set(4, 1, value=f"Tạo bởi boq_xd_tools {VERSION} ngày {_dt.date.today():%d/%m/%Y}")
    sh.set(6, 2, value="Hệ số chung VL", s=S("lbl_coef"))
    hv, hn = (old_coef or (1, 1))
    sh.set(6, 3, value=hv if is_num(hv) else 1, s=S("in_num"))
    sh.set(7, 2, value="Hệ số chung NC", s=S("lbl_coef"))
    sh.set(7, 3, value=hn if is_num(hn) else 1, s=S("in_num"))
    sh.set(6, 4, value="Trưởng phòng chốt (BKY dùng 1,12). Đơn giá khối XD = thành tiền phân tích × hệ số / KL.")
    hdr = ["STT", "Hạng mục (sheet)", "Thành tiền VL phân tích (trước HS)", "Thành tiền NC phân tích (trước HS)",
           "Thành tiền (sau HS)", "Ghi chú"]
    for i, t in enumerate(hdr, 1):
        sh.set(9, i, value=t, s=S("hdr"))
    r = 10
    for i, x in enumerate(infos, 1):
        sh.set(r, 1, value=i, s=S("txt_c"))
        sh.set(r, 2, value=x["sheet"], s=S("txt"))
        sh.set(r, 3, f="=" + x["vl"], cache=0, s=S("num0"))
        sh.set(r, 4, f="=" + x["nc"], cache=0, s=S("num0"))
        sh.set(r, 5, f="=" + x["tt"], cache=0, s=S("num0"))
        sh.set(r, 6, value=f"khối XD {x['block']}, dòng tổng {x['dong_tong']}", s=S("txt"))
        r += 1
    sh.set(r, 2, value="TỔNG CỘNG", s=S("tot_txt"))
    sh.set(r, 1, s=S("tot_txt"))
    sh.set(r, 6, s=S("tot_txt"))
    for c in (3, 4, 5):
        sh.set(r, c, f=f"=SUM({L(c)}10:{L(c)}{r - 1})", cache=0, s=S("tot_num"))
    sh.set_widths({1: 6, 2: 40, 3: 22, 4: 22, 5: 22, 6: 40})
    return sh.tobytes()


def tvgia_compare(boq_wb, lib):
    ws = boq_wb["TVGIA"]
    n = sg = sj = 0
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 2).value not in (None, ""):
            n += 1
            sg += to_num(ws.cell(r, 7).value) or 0
            sj += to_num(ws.cell(r, 10).value) or 0
    cur = {"dong_co_ten": n, "sum_G": round(sg), "sum_J": round(sj), "L1": ws["L1"].value}
    same = lib and lib.get("ok") and all(cur[k] == lib[k] for k in ("dong_co_ten", "sum_G", "sum_J"))
    return cur, bool(same)


def cmd_build(args):
    nguon = None
    if "--nguon" in args:
        i = args.index("--nguon")
        nguon = args[i + 1]
        args = args[:i] + args[i + 2:]
    src, out, cfgp = args[0], args[1], args[2]
    lib = args[3] if len(args) > 3 else None
    cfg = json.loads(Path(cfgp).read_text(encoding="utf-8"))
    if not cfg.get("sheets"):
        raise SystemExit("xdcfg.json không có sheet XD nào.")
    wbv = openpyxl.load_workbook(src, data_only=True)
    has_tv = "TVGIA" in wbv.sheetnames
    libres = None
    if lib:
        libres = check_lib(lib)
        if not libres["ok"] and not has_tv:
            print(json.dumps(libres, ensure_ascii=False, indent=1))
            raise SystemExit("Thư viện XD không đạt kiểm bố cục -> dừng, không build.")
    elif not has_tv:
        raise SystemExit("Chưa chỉ file thư viện XD và file BOQ chưa có TVGIA -> dừng.")
    pkg = Package(src)
    names = [n for n, _, _ in pkg.sheets()]
    st = Styles(pkg)
    sum_name = "Summary XD"
    infos = []
    for sc in cfg["sheets"]:
        if sc["sheet"] not in names:
            raise SystemExit(f"Không có sheet '{sc['sheet']}' trong file.")
        ws = wbv[sc["sheet"]]
        hb = has_block(ws, sc["header_rows"][1])
        if hb:
            print(f"- '{sc['sheet']}' đã có khối XD từ cột {L(hb - 1)}: bỏ qua (không chuyển hai lần).")
            continue
        infos.append(build_sheet(pkg, st, ws, sc, sum_name))
    if not infos:
        raise SystemExit("Không dựng được sheet nào.")
    # TVGIA
    tvinfo = {}
    if has_tv:
        cur, same = tvgia_compare(wbv, libres)
        tvinfo = {"trang_thai": "giữ TVGIA có sẵn trong file", "hien_tai": cur,
                  "nguon_moi": {k: libres.get(k) for k in ("file", "dong_co_ten", "sum_G", "sum_J")} if libres else None,
                  "khop_nguon": same}
        if libres and not same:
            tvinfo["canh_bao"] = "TVGIA trong file lệch file nguồn vừa chọn -> hỏi giữ ảnh chụp cũ hay thay (không tự thay)"
    else:
        note = nguon or f"Nguồn: {Path(lib).name} (sửa {libres['sua_ngay']}) · nạp {_dt.date.today():%d/%m/%Y} · boq_xd_tools {VERSION}"
        xml, nrow, nf = tvgia_xml(lib, note)
        pkg.add_sheet("TVGIA", xml)
        tvinfo = {"trang_thai": "đã nạp", "dong": nrow, "cong_thuc_giu": nf, "L1": note,
                  "dong_co_ten": libres["dong_co_ten"], "sum_G": libres["sum_G"], "sum_J": libres["sum_J"]}
    # Summary XD: mọi sheet có khối XD (dựng lần này + đã có từ trước), theo thứ tự sheet
    existing = dict((n, t) for n, t, _ in pkg.sheets())
    built = {x["sheet"]: x for x in infos}
    allinfo = []
    for n in wbv.sheetnames:
        if n in built:
            allinfo.append(built[n])
            continue
        b = find_block(wbv[n])
        if b:
            c0 = b["c0"]
            allinfo.append({"sheet": n, "block": f"{L(c0)}:{L(c0 + NCOL - 1)}", "dong_tong": b["tot"],
                            "tt": f"{qs(n)}!{L(c0 + C_TT)}{b['tot']}", "vl": f"{qs(n)}!{L(c0 + C_VTT)}{b['tot']}",
                            "nc": f"{qs(n)}!{L(c0 + C_NTT)}{b['tot']}"})
    old_coef = None
    if sum_name in wbv.sheetnames:
        old_coef = (wbv[sum_name]["C6"].value, wbv[sum_name]["C7"].value)
    sx = summary_xml(st, cfg.get("ma_da") or "", allinfo, old_coef)
    if sum_name in existing:
        pkg.parts[existing[sum_name]] = sx
    else:
        pkg.add_sheet(sum_name, sx)
    # bộ lọc: tên ẩn _FilterDatabase
    order = [n for n, _, _ in pkg.sheets()]
    for x in infos:
        if x.get("loc"):
            a, b = x["loc"].split(":")
            ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups()
            cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
            pkg.add_defined_name("_xlnm._FilterDatabase", f"{qs(x['sheet'])}!${ca}${ra}:${cb}${rb}",
                                 local_id=order.index(x["sheet"]), hidden=True)
    st.save()
    pkg.drop_calcchain()
    pkg.save(out)
    rep = {"tool": VERSION, "input": src, "output": out, "sheets": infos, "TVGIA": tvinfo, "summary": sum_name}
    Path(out + ".build.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    for x in infos:
        print(f"[{x['sheet']}] khối XD {x['block']}, hệ số hàng {x['hang_he_so']}, tiêu đề {x['hang_tieu_de']}, "
              f"{x['cong_viec']} dòng công việc / {x['tieu_de']} tiêu đề / {x['ghi_chu']} ghi chú, dòng tổng {x['dong_tong']}"
              + (f", lọc {x['loc']}" if x.get("loc") else "") + (f"  ! {x['canh_bao']}" if x.get("canh_bao") else ""))
    print("TVGIA:", json.dumps(tvinfo, ensure_ascii=False))
    print("Đã ghi", out)


# ----------------------------------------------------------------------------
# Đọc khối XD đã dựng
# ----------------------------------------------------------------------------
def find_block(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        for ce in row:
            if isinstance(ce.value, str) and ce.value.strip().upper() == MARK:
                c0, top = ce.column - 1, ce.row
                tot = None
                for r in range(top + 2, ws.max_row + 1):
                    v = ws.cell(r, c0 + C_DESC).value
                    if isinstance(v, str) and v.strip().upper() == "TỔNG CỘNG (TDI)":
                        tot = r
                        break
                if tot is None:
                    return None
                return {"c0": c0, "top": top, "hdr": top + 1, "coef": top - 1, "ds": top + 2, "de": tot - 2, "tot": tot}
    return None


_REF1 = re.compile(r"ISBLANK\(\$?([A-Z]+)\$?(\d+)\)")


def ref_col(formula):
    m = _REF1.search(formula or "")
    return m.group(1) if m else None


def block_rows(wf, wv, b):
    """Dòng của khối: kind work/title/note, kèm giá trị đọc được."""
    c0 = b["c0"]
    out, path = [], []
    prev_letter = None
    for r in range(b["ds"], b["de"] + 1):
        fv = wf.cell(r, c0 + C_VTT).value
        f1 = wf.cell(r, c0 + C_DESC).value
        if not (isinstance(f1, str) and f1.startswith("=")) and f1 in (None, ""):
            continue
        fill = wf.cell(r, c0 + C_DESC).fill
        is_title = fill is not None and fill.fill_type == "solid" and str(fill.fgColor.rgb or "").upper().endswith("FFFF00")
        val = lambda k: wv.cell(r, c0 + k).value  # noqa
        x = {"row": r, "stt": val(C_STT), "desc": val(C_DESC), "unit": val(C_UNIT), "kl": val(C_KL),
             "note": val(C_NOTE)}
        if x["desc"] is None and isinstance(f1, str):  # không có cache: đọc thẳng ô CĐT
            cd = ref_col(f1)
            x["desc"] = wv[f"{cd}{r}"].value if cd else None
            for k, kk in (("stt", C_STT), ("unit", C_UNIT), ("kl", C_KL), ("note", C_NOTE)):
                cc = ref_col(wf.cell(r, c0 + kk).value)
                if cc:
                    x[k] = wv[f"{cc}{r}"].value
        if is_title:
            x["kind"] = "title"
            lvl = seq_class(x["stt"], prev_letter)
            if lvl == "letter":
                prev_letter = str(x["stt"]).strip().rstrip(".")
            path = [p for p in path if p[0] != lvl and p[2] >= r]
            fsub = wf.cell(r, c0 + C_TT).value
            m = re.search(r":\$?[A-Z]+\$?(\d+)\)", fsub or "")
            end = int(m.group(1)) if m else r
            path.append((lvl, f"{x['stt'] or ''} {x['desc'] or ''}".strip(), end))
        elif isinstance(fv, str) and fv.upper().startswith("=N("):
            x["kind"] = "work"
        else:
            x["kind"] = "note"
        x["muc"] = " > ".join(p[1] for p in path if p[2] >= r and p[1] != f"{x['stt'] or ''} {x['desc'] or ''}".strip()) \
            if x["kind"] != "title" else ""
        if x["kind"] == "work":
            fvl, fnc = wf.cell(r, c0 + C_VDG).value, wf.cell(r, c0 + C_NDG).value
            x["vl_ct"], x["nc_ct"] = wv.cell(r, c0 + C_VCT).value, wv.cell(r, c0 + C_NCT).value
            x["f_vct"], x["f_nct"] = wf.cell(r, c0 + C_VCT).value, wf.cell(r, c0 + C_NCT).value
            x["f_vdg"], x["f_ndg"] = fvl, fnc
            x["src"] = wv.cell(r, c0 + C_SRC).value
            x["conf"] = wv.cell(r, c0 + C_CONF).value
            x["gnote"] = wv.cell(r, c0 + C_GNOTE).value
            hard = [v for v in (fvl, fnc) if is_num(v)]
            if EXCL.search(detone(x["note"])):
                x["status"] = "loai_tru"
            elif x["src"] == SRC_WEB:
                x["status"] = "gia_tham_khao"
            elif x["f_vct"] or x["f_nct"]:
                x["status"] = "da_gan"
            elif hard:
                x["status"] = "so_cung"
            else:
                x["status"] = "chua_gan"
        out.append(x)
    return out


def _load2(path):
    return openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)


def _sheets_with_block(wf, names):
    out = []
    for n in (names or wf.sheetnames):
        if n not in wf.sheetnames:
            print(f"! không có sheet '{n}'")
            continue
        b = find_block(wf[n])
        if b:
            out.append((n, b))
        elif names:
            print(f"! sheet '{n}' chưa có khối XD")
    return out


def cmd_todo(args):
    f, out, names = args[0], args[1], args[2:]
    wf, wv = _load2(f)
    res = {"file": f, "sheets": [], "can_gan": []}
    for n, b in _sheets_with_block(wf, names):
        rows = block_rows(wf[n], wv[n], b)
        work = [x for x in rows if x["kind"] == "work"]
        # số dòng có ĐVT của CĐT trong phạm vi dữ liệu
        cu = None
        for x in work:
            cu = ref_col(wf[n].cell(x["row"], b["c0"] + C_UNIT).value)
            if cu:
                break
        n_cdt = 0
        titles = {x["row"] for x in rows if x["kind"] == "title"}
        if cu:
            for r in range(b["ds"], b["de"] + 1):
                v = wv[n][f"{cu}{r}"].value
                if v not in (None, "") and str(v).strip() and r not in titles:
                    n_cdt += 1
        cnt = {}
        for x in work:
            cnt[x["status"]] = cnt.get(x["status"], 0) + 1
        res["sheets"].append({"sheet": n, "block": f"{L(b['c0'])}:{L(b['c0'] + NCOL - 1)}", "dong_tong": b["tot"],
                              "cong_viec": len(work), "co_dvt_cdt": n_cdt, "trang_thai": cnt,
                              "tieu_de": sum(1 for x in rows if x["kind"] == "title")})
        for x in work:
            if x["status"] == "chua_gan":
                res["can_gan"].append({"sheet": n, **{k: x[k] for k in ("row", "stt", "desc", "unit", "kl", "note", "muc")}})
        flag = "" if len(work) == n_cdt else f"  ! LỆCH: CĐT có {n_cdt} dòng có ĐVT"
        print(f"[{n}] {len(work)} dòng công việc ({cnt}), dòng có ĐVT của CĐT: {n_cdt}{flag}")
    Path(out).write_text(json.dumps(res, ensure_ascii=False, indent=0, default=str), encoding="utf-8")
    print(f"Đã ghi {out}: {len(res['can_gan'])} dòng cần gán.")


# ----------------------------------------------------------------------------
# recalc / cmp
# ----------------------------------------------------------------------------
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub RecalculateAndSave()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>"""


def _lo_env():
    for p in ("/mnt/skills/public/xlsx/scripts",):
        if os.path.isdir(p) and p not in sys.path:
            sys.path.append(p)
    try:
        from office.soffice import get_soffice_env  # type: ignore
        return get_soffice_env()
    except Exception:  # noqa
        env = os.environ.copy()
        env["SAL_USE_VCLPLUGIN"] = "svp"
        return env


def lo_recalc(path, timeout=180):
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return "không có LibreOffice (soffice) trên máy"
    env = _lo_env()
    with tempfile.TemporaryDirectory(prefix="xd-lo-") as prof:
        url = Path(prof).as_uri()
        subprocess.run([exe, "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       capture_output=True, timeout=timeout, env=env)
        md = Path(prof) / "user" / "basic" / "Standard"
        if not md.exists():
            return "LibreOffice không tạo được hồ sơ người dùng"
        (md / "Module1.xba").write_text(MACRO)
        before = os.stat(path).st_mtime_ns
        r = subprocess.run([exe, "--headless", "--norestore", f"-env:UserInstallation={url}",
                            "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
                            str(Path(path).resolve())], capture_output=True, text=True, timeout=timeout + 30, env=env)
        if r.returncode != 0:
            return "LibreOffice lỗi: " + (r.stderr or "")[:300]
        if os.stat(path).st_mtime_ns == before:
            return "LibreOffice không ghi lại file (có phiên LibreOffice khác đang chạy?)"
    return None


ERR_RX = re.compile(r"^#(VALUE!|REF!|NAME\?|DIV/0!|N/A|NULL!|NUM!)")
_MAC = re.compile(r"\b(m|mac|b|c)\s?(\d{2,3})\b")
_DIM = re.compile(r"\b(d|dn|phi|day|dày)\s?(\d+(?:[.,]\d+)?)")


def _marks(s):
    t = norm_key(s)
    return {m.group(2) for m in _MAC.finditer(t)}, {m.group(2).replace(",", ".") for m in _DIM.finditer(t)}


def cmd_recalc(args):
    f, out, names = args[0], args[1], args[2:]
    tmp = tempfile.mkdtemp(prefix="xd-rc-")
    cp = os.path.join(tmp, Path(f).name)
    shutil.copy(f, cp)
    err = lo_recalc(cp)
    if err:
        print("! Không tính lại được:", err, "-> đọc giá trị cache hiện có.")
    wf = openpyxl.load_workbook(f)
    wv = openpyxl.load_workbook(cp, data_only=True)
    tvv = wv["TVGIA"] if "TVGIA" in wv.sheetnames else None
    rep = {"file": f, "recalc": err or "LibreOffice OK", "sheets": []}
    total_err = 0
    for n, b in _sheets_with_block(wf, names):
        ws, wsv = wf[n], wv[n]
        c0 = b["c0"]
        errs, unit_bad, mark_bad, loops = [], [], [], []
        for r in range(b["coef"], b["tot"] + 1):
            for k in range(NCOL):
                v = wsv.cell(r, c0 + k).value
                if isinstance(v, str) and ERR_RX.match(v):
                    errs.append(f"{L(c0 + k)}{r}={v}")
            fsub = ws.cell(r, c0 + C_TT).value
            if isinstance(fsub, str) and fsub.upper().startswith("=SUBTOTAL") and r != b["tot"]:
                m = re.search(r"(\d+):\$?[A-Z]+\$?(\d+)\)", fsub)
                if m and (int(m.group(2)) >= b["tot"] or int(m.group(1)) <= r):
                    loops.append(f"{L(c0 + C_TT)}{r}: {fsub}")
            ct = ws.cell(r, c0 + C_VCT).value
            m = re.match(r"^=\s*'?TVGIA'?!\$?B\$?(\d+)$", str(ct or ""))
            if m and tvv is not None:
                tr = int(m.group(1))
                u1, u2 = norm_unit(wsv.cell(r, c0 + C_UNIT).value), norm_unit(tvv.cell(tr, 3).value)
                if u1 and u2 and u1 != u2:
                    unit_bad.append(f"dòng {r}: ĐVT BOQ '{wsv.cell(r, c0 + C_UNIT).value}' ≠ TVGIA!C{tr} '{tvv.cell(tr, 3).value}'")
                m1, d1 = _marks(wsv.cell(r, c0 + C_DESC).value)
                m2, d2 = _marks(tvv.cell(tr, 2).value)
                if (m1 and m2 and not (m1 & m2)) or (d1 and d2 and not (d1 & d2)):
                    mark_bad.append(f"dòng {r}: mác/kích thước BOQ {sorted(m1 | d1)} ≠ TVGIA dòng {tr} {sorted(m2 | d2)}")
        tot = {k: wsv.cell(b["tot"], c0 + kk).value for k, kk in (("thanh_tien", C_TT), ("vl_phan_tich", C_VTT), ("nc_phan_tich", C_NTT))}
        total_err += len(errs)
        rep["sheets"].append({"sheet": n, "o_loi": len(errs), "loi": errs[:50], "tong": tot, "vong_tham_chieu": loops,
                              "lech_dvt": unit_bad, "lech_mac_kich_thuoc": mark_bad})
        print(f"[{n}] ô lỗi {len(errs)}, tổng TDI {tot}, vòng tham chiếu {len(loops)}, lệch ĐVT {len(unit_bad)}, "
              f"lệch mác/kích thước {len(mark_bad)}")
        for e in (errs[:10] + loops[:5] + unit_bad[:10] + mark_bad[:10]):
            print("   ", e)
    if "Summary XD" in wv.sheetnames:
        s = wv["Summary XD"]
        rep["summary"] = {"HS_VL": s["C6"].value, "HS_NC": s["C7"].value}
        for r in range(10, s.max_row + 1):
            if s.cell(r, 2).value == "TỔNG CỘNG":
                rep["summary"]["tong"] = [s.cell(r, c).value for c in (3, 4, 5)]
        print("Summary XD:", rep["summary"])
    rep["tong_o_loi"] = total_err
    Path(out).write_text(json.dumps(rep, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    shutil.rmtree(tmp, ignore_errors=True)
    print("Đã ghi", out)


def _sig(ce):
    b = ce.border
    return (ce.value if not hasattr(ce.value, "text") else ce.value.text, bool(ce.font and ce.font.b),
            str(ce.fill.fgColor.rgb) if ce.fill and ce.fill.fill_type else None,
            tuple(getattr(getattr(b, s), "style", None) for s in ("left", "right", "top", "bottom")) if b else None)


def cmd_cmp(args):
    a, b, names = args[0], args[1], args[2:]
    wa, wb_ = openpyxl.load_workbook(a), openpyxl.load_workbook(b)
    total = 0
    for n in (names or wa.sheetnames):
        if n not in wb_.sheetnames:
            print(f"! '{n}' không có trong file ra")
            total += 1
            continue
        x, y = wa[n], wb_[n]
        blk = find_block(y)
        maxc = (blk["c0"] - 1) if blk else x.max_column
        diffs = []
        for row in x.iter_rows(min_row=1, max_row=x.max_row, max_col=min(x.max_column, maxc)):
            for ce in row:
                if _sig(ce) != _sig(y.cell(ce.row, ce.column)):
                    diffs.append(ce.coordinate)
        mx = {str(m) for m in x.merged_cells.ranges}
        my = {str(m) for m in y.merged_cells.ranges}
        lost = mx - my
        dw = [c for c, d in x.column_dimensions.items() if CI(c) <= maxc and
              abs((d.width or 0) - (y.column_dimensions[c].width or 0)) > 0.01]
        total += len(diffs) + len(lost) + len(dw)
        print(f"[{n}] ô CĐT khác: {len(diffs)}" + (f" {diffs[:20]}" if diffs else "") +
              f"; merge mất: {len(lost)}; độ rộng cột đổi: {len(dw)}")
    print("KẾT QUẢ:", "không khác ô nào" if total == 0 else f"{total} khác biệt -> kiểm lại")
    return 0 if total == 0 else 1


# ----------------------------------------------------------------------------
# Bảng khớp XD: apply / export
# ----------------------------------------------------------------------------
KHOP_COLS = ["Dự án", "Sheet", "Dòng", "STT", "Mô tả CĐT", "ĐVT", "Khóa", "Tên TVGIA", "ĐVT TVGIA", "Dòng TVGIA (bản lưu)",
             "Cộng thêm VL (tên TVGIA)", "NC từ (tên TVGIA)", "HS VL", "HS NC", "Có VL", "Có NC", "Độ tin cậy",
             "Ghi chú gán giá", "Ngày lưu"]


def khop_key(desc, unit):
    return norm_key(desc) + "|" + norm_unit(unit)


def read_khop(path):
    if not path or not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Khop_XD" not in wb.sheetnames:
        return []
    ws = wb["Khop_XD"]
    hdr = [c.value for c in ws[1]]
    out = []
    for r in range(2, ws.max_row + 1):
        d = {hdr[i]: ws.cell(r, i + 1).value for i in range(len(hdr)) if hdr[i]}
        if d.get("Khóa"):
            out.append(d)
    return out


def tv_index(wv):
    ws = wv["TVGIA"]
    idx = {}
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if v not in (None, ""):
            idx.setdefault(norm_key(v), []).append(r)
    return ws, idx


def resolve_tv(ws, idx, name, unit=None, row_hint=None):
    if not name:
        return None, None
    c = idx.get(norm_key(name), [])
    if not c:
        return None, "tên TVGIA không còn trong bản đang dùng"
    if len(c) == 1:
        return c[0], None
    if row_hint and int(row_hint) in c:
        return int(row_hint), None
    cu = [r for r in c if unit and norm_unit(ws.cell(r, 3).value) == norm_unit(unit)]
    return (cu or c)[0], f"tên trùng ở các dòng {c}: chọn dòng {(cu or c)[0]}, cần kiểm"


def cmd_apply(args):
    f, khop, todo, out, ma = args[0], args[1], args[2], args[3], (args[4] if len(args) > 4 else "")
    wv = openpyxl.load_workbook(f, data_only=True)
    if "TVGIA" not in wv.sheetnames:
        raise SystemExit("File chưa có TVGIA -> dừng.")
    ws, idx = tv_index(wv)
    items = json.loads(Path(todo).read_text(encoding="utf-8"))
    items = items.get("can_gan", items) if isinstance(items, dict) else items
    rows = read_khop(khop)
    if not rows:
        print(f"! Chưa có Bảng khớp XD ({khop}) hoặc bảng rỗng: chạy không có Bảng khớp, kết quả có thể lệch lần trước.")
    by_key = {}
    for d in rows:
        by_key.setdefault(d["Khóa"], []).append(d)
    assign, left, warn = [], [], []
    for it in items:
        k = khop_key(it.get("desc"), it.get("unit"))
        cand = by_key.get(k, [])
        same = [d for d in cand if str(d.get("Dự án") or "") == str(ma)]
        exact = [d for d in same if d.get("Sheet") == it["sheet"] and str(d.get("Dòng")) == str(it["row"])]
        pick = exact or same or cand
        if not pick:
            left.append(it)
            continue
        sigs = {}
        for d in pick:
            sg = tuple(str(d.get(h) or "") for h in ("Tên TVGIA", "Cộng thêm VL (tên TVGIA)", "NC từ (tên TVGIA)",
                                                      "HS VL", "HS NC", "Có VL", "Có NC"))
            sigs.setdefault(sg, d)
        if len(sigs) > 1:
            opts = [f"{d.get('Sheet')}!{d.get('Dòng')}: {d.get('Tên TVGIA')}" +
                    (f" + {d.get('Cộng thêm VL (tên TVGIA)')}" if d.get("Cộng thêm VL (tên TVGIA)") else "") for d in sigs.values()]
            left.append({**it, "ly_do": "Bảng khớp có nhiều cách gán cho cùng mô tả: " + " | ".join(opts)})
            continue
        same = same if exact or same else []
        d = pick[-1]
        tv, w1 = resolve_tv(ws, idx, d.get("Tên TVGIA"), d.get("ĐVT TVGIA"), d.get("Dòng TVGIA (bản lưu)"))
        if d.get("Có VL") not in (False, 0, "0", "Không") and d.get("Tên TVGIA") and tv is None:
            warn.append(f"{it['sheet']}!{it['row']}: {w1}")
            left.append({**it, "ly_do": w1})
            continue
        a = {"sheet": it["sheet"], "row": it["row"], "tv": tv,
             "conf": d.get("Độ tin cậy") or ("Bảng khớp" if same else "Bảng khớp (DA khác)"),
             "note": ("Bảng khớp XD" + ("" if same else f" (DA {d.get('Dự án')})") +
                      (f"; {d.get('Ghi chú gán giá')}" if d.get("Ghi chú gán giá") else ""))}
        if d.get("Cộng thêm VL (tên TVGIA)"):
            a["plus_vl"], w = resolve_tv(ws, idx, d["Cộng thêm VL (tên TVGIA)"])
            if w:
                warn.append(f"{it['sheet']}!{it['row']} cộng thêm: {w}")
        if d.get("NC từ (tên TVGIA)"):
            a["tv_nc"], w = resolve_tv(ws, idx, d["NC từ (tên TVGIA)"])
            if w:
                warn.append(f"{it['sheet']}!{it['row']} NC: {w}")
        for kk, col in (("coef_vl", "HS VL"), ("coef_nc", "HS NC")):
            if to_num(d.get(col)) not in (None, 1.0):
                a[kk] = to_num(d.get(col))
        for kk, col in (("vl", "Có VL"), ("nc", "Có NC")):
            if d.get(col) in (False, 0, "0", "Không"):
                a[kk] = False
        if w1:
            warn.append(f"{it['sheet']}!{it['row']}: {w1}")
        assign.append(a)
    Path(out).write_text(json.dumps({"assign": assign, "left": left, "canh_bao": warn}, ensure_ascii=False, indent=0, default=str),
                         encoding="utf-8")
    print(f"Bảng khớp XD: {len(rows)} dòng. Áp được {len(assign)}, còn {len(left)} dòng 'left' để Claude gán. Cảnh báo {len(warn)}.")
    for w in warn[:20]:
        print("  !", w)
    print("Đã ghi", out)


def cmd_export(args):
    f, khop, ma, names = args[0], args[1], args[2], args[3:]
    wf, wv = _load2(f)
    tvv = wv["TVGIA"]
    new = []
    for n, b in _sheets_with_block(wf, names):
        for x in block_rows(wf[n], wv[n], b):
            if x["kind"] != "work" or x["status"] not in ("da_gan",):
                continue
            m = re.match(r"^=\s*'?TVGIA'?!\$?B\$?(\d+)$", str(x["f_vct"] or ""))
            mn = re.match(r"^=\s*'?TVGIA'?!\$?B\$?(\d+)$", str(x["f_nct"] or ""))
            tv = int(m.group(1)) if m else None
            tn = int(mn.group(1)) if mn else None
            plus = None
            if isinstance(x["f_vdg"], str):
                gs = [int(g) for g in re.findall(r"TVGIA'?!\$?G\$?(\d+)", x["f_vdg"])]
                extra = [g for g in gs if g != tv]
                if "VLOOKUP" in x["f_vdg"].upper():
                    extra = gs
                plus = extra[0] if extra else None
            cv = re.search(r"\)\s*\*\s*([\d.]+)\s*$", str(x["f_vdg"] or ""))
            cn = re.search(r"\)\s*\*\s*([\d.]+)\s*$", str(x["f_ndg"] or ""))
            new.append({"Dự án": ma, "Sheet": n, "Dòng": x["row"], "STT": x["stt"], "Mô tả CĐT": x["desc"], "ĐVT": x["unit"],
                        "Khóa": khop_key(x["desc"], x["unit"]),
                        "Tên TVGIA": tvv.cell(tv or tn, 2).value if (tv or tn) else None,
                        "ĐVT TVGIA": tvv.cell(tv or tn, 3).value if (tv or tn) else None,
                        "Dòng TVGIA (bản lưu)": tv or tn,
                        "Cộng thêm VL (tên TVGIA)": tvv.cell(plus, 2).value if plus else None,
                        "NC từ (tên TVGIA)": tvv.cell(tn, 2).value if (tn and tn != tv) else None,
                        "HS VL": float(cv.group(1)) if cv else 1, "HS NC": float(cn.group(1)) if cn else 1,
                        "Có VL": "Có" if tv else "Không", "Có NC": "Có" if tn else "Không",
                        "Độ tin cậy": x["conf"], "Ghi chú gán giá": x["gnote"], "Ngày lưu": f"{_dt.date.today():%d/%m/%Y}"})
    if os.path.exists(khop):
        wb = openpyxl.load_workbook(khop)
    else:
        wb = openpyxl.Workbook()
        wb.active.title = "Khop_XD"
    if "Khop_XD" not in wb.sheetnames:
        wb.create_sheet("Khop_XD")
    ws = wb["Khop_XD"]
    if ws.max_row < 1 or ws.cell(1, 1).value != KHOP_COLS[0]:
        for i, h in enumerate(KHOP_COLS, 1):
            ws.cell(1, i, h)
    hdr = [c.value for c in ws[1]]
    pos = {}
    for r in range(2, ws.max_row + 1):
        key = (str(ws.cell(r, hdr.index("Dự án") + 1).value), ws.cell(r, hdr.index("Sheet") + 1).value,
               ws.cell(r, hdr.index("Dòng") + 1).value)
        pos[key] = r
    cmpk = ["Khóa", "Tên TVGIA", "Cộng thêm VL (tên TVGIA)", "NC từ (tên TVGIA)", "HS VL", "HS NC", "Có VL", "Có NC"]
    n_new = n_chg = n_same = 0
    changed = []
    for d in new:
        key = (str(d["Dự án"]), d["Sheet"], d["Dòng"])
        if key in pos:
            r = pos[key]
            old = {h: ws.cell(r, hdr.index(h) + 1).value for h in cmpk}
            if all(str(old[h] or "") == str(d[h] or "") for h in cmpk):
                n_same += 1
                continue
            n_chg += 1
            changed.append(f"{d['Sheet']}!{d['Dòng']}: {old['Tên TVGIA']} -> {d['Tên TVGIA']}")
        else:
            r = ws.max_row + 1
            n_new += 1
        for h, v in d.items():
            ws.cell(r, hdr.index(h) + 1, v)
    wb.save(khop)
    print(f"Bảng khớp XD {khop}: MỚI {n_new}, ĐỔI {n_chg}, GIỮ {n_same} (tổng {ws.max_row - 1} dòng).")
    for c in changed[:30]:
        print("  ĐỔI", c)


# ----------------------------------------------------------------------------
# write
# ----------------------------------------------------------------------------
def cmd_write(args):
    f, out, ap = args[0], args[1], args[2]
    data = json.loads(Path(ap).read_text(encoding="utf-8"))
    items = data.get("assign", []) if isinstance(data, dict) else data
    wf, wv = _load2(f)
    if "TVGIA" not in wv.sheetnames:
        raise SystemExit("File chưa có TVGIA -> dừng.")
    tvv, idx = tv_index(wv)
    dup_rows = {r for rs in idx.values() if len(rs) > 1 for r in rs}
    pkg = Package(f)
    st = Styles(pkg)
    sheets, blocks = {}, {}
    seen = set()
    rep = {"lib": 0, "web": 0, "bo_qua": [], "canh_bao": [], "ep_tham_chieu_truc_tiep": 0}

    def tvn(r, c):
        return tvv.cell(r, c).value

    for it in items:
        n, r = it["sheet"], int(it["row"])
        if (n, r) in seen:
            rep["canh_bao"].append(f"{n}!{r}: gán hai lần, lấy mục sau cùng")
        seen.add((n, r))
        if n not in blocks:
            if n not in wf.sheetnames:
                rep["bo_qua"].append(f"{n}!{r}: không có sheet")
                continue
            blocks[n] = find_block(wf[n])
            sheets[n] = Sheet(pkg.parts[pkg.sheet_part(n)])
        b = blocks[n]
        if not b:
            rep["bo_qua"].append(f"{n}!{r}: sheet chưa có khối XD")
            continue
        sh, c0 = sheets[n], b["c0"]
        C = lambda k: L(c0 + k)  # noqa
        fv = wf[n].cell(r, c0 + C_VTT).value
        if not (isinstance(fv, str) and fv.upper().startswith("=N(")):
            rep["bo_qua"].append(f"{n}!{r}: không phải dòng công việc của khối XD")
            continue
        old_vdg, old_ndg = wf[n].cell(r, c0 + C_VDG).value, wf[n].cell(r, c0 + C_NDG).value
        if not it.get("force") and (is_num(old_vdg) and wv[n].cell(r, c0 + C_SRC).value != SRC_WEB or is_num(old_ndg)):
            rep["bo_qua"].append(f"{n}!{r}: có số cứng ở Đơn giá (giá thầu phụ?) -> giữ nguyên, không ghi đè")
            continue
        boq_unit = wv[n].cell(r, c0 + C_UNIT).value
        src = SRC_NONE
        notes = [it.get("note")] if it.get("note") else []
        # ---- VL
        tv = it.get("tv")
        if it.get("vl") is False:
            sh.set(r, c0 + C_VCT)
            sh.set(r, c0 + C_VDG)
        elif it.get("web_vl") is not None:
            sh.set(r, c0 + C_VCT)
            sh.set(r, c0 + C_VDG, value=float(it["web_vl"]), s=st.get("web_num"))
            src = SRC_WEB
            rep["web"] += 1
            if it.get("web_src"):
                notes.append("Giá tham khảo: " + str(it["web_src"]))
        elif tv:
            tv = int(tv)
            if tvn(tv, 2) in (None, ""):
                rep["bo_qua"].append(f"{n}!{r}: TVGIA dòng {tv} không có tên")
                continue
            sh.set(r, c0 + C_VCT, f=f"=TVGIA!B{tv}", cache=tvn(tv, 2))
            direct = tv in dup_rows or it.get("direct")
            rep["ep_tham_chieu_truc_tiep"] += 1 if tv in dup_rows else 0
            e = f"TVGIA!G{tv}" if direct else f"VLOOKUP({C(C_VCT)}{r},TVGIA!$B:$J,6,0)"
            val = to_num(tvn(tv, 7)) or 0
            if it.get("plus_vl"):
                p = int(it["plus_vl"])
                e += f"+TVGIA!G{p}"
                val += to_num(tvn(p, 7)) or 0
            if to_num(it.get("coef_vl")) not in (None, 1.0):
                e = f"({e})*{to_num(it['coef_vl']):g}"
                val *= to_num(it["coef_vl"])
                notes.append(f"quy đổi VL ×{to_num(it['coef_vl']):g}")
            sh.set(r, c0 + C_VDG, f="=" + e, cache=val, s=st.get("num0"))
            src = SRC_LIB
            u2 = tvn(tv, 3)
            if norm_unit(boq_unit) and norm_unit(u2) and norm_unit(boq_unit) != norm_unit(u2) and not it.get("coef_vl"):
                rep["canh_bao"].append(f"{n}!{r}: ĐVT BOQ '{boq_unit}' ≠ TVGIA dòng {tv} '{u2}'")
        # ---- NC
        tn = it.get("tv_nc") or tv if it.get("web_vl") is None else it.get("tv_nc")
        if it.get("nc") is False or not tn:
            if it.get("nc") is False:
                sh.set(r, c0 + C_NCT)
                sh.set(r, c0 + C_NDG)
        else:
            tn = int(tn)
            sh.set(r, c0 + C_NCT, f=f"=TVGIA!B{tn}", cache=tvn(tn, 2))
            direct = tn in dup_rows or it.get("direct")
            e = f"TVGIA!J{tn}" if direct else f"VLOOKUP({C(C_NCT)}{r},TVGIA!$B:$J,9,0)"
            val = to_num(tvn(tn, 10)) or 0
            if to_num(it.get("coef_nc")) not in (None, 1.0):
                e = f"({e})*{to_num(it['coef_nc']):g}"
                val *= to_num(it["coef_nc"])
                notes.append(f"quy đổi NC ×{to_num(it['coef_nc']):g}")
            sh.set(r, c0 + C_NDG, f="=" + e, cache=val, s=st.get("num0"))
            if src == SRC_NONE:
                src = SRC_LIB
            if it.get("tv_nc") and int(it["tv_nc"]) != (tv or 0):
                notes.append(f"NC lấy từ TVGIA dòng {tn}")
        if src == SRC_LIB:
            rep["lib"] += 1
        sh.set(r, c0 + C_SRC, value=src)
        sh.set(r, c0 + C_CONF, value=it.get("conf") or "")
        sh.set(r, c0 + C_GNOTE, value="; ".join(x for x in notes if x))
    for n, sh in sheets.items():
        pkg.parts[pkg.sheet_part(n)] = sh.tobytes()
    st.save()
    pkg.drop_calcchain()
    pkg.save(out)
    Path(out + ".write.json").write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Đã ghi {out}: giá thư viện {rep['lib']} dòng, giá tham khảo {rep['web']} dòng, "
          f"ép tham chiếu trực tiếp (tên trùng) {rep['ep_tham_chieu_truc_tiep']}, bỏ qua {len(rep['bo_qua'])}, cảnh báo {len(rep['canh_bao'])}.")
    for x in rep["bo_qua"][:30] + rep["canh_bao"][:30]:
        print("  !", x)


# ----------------------------------------------------------------------------
CMDS = {"checklib": (cmd_checklib, 1), "auto": (cmd_auto, 2), "build": (cmd_build, 3), "todo": (cmd_todo, 2),
        "recalc": (cmd_recalc, 2), "cmp": (cmd_cmp, 2), "tv": (cmd_tv, 2), "apply": (cmd_apply, 4),
        "write": (cmd_write, 3), "export": (cmd_export, 3)}


def main():
    if len(sys.argv) < 2 or sys.argv[1] in ("-h", "--help") or sys.argv[1] not in CMDS:
        print(__doc__)
        print("Phiên bản:", VERSION)
        return 0 if len(sys.argv) >= 2 and sys.argv[1] in ("-h", "--help") else 2
    fn, need = CMDS[sys.argv[1]]
    args = sys.argv[2:]
    if len([a for a in args if not a.startswith("--")]) < need:
        print(f"Thiếu tham số cho '{sys.argv[1]}'. Xem: python3 boq_xd_tools.py --help")
        return 2
    rc = fn(args)
    return rc or 0


if __name__ == "__main__":
    sys.exit(main())
